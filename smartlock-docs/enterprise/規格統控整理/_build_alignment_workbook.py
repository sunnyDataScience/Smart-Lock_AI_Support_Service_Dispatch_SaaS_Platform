#!/usr/bin/env python3
"""Build the lightweight Excel alignment template used before Plane intake.

The workbook deliberately stops at definition and design.  Test execution,
append-only results, defects, and business sign-off belong in Plane.
"""

from __future__ import annotations

import re
import zipfile
from collections import defaultdict
from pathlib import Path
from typing import Iterable

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

from _task_breakdown import build_enablers, build_tasks


FONT = "Noto Sans CJK TC"
NAVY = "1F3864"
WHITE = "FFFFFF"
HUMAN = "FFF2CC"
DERIVED = "E7E6E6"
BLOCKED = "FCE4D6"
# 第四種底色：生成器寫了草稿，但人**必須**覆寫才算數。
# 沒有這個色，Task 的「完成定義」只能二選一：留白（表看起來沒做完）或塗灰
# （看起來是生成欄，於是沒人去改）。兩個都會讓那一欄失去意義。
DRAFT = "DDEBF7"
THIN = Side(style="thin", color="D9E1F2")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")


SHEET_SPECS = {
    "01_需求收斂": "4472C4",
    "02_需求與驗收": "548235",
    "03_交付切片": "BF9000",
    "04_測試設計": "7030A0",
    "05_技術地基": "808080",
}


def _unique(values: Iterable[str]) -> list[str]:
    seen: set[str] = set()
    result: list[str] = []
    for value in values:
        value = str(value or "").strip()
        if value and value not in seen:
            seen.add(value)
            result.append(value)
    return result


def _joined(values: Iterable[str], separator: str = "; ") -> str:
    return separator.join(_unique(values))


def _new_book() -> Workbook:
    wb = Workbook()
    wb.remove(wb.active)
    wb.properties.creator = "Smart Lock alignment workbook generator"
    wb.properties.title = "Smart Lock 需求收斂與驗收設計"
    wb.properties.subject = "Excel 快速對焦；Plane 承接排程、版本、執行證據與簽核"
    wb.properties.description = (
        "Story 需求收斂、AC 驗收定義、Task 交付切片、Test Case 測試設計，"
        "外加不掛 Story 的技術／品質地基（Enabler）。"
        "不含 Test Run、Test Result、Defect 或業務簽核狀態。"
    )
    return wb


def _sheet(wb: Workbook, name: str, headers: list[tuple[str, int, str, str]]) -> object:
    """Create a sheet. Header tuples are (name, width, kind, comment)."""
    ws = wb.create_sheet(name)
    ws.sheet_properties.tabColor = SHEET_SPECS[name]
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A2"
    for col, (label, width, _kind, note) in enumerate(headers, 1):
        cell = ws.cell(1, col, label)
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(name=FONT, color=WHITE, bold=True, size=10)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BOX
        if note:
            cell.comment = Comment(note, "Codex")
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.row_dimensions[1].height = 36
    return ws


def _write_row(ws, row_number: int, values: list[object], kinds: list[str], height: int = 36) -> None:
    for col, value in enumerate(values, 1):
        cell = ws.cell(row_number, col, value)
        cell.font = Font(name=FONT, size=10)
        cell.alignment = WRAP
        cell.border = BOX
        if kinds[col - 1] == "human":
            cell.fill = PatternFill("solid", fgColor=HUMAN)
        elif kinds[col - 1] == "derived":
            cell.fill = PatternFill("solid", fgColor=DERIVED)
        elif kinds[col - 1] == "draft":
            cell.fill = PatternFill("solid", fgColor=DRAFT)
        elif kinds[col - 1] == "blocked":
            cell.fill = PatternFill("solid", fgColor=BLOCKED)
    ws.row_dimensions[row_number].height = height


def _finish_table(ws, table_name: str, last_row: int, last_col: int) -> None:
    ref = f"A1:{get_column_letter(last_col)}{max(last_row, 2)}"
    tab = Table(displayName=table_name, ref=ref)
    tab.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    ws.add_table(tab)
    ws.auto_filter.ref = ref


def _dropdown(ws, header_names: list[str], header: str, last_row: int, values: str) -> None:
    col = get_column_letter(header_names.index(header) + 1)
    validation = DataValidation(
        type="list",
        formula1=f'"{values}"',
        allow_blank=True,
        showErrorMessage=True,
        errorTitle="值域不符",
        error="請從下拉選單選擇，不要自行發明狀態。",
    )
    ws.add_data_validation(validation)
    validation.add(f"{col}2:{col}{max(last_row, 2)}")


def _requirement_context(model, requirement_id: str) -> tuple[str, str, str, str]:
    scenario_ids = model.rel.scenarios_of(requirement_id)
    scenarios = [model.sc_by_id[sid] for sid in scenario_ids if sid in model.sc_by_id]
    persona_ids = _unique(
        pid for scenario in scenarios for pid in model.rel.personas_of(scenario.sc_id)
    )
    personas = [model.per_by_id[pid] for pid in persona_ids if pid in model.per_by_id]
    actors = _joined([scenario.actor for scenario in scenarios] + [p.name for p in personas])
    pains = _joined(p.pain for p in personas)
    goals = _joined(p.goal for p in personas)
    legacy = _joined(
        f"{scenario.sc_id} {scenario.name}" for scenario in scenarios
    )
    return actors, pains, goals, legacy


def _coverage_by_requirement(model) -> dict[str, set[str]]:
    coverage: dict[str, set[str]] = defaultdict(set)
    for link in model.rel.rq_tc:
        coverage[link["requirement"]].add(link.get("kind") or "")
    return coverage


def _coverage_check(kinds: set[str]) -> str:
    clean = {kind for kind in kinds if kind}
    if not clean:
        return "缺 TC"
    if "happy" not in clean:
        return "缺正常"
    if not clean & {"boundary", "failure", "recovery"}:
        return "缺非正常"
    return "可進開發"


def _build_story_sheet(wb: Workbook, model) -> None:
    headers = [
        ("Epic ID（選填）", 14, "human", "Plane level 0。沒有確認的策略成果就留白。"),
        ("Epic 名稱", 22, "human", "同一 Epic ID 的名稱必須一致。"),
        ("Feature ID（必要）", 16, "human", "Plane level 1，是 Story 的唯一 parent；不得由舊 SC 自動帶入。"),
        ("Feature 名稱", 24, "human", "以功能能力分組，不是測試情境或執行批次。"),
        ("Story ID", 18, "", "Plane 的唯一直接量測點；沿用現有 FR/NFR 編號作穩定外部 ID。"),
        ("Story 標題", 28, "", "一列一個可驗收的價值切片。"),
        ("目標角色", 24, "human", "從舊業務旅程與 Persona 預填，由 PM/SA 確認。"),
        ("問題／痛點", 42, "human", "說明為什麼要做；不寫技術解法。"),
        ("Story 價值句", 52, "human", "功能用『作為…我需要…以便…』；品質需求可直寫業務保障，不強迫套句型。"),
        ("FR/NFR", 10, "derived", "匯入 Plane 時映射 Issue.requirement_kind=functional/quality。"),
        ("需求敘述／目標值", 52, "", "FR 寫系統行為；NFR 寫場景、指標與目標值。"),
        ("範圍內", 40, "human", "本次承諾的行為或品質邊界。"),
        ("範圍外", 32, "human", "本版明確不做什麼。"),
        ("優先序", 10, "human", "P0/P1/P2/P3/未定；匯入時再映射 Plane priority。"),
        ("PM Owner", 16, "human", "負責價值、範圍與優先序。"),
        ("SA Owner", 16, "human", "負責規則清楚且可測。"),
        ("收斂狀態", 13, "human", "只是需求軸；不得由 Task、TC 或簽核狀態自動推進。"),
        ("舊業務旅程參考（不匯入）", 44, "derived", "只保留語意背景；SC 不再是 Feature、Test Run 或簽核單位。"),
    ]
    ws = _sheet(wb, "01_需求收斂", headers)
    kinds = [header[2] for header in headers]
    p0_requirements = model.p0_requirements()
    all_requirements = [("FR", item) for item in model.frs] + [("NFR", item) for item in model.nfrs]
    for row_number, (kind, item) in enumerate(all_requirements, 2):
        actors, pains, goals, legacy = _requirement_context(model, item.req_id)
        target_actor = actors or ("平台／全體使用角色" if kind == "NFR" else "待 PM/SA 確認")
        if kind == "FR":
            value = f"作為 {target_actor}，我需要「{item.name}」，以便 {goals or item.acceptance}"
            requirement = item.flow
            scope_in = item.flow
        else:
            value = f"為了保障 {target_actor} 在「{item.name}」的服務品質，系統需達成 {item.target}"
            requirement = item.target
            scope_in = item.target
        _write_row(ws, row_number, [
            "", "", "", "", item.req_id, item.name, target_actor,
            pains or "待 PM 補業務損失／痛點", value, kind,
            requirement, scope_in, "", "P0" if item.req_id in p0_requirements else "未定",
            "", "", "待釐清", legacy or "global（跨情境品質地板）",
        ], kinds, height=54)
    last_row = len(all_requirements) + 1
    names = [header[0] for header in headers]
    _dropdown(ws, names, "FR/NFR", last_row, "FR,NFR")
    _dropdown(ws, names, "優先序", last_row, "P0,P1,P2,P3,未定")
    _dropdown(ws, names, "收斂狀態", last_row, "待釐清,已對焦,可開發,已凍結")
    _finish_table(ws, "StoryAlignment", last_row, len(headers))


def _build_acceptance_sheet(wb: Workbook, model) -> None:
    headers = [
        ("AC ID（草案）", 24, "", "AC 不是 Plane 獨立物件；匯入時收進 Story 敘述與 TC 追溯資訊。"),
        ("Story ID", 18, "", "必須存在於 01_需求收斂。"),
        ("FR/NFR", 10, "derived", "與 Story requirement_kind 一致。"),
        ("驗收類型", 16, "human", "只分驗收意圖；測試路徑在 04 另列。"),
        ("需求名稱", 30, "", "來自 SRS/NFR 正典。"),
        ("Given／前置", 38, "human", "情境開始前必須成立什麼。"),
        ("When／行為或量測", 48, "human", "觸發行為，或 NFR 的量測方式與資料量。"),
        ("Then／可觀察結果", 52, "human", "明確的過／不過判準。"),
        ("規則／例外", 38, "human", "商業規則、邊界、拒絕或回復行為。"),
        ("指標", 18, "human", "NFR 必填；FR 可留白。"),
        ("目標值", 32, "human", "NFR 必須可判定；沒數字就明寫本版不承諾。"),
        ("驗證方式", 38, "human", "在哪個環境、用什麼資料與量測點驗證。"),
        ("SA Owner", 16, "human", "負責可測表述。"),
        ("業務確認人", 16, "human", "確認這個結果是不是業務要的。"),
        ("驗收設計檢查", 16, "derived", "只顯示 TC 設計是否同時有正常與非正常路徑；不會自動推進需求狀態。"),
        ("可開發（人工）", 16, "human", "PM/SA/QA 對焦後才能選『是』，不由公式帶出。"),
        ("來源追溯", 32, "derived", "舊 FR/NFR 文件行號；不是 Plane object。"),
    ]
    ws = _sheet(wb, "02_需求與驗收", headers)
    kinds = [header[2] for header in headers]
    coverage = _coverage_by_requirement(model)
    all_requirements = [("FR", item) for item in model.frs] + [("NFR", item) for item in model.nfrs]
    for row_number, (kind, item) in enumerate(all_requirements, 2):
        if kind == "FR":
            given = item.precondition
            when = item.flow
            then = item.acceptance
            rule = item.trace
            metric = ""
            target = ""
            verification = ""
            acceptance_type = "功能結果"
            source = f"04_SRS.md:{item.source_line}"
        else:
            given = item.name
            when = item.verification
            then = item.target
            rule = item.tier
            metric = item.category
            target = item.target
            verification = item.verification
            acceptance_type = "品質目標"
            source = f"05_NFR.md:{item.source_line}"
        _write_row(ws, row_number, [
            f"AC-{item.req_id}-01-DRAFT", item.req_id, kind, acceptance_type,
            item.name, given, when, then, rule, metric, target, verification,
            "", "", _coverage_check(coverage.get(item.req_id, set())), "", source,
        ], kinds, height=58)
    last_row = len(all_requirements) + 1
    names = [header[0] for header in headers]
    _dropdown(ws, names, "FR/NFR", last_row, "FR,NFR")
    _dropdown(ws, names, "驗收類型", last_row, "功能結果,業務規則,例外/邊界,品質目標")
    _dropdown(ws, names, "可開發（人工）", last_row, "是,否")
    _finish_table(ws, "AcceptanceCriteria", last_row, len(headers))


def _build_task_sheet(wb: Workbook, model) -> None:
    """一列＝一個 Task（Plane level 3），從 Story 往下拆，不是舊 WBS 的殘骸。

    2026-08-07 改版理由：舊版一列一個舊 WBS 工作包，16 列對 171 個 Story——
    Task 少於 Story 一個數量級，在任何拆解模型下都不成立（Story 拆成 Task，數量只會
    變多）。查證那 16 列後確認它們本來就不是 Task：8 列不掛任何 Story，3 列橫跨多個
    Story。它們改列在 `05_技術地基`，正名為 Enabler。拆解規則見 `_task_breakdown.py`。
    """
    headers = [
        ("Task ID（草案）", 22, "derived",
         "格式 TASK-<Story ID>-<序>。一列＝一個工程層面的工作，不是一個舊 WBS 工作包。"
         "『草案』＝尚未進 Plane；確認後才建立 level 3 work item。"),
        ("Story Parent（唯一）", 18, "derived",
         "Task 只能有一個 Story parent，且必存在於 01_需求收斂。"
         "拆不出唯一 parent 的東西不是 Task——那是 Enabler，去 05_技術地基。"),
        ("價值線 Epic", 26, "derived",
         "由 Story 反查 SC → 價值線；全域需求歸 E-GLB。"
         "『待定』＝該 Story 還沒掛旅程，是 01 的 Feature 欄整欄空白的同一個缺口。"),
        ("任務類型", 14, "derived",
         "這張卡屬於哪個工程層面：資料模型／API 契約／核心邏輯／事件／授權／前端／整合。"
         "同一個 Story 穿過幾層就有幾張卡。"),
        ("交付內容", 52, "draft",
         "工程師明天可以開工的切片，附觸發它的正典原文。SD 需在此具體化到可估點，"
         "但不得改寫『依據』那一行——那是這張卡存在的證據。"),
        ("Owner", 12, "human", "SD/DEV 實作 owner。生成器只給層面預設（BE/FE），排到人是人的事。"),
        ("State", 12, "human", "匯入前先讀 Plane project states，不在 Excel 猜 UUID。"),
        ("Cycle", 14, "human", "時間箱，與 parent 無關；Excel 不預設，排程是 Plane 的職責。"),
        ("Module(s)", 26, "derived", "子系統 / 能力群，由 Story 的 prefix 與 MODULES 推得。"),
        ("Milestone", 24, "derived", "由 _canon.phase_for(Story) 推得的交付檢查點。"),
        ("相依 Task ID", 22, "derived",
         "同一 Story 內的層面順序：資料 → 契約 → 核心 → 其餘。跨 Story 相依請人工補。"),
        ("完成定義（草案）", 52, "draft",
         "**工程判準**，不等於 TC passed，更不等於業務驗收通過。"
         "核心邏輯用 04_SRS 的後置與驗收原文；其餘是層面樣板，SD 必須替換成本卡的具體判準。"),
        ("驗證線索（TC）", 26, "derived",
         "驗這個 Story 的 TC。**TC 不是 Task 的完成定義**——它驗的是 Story 對不對，"
         "不是這張卡做完沒。放這裡只是讓 SD 知道下游會怎麼驗。"),
        ("拆解檢查", 18, "derived", "生成器看得出來的缺口：待掛旅程／待補 TC／待補驗收。"),
        ("可匯入（人工）", 14, "human", "交付內容與完成定義都被 SD 具體化過，才可以選『是』。"),
        ("來源追溯", 44, "derived",
         "04_SRS 行號 | 觸發這張卡的關鍵詞 | 對應元件。"
         "關鍵詞判錯就是規則錯，請改 _task_breakdown.py 的 ASPECTS，不要手改這一欄。"),
    ]
    ws = _sheet(wb, "03_交付切片", headers)
    kinds = [header[2] for header in headers]
    tasks = build_tasks(model)
    for row_number, task in enumerate(tasks, 2):
        _write_row(ws, row_number, [
            task.task_id, task.story, task.epic, task.label, task.content,
            task.owner, "", "", task.module, task.milestone, task.depends,
            task.done, task.verified_by, task.check, "", task.trace,
        ], kinds, height=64)
    last_row = len(tasks) + 1
    names = [header[0] for header in headers]
    _dropdown(ws, names, "Owner", last_row, "BE,FE,BE+FE,OPS,QA,PM")
    _dropdown(ws, names, "可匯入（人工）", last_row, "是,否")
    _finish_table(ws, "DeliverySlices", last_row, len(headers))


def _build_enabler_sheet(wb: Workbook, model, base_dir: Path) -> None:
    """不掛 Story 的地基工作。它們是真工作，只是**不是 Task**。

    兩類：舊 WBS 留下的技術地基（CD、雲端拓撲、cutover…），與 106 條 NFR 依驗證形態
    收斂出的品質地基。後者刻意不逐條產卡——一套 k6 骨架服務 45 條門檻量測型 NFR，
    逐條產會得到 106 張互相重複的卡。
    """
    headers = [
        ("Enabler ID", 20, "derived", "ENB-WBS-* 來自舊 WBS；ENB-NFR-* 由 NFR 驗證形態收斂。"),
        ("名稱", 42, "derived", "這件地基工作本身。"),
        ("類型", 12, "derived", "技術地基（跑得起來）／品質地基（證得出來）。"),
        ("為何不是 Task", 46, "derived",
         "Task 的定義是『某個 Story 的零件』。掛不到唯一 Story 的東西放進 Task 表，"
         "會讓 Story 覆蓋率與工程進度兩個數字同時失真。"),
        ("涵蓋範圍", 40, "derived", "它服務哪些 Story 或 NFR；空的代表全體共用。"),
        ("Owner", 14, "human", "地基通常是 OPS/平台組，不是功能組。"),
        ("Milestone", 26, "derived", "來自舊 WBS 或預設 M1。"),
        ("完成定義（草案）", 50, "draft", "工程判準；同樣需要負責人具體化。"),
        ("可匯入（人工）", 14, "human", "進 Plane 時建議掛在 E-GLB 底下，不掛任何旅程。"),
        ("來源追溯", 40, "derived", "舊 WBS ID 與處置，或 NFR 驗證形態。"),
    ]
    ws = _sheet(wb, "05_技術地基", headers)
    kinds = [header[2] for header in headers]
    enablers = build_enablers(model, base_dir)
    for row_number, enabler in enumerate(enablers, 2):
        _write_row(ws, row_number, [
            enabler.enabler_id, enabler.name, enabler.kind, enabler.why,
            enabler.covers, enabler.owner, enabler.milestone, enabler.done,
            "", enabler.trace,
        ], kinds, height=58)
    last_row = len(enablers) + 1
    names = [header[0] for header in headers]
    _dropdown(ws, names, "可匯入（人工）", last_row, "是,否")
    _finish_table(ws, "TechnicalEnablers", last_row, len(headers))


ASPECT_BY_NFR_CATEGORY = {
    "Perf": "效能",
    "Scal": "效能",
    "Avail": "可靠性",
    "Rel": "可靠性",
    "SLA": "可靠性",
    "Sch": "可靠性",
    "Obs": "可靠性",
    "Sec": "安全",
    "Priv": "安全",
    "Comp": "安全",
    "Aud": "安全",
    "DQ": "資料/整合",
    "Rep": "資料/整合",
    "A11y": "可用性/無障礙",
    "Maint": "維護/交付",
    "PUB": "維護/交付",
    "DORA": "維護/交付",
}


def _case_aspect(model, case, requirement_ids: list[str]) -> str:
    aspects: set[str] = set()
    if any(requirement_id.startswith("FR-") for requirement_id in requirement_ids):
        if "權限" in (case.kind or ""):
            aspects.add("權限")
        elif "冪等" in (case.kind or ""):
            aspects.add("業務規則")
        else:
            aspects.add("功能")
    for requirement_id in requirement_ids:
        nfr = model.nfr_by_id.get(requirement_id)
        if nfr:
            aspects.add(ASPECT_BY_NFR_CATEGORY.get(nfr.category, "維護/交付"))
    return next(iter(aspects)) if len(aspects) == 1 else ("待拆分" if aspects else "")


def _build_test_sheet(wb: Workbook, model) -> None:
    headers = [
        ("TC ID", 18, "", "Plane Test Case 外部穩定識別；建議放在 title 前綴與 tag。"),
        ("Linked Story IDs（M:N）", 38, "derived", "用 ; 分隔；匯入時逐一建 TestCaseWorkItemLink。"),
        ("AC IDs（M:N）", 42, "derived", "需對應 02 中的 AC；舊資料為決定性草案 ID。"),
        ("測試領域／Folder", 30, "derived", "舊 TS 只當領域分組，不是 Test Scenario 物件。"),
        ("測試情境（草案）", 44, "human", "回答『在什麼狀況下驗』；不含實際步驟與證據。"),
        ("路徑類型", 14, "human", "只有 happy/boundary/failure/recovery；多路徑的舊案例應拆分。"),
        ("驗證面向", 18, "human", "與路徑正交；回答這個案例在驗什麼性質。"),
        ("前置條件／資料", 38, "human", "具體帳號、狀態、環境與輸入資料。"),
        ("操作步驟", 48, "human", "可執行的具體步驟；不是業務旅程大綱。"),
        ("預期結果", 54, "human", "每一步或案例結束時的可觀察判準。"),
        ("優先序", 10, "human", "P0/P1/P2/P3/未定。"),
        ("QA Owner", 14, "human", "負責測試設計；實際執行人與結果進 Plane。"),
        ("設計檢查", 16, "derived", "只顯示關聯/路徑/面向/前置的設計問題。"),
        ("可匯入（人工）", 16, "human", "人工確認後才能匯入。"),
        ("來源追溯", 34, "derived", "舊 case kind/章節/文件行號；不包含舊 SC/UAT 物件。"),
    ]
    ws = _sheet(wb, "04_測試設計", headers)
    kinds = [header[2] for header in headers]
    edges_by_case: dict[str, list[dict]] = defaultdict(list)
    for edge in model.rel.rq_tc:
        edges_by_case[edge["case"]].append(edge)
    for row_number, case in enumerate(model.cases, 2):
        edges = edges_by_case.get(case.tc_id, [])
        requirement_ids = _unique(edge["requirement"] for edge in edges)
        ac_ids = [f"AC-{requirement_id}-01-DRAFT" for requirement_id in requirement_ids]
        path_kinds = _unique(edge.get("kind", "") for edge in edges)
        path_type = path_kinds[0] if len(path_kinds) == 1 and path_kinds[0] in {
            "happy", "boundary", "failure", "recovery"
        } else ""
        aspect = _case_aspect(model, case, requirement_ids)
        ts_ids = _unique(edge.get("ts", "") for edge in edges)
        folders = _joined(
            f"{ts_id} {model.ts.get(ts_id, {}).get('name', '')}" for ts_id in ts_ids
        ) or case.heading
        notes = _unique(edge.get("note", "") for edge in edges)
        scenario = _joined(notes) or (
            "待重新定義：這不是把舊 SC/TS 改名" if not requirement_ids else case.heading
        )
        checks: list[str] = []
        if not requirement_ids:
            checks.append("待重新綁定")
        if len(path_kinds) != 1 or not path_type:
            checks.append("待拆分路徑")
        if aspect == "待拆分":
            checks.append("待拆分面向")
        if not case.precondition:
            checks.append("待補前置")
        design_check = "、".join(checks) or "可匯入"
        _write_row(ws, row_number, [
            case.tc_id, "; ".join(requirement_ids), "; ".join(ac_ids), folders,
            scenario, path_type, aspect, case.precondition, case.steps, case.expected,
            case.priority, "", design_check, "",
            f"20_Test_Cases.md:{case.source_line} | 舊 kind: {case.kind or '空白'}",
        ], kinds, height=64)
    last_row = len(model.cases) + 1
    names = [header[0] for header in headers]
    _dropdown(ws, names, "路徑類型", last_row, "happy,boundary,failure,recovery")
    _dropdown(
        ws,
        names,
        "驗證面向",
        last_row,
        "功能,業務規則,權限,資料/整合,效能,可靠性,安全,可用性/無障礙,維護/交付,待拆分",
    )
    _dropdown(ws, names, "優先序", last_row, "P0,P1,P2,P3,未定")
    _dropdown(ws, names, "可匯入（人工）", last_row, "是,否")
    _finish_table(ws, "TestCaseDesign", last_row, len(headers))


def _column_values(ws, header: str) -> list[str]:
    headers = {cell.value: index + 1 for index, cell in enumerate(ws[1])}
    column = headers[header]
    return [str(ws.cell(row, column).value or "").strip() for row in range(2, ws.max_row + 1)]


def _validate_workbook(wb: Workbook, model) -> None:
    errors: list[str] = []
    expected_sheets = list(SHEET_SPECS)
    if wb.sheetnames != expected_sheets:
        errors.append(f"sheet order {wb.sheetnames!r} != {expected_sheets!r}")

    story_ids = _column_values(wb["01_需求收斂"], "Story ID")
    if len(story_ids) != len(model.frs) + len(model.nfrs):
        errors.append("Story row count does not match FR+NFR source count")
    if not all(story_ids) or len(story_ids) != len(set(story_ids)):
        errors.append("Story IDs must be non-empty and unique")

    ac_ids = _column_values(wb["02_需求與驗收"], "AC ID（草案）")
    ac_story_ids = _column_values(wb["02_需求與驗收"], "Story ID")
    if len(ac_ids) != len(set(ac_ids)) or set(ac_story_ids) - set(story_ids):
        errors.append("AC IDs must be unique and every AC must reference a known Story")

    # Task 層的守線。這幾條是 2026-08-07 那次「16 個 Task 對 171 個 Story」跑掉的原因——
    # 舊 validator 只檢查了 ID 唯一與 parent 存在，對「數量關係整個反了」完全沉默。
    task_ids = _column_values(wb["03_交付切片"], "Task ID（草案）")
    task_parents = _column_values(wb["03_交付切片"], "Story Parent（唯一）")
    task_depends = _column_values(wb["03_交付切片"], "相依 Task ID")
    if not all(task_ids) or len(task_ids) != len(set(task_ids)):
        errors.append("Task draft IDs must be non-empty and unique")
    if not all(task_parents):
        errors.append("Every Task must name exactly one Story parent")
    if set(task_parents) - set(story_ids):
        errors.append("A Task references an unknown Story parent")
    # Story 拆成 Task，數量只會變多。反過來就代表那張表裝的不是 Task。
    parented_stories = set(task_parents)
    if len(task_ids) <= len(parented_stories):
        errors.append(
            f"Task count ({len(task_ids)}) must exceed the number of parented Stories "
            f"({len(parented_stories)}) — a Task table smaller than its Story set is not "
            f"a decomposition"
        )
    missing = {fr.req_id for fr in model.frs} - parented_stories
    if missing:
        errors.append(f"{len(missing)} FR Stories have no Task at all: {sorted(missing)[:5]}")
    for task_id, depends in zip(task_ids, task_depends):
        for dependency in (item.strip() for item in depends.split(";") if item.strip()):
            if dependency == task_id:
                errors.append(f"Task {task_id} depends on itself")
            elif dependency not in set(task_ids):
                errors.append(f"Task {task_id} depends on unknown {dependency}")

    enabler_ids = _column_values(wb["05_技術地基"], "Enabler ID")
    if not all(enabler_ids) or len(enabler_ids) != len(set(enabler_ids)):
        errors.append("Enabler IDs must be non-empty and unique")
    if set(enabler_ids) & set(task_ids):
        errors.append("An Enabler ID collides with a Task ID")

    tc_ids = _column_values(wb["04_測試設計"], "TC ID")
    if len(tc_ids) != len(model.cases) or len(tc_ids) != len(set(tc_ids)):
        errors.append("TC row count/uniqueness does not match the source library")
    linked_story_cells = _column_values(wb["04_測試設計"], "Linked Story IDs（M:N）")
    linked_story_ids = [
        item.strip()
        for cell in linked_story_cells
        for item in cell.split(";")
        if item.strip()
    ]
    if len(linked_story_ids) != len(model.rel.rq_tc):
        errors.append("Story–TC link count does not match rq_verified_by_tc.yaml")
    if set(linked_story_ids) - set(story_ids):
        errors.append("A Test Case references an unknown Story ID")
    linked_ac_cells = _column_values(wb["04_測試設計"], "AC IDs（M:N）")
    linked_ac_ids = {
        item.strip()
        for cell in linked_ac_cells
        for item in cell.split(";")
        if item.strip()
    }
    if linked_ac_ids - set(ac_ids):
        errors.append("A Test Case references an unknown AC ID")

    if errors:
        raise ValueError("Alignment workbook preflight failed: " + "; ".join(errors))


# openpyxl emits comment parts and worksheet relationships in a layout Excel
# accepts but several JS readers do not.  This workbook is the only one of the
# four carrying tables and cell comments, so it is the only one that trips them.
#
# Three divergences, all fixed by rewriting the archive in place:
#   xl/comments/commentN.xml           -> xl/commentsN.xml
#   xl/drawings/commentsDrawingN.vml   -> xl/drawings/vmlDrawingN.vml
#   Target="/xl/..."  (absolute)       -> Target="../..."  (relative)
#
# The relative form is what Excel itself writes; openpyxl round-trips either.
_PART_RENAMES = (
    (re.compile(r"xl/comments/comment(\d+)\.xml"), r"xl/comments\1.xml"),
    (re.compile(r"xl/drawings/commentsDrawing(\d+)\.vml"), r"xl/drawings/vmlDrawing\1.vml"),
)


def _normalize_ooxml_layout(path: Path) -> None:
    """Rewrite the saved workbook into the conventional OOXML part layout."""
    with zipfile.ZipFile(path) as zin:
        items = [(item, zin.read(item.filename)) for item in zin.infolist()]

    renames: dict[str, str] = {}
    for item, _ in items:
        for pattern, replacement in _PART_RENAMES:
            if pattern.fullmatch(item.filename):
                renames[item.filename] = pattern.sub(replacement, item.filename)
                break

    def _relative(match: "re.Match[str]") -> str:
        target = renames.get(match.group(1).lstrip("/"), match.group(1).lstrip("/"))
        # worksheet rels live in xl/worksheets/_rels/, so xl/ resolves to ../
        return f'Target="../{target[3:] if target.startswith("xl/") else target}"'

    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zout:
        for item, data in items:
            name = item.filename
            if name == "[Content_Types].xml":
                body = data.decode("utf-8")
                for old, new in renames.items():
                    body = body.replace(f'PartName="/{old}"', f'PartName="/{new}"')
                data = body.encode("utf-8")
            elif name.startswith("xl/worksheets/_rels/") and name.endswith(".rels"):
                body = re.sub(r'Target="(/[^"]+)"', _relative, data.decode("utf-8"))
                data = body.encode("utf-8")
            zout.writestr(renames.get(name, name), data)


def build_alignment_workbook(model, output: Path, base_dir: Path) -> None:
    wb = _new_book()
    _build_story_sheet(wb, model)
    _build_acceptance_sheet(wb, model)
    _build_task_sheet(wb, model)
    _build_test_sheet(wb, model)
    _build_enabler_sheet(wb, model, base_dir)
    _validate_workbook(wb, model)
    wb.save(output)
    _normalize_ooxml_layout(output)
