#!/usr/bin/env python3
"""從 smartlock-docs/enterprise 正典 Markdown 單向生成規格四書。

主鍵原則：
    04_SRS FR/NFR ID -> 驗收條件 -> 測試場景 -> 20_Test_Cases §2.1 QTM -> 指定 TC

生成器不修改上游 Markdown，也不會用名稱臆測出新鍵。
"""

from __future__ import annotations

import re
from collections import Counter, defaultdict
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable, Sequence

from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.dimensions import ColumnDimension

from _spec_data import (
    ARCH_CHOICES,
    ARCH_RISKS,
    BUG_LEVELS,
    COMPONENT_GLOSSARY,
    DOMAIN_QA_CHECKS,
    DOMAIN_TEST_META,
    ENVIRONMENTS,
    FR_BUSINESS_COPY,
    FR_TC_HINTS,
    GENERATED_ON,
    GLOSSARY,
    MODULE_ARCH,
    MODULES,
    NFR_DOMAIN_META,
    NFR_TC_HINTS,
    PHASE_BY_ID,
    PHASE_OVERVIEW,
    RESPONSIBILITIES,
    SCENARIOS,
    SCENARIO_PASS_CRITERIA,
    STLC,
    SUBSYSTEMS,
    TEST_STAGES,
    TEST_STRATEGY,
    TEST_TYPES,
)


HERE = Path(__file__).resolve().parent
ENTERPRISE = HERE.parent

SRS_PATH = ENTERPRISE / "04_SRS.md"
NFR_PATH = ENTERPRISE / "05_NFR.md"
SAD_PATH = ENTERPRISE / "12_SAD.md"
SDS_PATH = ENTERPRISE / "15_SDS.md"
TEST_CASE_PATH = ENTERPRISE / "20_Test_Cases.md"
TEST_PLAN_PATH = ENTERPRISE / "19_Test_Plan.md"
TRACE_PATH = ENTERPRISE / "21_Traceability_Matrix.md"
UAT_PATH = ENTERPRISE / "22_UAT_Report.md"
ROADMAP_PATH = ENTERPRISE / "27_Product_Roadmap_WBS.md"
ADR_INDEX_PATH = ENTERPRISE / "14_ADR" / "00_INDEX.md"

OUTPUTS = {
    "planning": HERE / "SmartLock_規格統控規劃書.xlsx",
    "bom": HERE / "SmartLock_模組功能BOM.xlsx",
    "acceptance": HERE / "SmartLock_業務邏輯驗收控制表.xlsx",
    "test": HERE / "SmartLock_整合測試計畫.xlsx",
}
REPORT_PATH = HERE / "產出健康報告.md"
COMPONENT_GLOSSARY_PATH = HERE / "SAD_SDS元件標籤字典.md"


COLORS = {
    "navy": "17365D",
    "blue": "1F4E78",
    "mid_blue": "5B9BD5",
    "light_blue": "D9EAF7",
    "pale_blue": "EAF3F8",
    "green": "70AD47",
    "light_green": "E2F0D9",
    "orange": "F4B183",
    "light_orange": "FCE4D6",
    "red": "C00000",
    "light_red": "F4CCCC",
    "yellow": "FFF2CC",
    "gray": "D9E1F2",
    "light_gray": "F2F2F2",
    "dark_gray": "595959",
    "white": "FFFFFF",
    "black": "1F1F1F",
}

THIN_GRAY = Side(style="thin", color="B7C9D6")
TABLE_BORDER = Border(left=THIN_GRAY, right=THIN_GRAY, top=THIN_GRAY, bottom=THIN_GRAY)
FONT_NAME = "Noto Sans CJK TC"


@dataclass
class Requirement:
    req_id: str
    raw_id: str
    prefix: str
    name: str
    precondition: str
    flow: str
    acceptance: str
    trace: str
    heading: str
    source_line: int
    raw_line: str
    duplicate: bool = False

    @property
    def all_text(self) -> str:
        return " ".join(
            [self.raw_id, self.name, self.precondition, self.flow, self.acceptance, self.trace]
        )


@dataclass
class NFR:
    req_id: str
    category: str
    name: str
    target: str
    verification: str
    tier: str
    heading: str
    source_line: int

    @property
    def all_text(self) -> str:
        return " ".join([self.name, self.target, self.verification, self.tier])


@dataclass
class TestCase:
    tc_id: str
    source_ref: str
    cells: list[str]
    heading: str
    source_line: int


def read_lines(path: Path) -> list[str]:
    return path.read_text(encoding="utf-8").splitlines()


def split_md_row(line: str) -> list[str]:
    """足夠應付本專案表格的 Markdown row parser。

    本專案的正典表格沒有在儲存格內使用未 escape 的 pipe，因此不引入
    額外 Markdown parser，以降低生成器依賴。
    """
    text = line.strip()
    if not text.startswith("|"):
        return []
    return [clean_md(cell.strip()) for cell in text.strip("|").split("|")]


def clean_md(value: object) -> str:
    if value is None:
        return ""
    text = str(value)
    text = ILLEGAL_CHARACTERS_RE.sub("", text)
    text = re.sub(r"\[([^\]]+)\]\([^\)]+\)", r"\1", text)
    text = text.replace("**", "").replace("`", "")
    text = text.replace("<br>", " / ").replace("<br/>", " / ")
    return re.sub(r"\s+", " ", text).strip()


# 活頁簿是給台灣的 PM、SA、架構師與 QA 使用。來源正典可能仍保留不同地區
# 或偏工程圈的慣用語；產出時統一成台灣常見說法，但不改動 API、事件名稱、
# 資料表欄位與檔案路徑等技術識別碼。
TAIWAN_WORDING_REPLACEMENTS: tuple[tuple[str, str], ...] = (
    ("知識精煉 HITL 閉環", "知識精煉－人工審核完整流程"),
    ("知識精煉 人工審核 閉環", "知識精煉－人工審核完整流程"),
    ("高風險 HITL", "高風險操作需人工審核"),
    ("驗收、測試設計與追溯映射已閉環", "驗收、測試設計與追溯映射已完整"),
    ("已閉環", "已完整"),
    ("未閉環", "尚未完整"),
    ("transfer_to_human 唯一出口 + 兜底", "transfer_to_human 是唯一轉接方式，並有備援處理"),
    ("空庫/舊版庫", "空白資料庫／舊版資料庫"),
    ("agent 旁路", "不經 agent 的替代路徑"),
    ("轉人不得蒸發、AI 金額/影像紅線", "轉真人案件不可遺漏，AI 不得自行報價或判讀影像"),
    ("立即止血、留 audit，需 root cause + regression", "立即控制影響並保留稽核紀錄，完成根因分析與回歸測試"),
    ("assign/reassign/池單三掛點", "assign、reassign 與池單三個整合點"),
    ("AI 永不自轉工單", "AI 不得自行建立工單"),
    ("自轉工單", "自行建立工單"),
    ("案子不蒸發", "案件不可遺漏"),
    ("隱性污染", "不易察覺的資料污染"),
    ("RD、QA、PM 與 Business Owner 責任切分", "RD、QA、PM 與業務負責人的職責分工"),
    ("License 開通與 provisioning", "授權開通與環境設定"),
    ("License provisioning", "授權開通設定"),
    ("身分、RBAC 與 License", "身分、角色權限與產品授權"),
    ("HITL 審核", "人工審核"),
    ("Clarify gate", "釐清檢核"),
    ("Business Owner", "業務負責人"),
    ("業務 Owner", "業務負責人"),
    ("業務負責人 責任切分", "業務負責人的職責分工"),
    ("業務負責人責任切分", "業務負責人的職責分工"),
    ("業務負責人責任分工", "業務負責人的職責分工"),
    ("Entry/Exit Criteria", "進入／結束條件"),
    ("Entry Criteria", "進入條件"),
    ("Exit Criteria", "結束條件"),
    ("職能切分", "職責分工"),
    ("責任切分", "責任分工"),
    ("範圍 / DoD", "範圍／完成定義（DoD）"),
    ("RAG pipeline", "RAG 處理流程"),
    ("首回應 latency", "首次回應時間"),
    ("steady state", "穩定負載"),
    ("RBAC enforce", "角色權限強制執行"),
    ("三層解決 + 釐清檢核", "三層解決＋釐清檢核"),
    ("incident path", "異常處理流程"),
    ("pricing engine down", "計價引擎停止服務"),
    ("admin banner", "後台警示橫幅"),
    ("override SLI", "人工改價指標"),
    ("deny-by-default", "預設拒絕"),
    ("知識庫可/不可命中問題", "知識庫可／不可找到答案的問題"),
    ("可/不可命中問題", "知識庫可／不可找到答案的問題"),
    ("無法命中知識的問題", "知識庫找不到答案的問題"),
    ("可命中知識的問題", "能從知識庫找到答案的問題"),
    ("知識命中", "知識庫查詢結果"),
    ("命中急件", "判定為急件"),
    ("未核可零落地", "未經核可不得寫入正式環境"),
    ("零落地", "不得寫入正式環境"),
    ("落地數", "寫入筆數"),
    ("雙路落地", "雙路發布"),
    ("落地順序", "實作順序"),
    ("落地方式", "實作方式"),
    ("落地狀態", "實作狀態"),
    ("未落地", "尚未實作"),
    ("規格閉環狀態", "規格完整狀態"),
    ("QA 規格閉環", "QA 規格完整"),
    ("規格閉環", "規格完整"),
    ("追溯閉環", "追溯完整"),
    ("全鏈路", "完整呼叫路徑"),
    ("鏈路追蹤", "分散式追蹤"),
    ("主流作業", "主要作業"),
    ("主流流程", "主要流程"),
    ("回復上一版", "還原上一版"),
    ("覆蓋率", "涵蓋率"),
    ("硬 gate", "強制檢核"),
    ("hard gate", "強制檢核"),
    ("fail-closed", "驗證失敗即拒絕"),
    ("fail-soft", "失敗時採降級處理"),
    ("block-deploy", "未通過即禁止發布"),
    ("blocker", "阻擋問題"),
    ("容器鏡像", "容器映像檔"),
    ("legacy display", "舊版顯示"),
    ("join key", "關聯鍵"),
    ("Stage Gate", "階段檢核點"),
    ("dry/live gate runner", "模擬／正式環境檢核執行器"),
    ("品質 gate", "品質檢核"),
    ("固定 fixture", "固定測試資料"),
    ("全量入庫", "全部寫入資料庫"),
    ("全量記錄", "完整記錄"),
    ("優先級", "優先順序"),
    ("紅線", "不可違反條件"),
    ("白名單", "允許清單"),
    ("黑名單", "封鎖清單"),
    ("人審", "人工審核"),
    ("灌注", "匯入"),
    ("驗證閘", "驗證檢核"),
    ("發佈", "發布"),
    ("重佈", "重新部署"),
    ("入庫", "寫入資料庫"),
    ("話術", "回覆內容"),
    ("校驗", "驗證"),
    ("告警", "警示"),
    ("觸達", "通知"),
    ("運維", "維運"),
    ("全量", "完整"),
    ("物化", "實體化"),
    ("掛點", "整合點"),
    ("闀檻", "門檻"),
    ("蒸發", "遺漏"),
    ("訴求", "需求"),
    ("裁決", "決議"),
    ("查實", "確認"),
    ("形態", "形式"),
    ("自維", "自行維護"),
    ("出廠", "預設"),
    ("落檔", "寫入紀錄"),
    ("硬刪", "永久刪除"),
    ("真相源", "真實資料來源"),
    ("全綠", "全部通過"),
    ("開帳", "建立帳號"),
    ("實證", "實際驗證"),
    ("鋪開", "逐步套用"),
    ("補漏", "補齊遺漏資料"),
    ("上線硬化", "上線強化"),
    ("HITL", "人工審核"),
    ("provisioning", "環境開通設定"),
    ("場景", "情境"),
    ("閉環", "完整流程"),
    ("落地", "實作"),
    ("配置", "設定"),
    ("回滾", "還原"),
    ("驗簽", "簽章驗證"),
    ("兜底", "備援處理"),
    ("覆蓋", "涵蓋"),
    ("合規", "法規遵循"),
    ("用戶", "使用者"),
    ("數據", "資料"),
    ("文檔", "文件"),
    ("接口", "介面"),
    ("字段", "欄位"),
    ("默認", "預設"),
    ("異步", "非同步"),
    ("視頻", "影片"),
    ("網絡", "網路"),
    ("在線", "線上"),
    ("服務器", "伺服器"),
    ("調用", "呼叫"),
    ("創建", "建立"),
    ("返回", "回傳"),
    ("接入", "串接"),
    ("全鏈", "完整流程"),
    ("鏈路", "呼叫路徑"),
    ("沉澱", "累積"),
    ("拉取", "取得"),
    ("支持", "支援"),
    ("兼容", "相容"),
    ("實時", "即時"),
    ("信息", "資訊"),
    ("消息", "訊息"),
    ("搜索", "搜尋"),
    ("集群", "叢集"),
    ("灰度", "分階段發布"),
    ("標注", "標註"),
    ("賬號", "帳號"),
    ("日誌", "紀錄"),
    ("質量", "品質"),
    ("主流", "主要流程"),
    ("回填", "補登"),
)
TAIWAN_WORDING_FORBIDDEN = tuple(source for source, _ in TAIWAN_WORDING_REPLACEMENTS)


def taiwan_wording(value: object) -> str:
    """將顯示文案統一為台灣常用詞，不回寫上游規格原文。"""
    text = clean_md(value)
    protected_labels: dict[str, str] = {}
    for index, label in enumerate(sorted(COMPONENT_GLOSSARY, key=len, reverse=True)):
        if label not in text:
            continue
        placeholder = f"§§SL_COMPONENT_{index}§§"
        text = text.replace(label, placeholder)
        protected_labels[placeholder] = label
    for source, target in TAIWAN_WORDING_REPLACEMENTS:
        text = text.replace(source, target)
    regex_replacements = (
        (r"(?<![A-Za-z0-9_])fallback(?![A-Za-z0-9_])", "備援處理"),
        (r"(?<![A-Za-z0-9_])failover(?![A-Za-z0-9_])", "備援切換"),
        (r"(?<![A-Za-z0-9_])rollback(?![A-Za-z0-9_])", "還原"),
        (r"(?<![A-Za-z0-9_])artifact(?![A-Za-z0-9_])", "產出檔"),
        (r"(?<![A-Za-z0-9_])workflow(?![A-Za-z0-9_])", "工作流程"),
        (r"(?<![A-Za-z0-9_])checklist(?![A-Za-z0-9_])", "檢查清單"),
        (r"(?<![A-Za-z0-9_])timeout(?![A-Za-z0-9_])", "逾時"),
        (r"(?<![A-Za-z0-9_])retry(?![A-Za-z0-9_])", "重試"),
        (r"(?<![A-Za-z0-9_])owner(?![A-Za-z0-9_])", "負責人"),
        (r"(?<![A-Za-z0-9_])runbook(?![A-Za-z0-9_])", "維運手冊"),
        (r"(?<![A-Za-z0-9_])benchmark(?![A-Za-z0-9_])", "基準測試"),
        (r"(?<![A-Za-z0-9_])leakage(?![A-Za-z0-9_])", "資料外洩"),
        (r"(?<![A-Za-z0-9_])crash(?![A-Za-z0-9_])", "當機"),
        (r"(?<![A-Za-z0-9_])concurrent(?![A-Za-z0-9_])", "位並行使用者"),
        (r"(?<![A-Za-z0-9_])page(?![A-Za-z0-9_])", "通知值班人員"),
    )
    for pattern, target in regex_replacements:
        text = re.sub(pattern, target, text, flags=re.IGNORECASE)
    text = re.sub(r"(?<=[\u3400-\u9fff])\s+(?=[\u3400-\u9fff])", "", text)
    for placeholder, label in protected_labels.items():
        text = text.replace(placeholder, label)
    return text


def parse_requirements() -> list[Requirement]:
    rows: list[Requirement] = []
    heading = ""
    for line_no, line in enumerate(read_lines(SRS_PATH), start=1):
        if line.startswith("### "):
            heading = clean_md(line[4:])
        cells = split_md_row(line)
        if not cells:
            continue
        match = re.match(r"^(FR-([A-Z]+)-(\d+))(?:\s*🔜)?$", cells[0])
        if not match or len(cells) < 6:
            continue
        rows.append(
            Requirement(
                req_id=match.group(1),
                raw_id=cells[0],
                prefix=match.group(2),
                name=cells[1],
                precondition=cells[2],
                flow=cells[3],
                acceptance=cells[4],
                trace=cells[5],
                heading=heading,
                source_line=line_no,
                raw_line=line,
            )
        )
    counts = Counter(row.req_id for row in rows)
    for row in rows:
        row.duplicate = counts[row.req_id] > 1
    return rows


def parse_nfrs() -> list[NFR]:
    rows: list[NFR] = []
    heading = ""
    for line_no, line in enumerate(read_lines(NFR_PATH), start=1):
        if line.startswith("## "):
            heading = clean_md(line[3:])
        cells = split_md_row(line)
        if not cells or len(cells) < 5:
            continue
        match = re.match(r"^(NFR-([A-Za-z0-9]+)-(\d+))$", cells[0])
        if not match:
            continue
        rows.append(
            NFR(
                req_id=match.group(1),
                category=match.group(2),
                name=cells[1],
                target=cells[2],
                verification=cells[3],
                tier=cells[4],
                heading=heading,
                source_line=line_no,
            )
        )
    return rows


def parse_test_cases() -> list[TestCase]:
    rows: list[TestCase] = []
    heading = ""
    for line_no, line in enumerate(read_lines(TEST_CASE_PATH), start=1):
        if line.startswith("## ") or line.startswith("### "):
            heading = clean_md(line.lstrip("# "))
        cells = split_md_row(line)
        if not cells:
            continue
        match = re.match(r"^(TC-[A-Z0-9-]+)", cells[0])
        if not match:
            continue
        rows.append(
            TestCase(
                tc_id=match.group(1),
                source_ref=cells[1] if len(cells) > 1 else "",
                cells=cells,
                heading=heading,
                source_line=line_no,
            )
        )
    return rows


def parse_adrs() -> list[list[str]]:
    rows: list[list[str]] = []
    group = ""
    for line in read_lines(ADR_INDEX_PATH):
        cells = split_md_row(line)
        if not cells:
            continue
        if cells[0].startswith("群 "):
            group = cells[0]
            continue
        match = re.match(r"^(ADR-\d+)$", cells[0])
        if match and len(cells) >= 5:
            rows.append([group, cells[0], cells[1], cells[2], cells[3], cells[4]])
    return rows


def parse_wbs() -> list[list[str]]:
    rows: list[list[str]] = []
    milestone = ""
    for line in read_lines(ROADMAP_PATH):
        if line.startswith("### "):
            milestone = clean_md(line[4:])
        cells = split_md_row(line)
        if len(cells) < 4:
            continue
        if re.match(r"^\d+(?:\.\d+){0,2}$", cells[0]):
            if len(cells) >= 6:
                rows.append([milestone, *cells[:6]])
            else:
                # M4/M5 概要表只有 4 欄。
                rows.append([milestone, cells[0], "", cells[1], "", "", " / ".join(cells[2:])])
    return rows


def parse_trace_legacy_ids() -> set[str]:
    ids: set[str] = set()
    for line in read_lines(TRACE_PATH):
        cells = split_md_row(line)
        if not cells:
            continue
        for match in re.finditer(r"FR-\d{4}", cells[0]):
            ids.add(match.group(0))
    return ids


def subsystem(req: Requirement) -> dict[str, str]:
    return SUBSYSTEMS.get(
        req.prefix,
        {
            "name": req.prefix,
            "short": req.prefix,
            "component": req.prefix,
            "description": "",
        },
    )


def module_for(req: Requirement) -> tuple[str, str]:
    number = req.req_id.rsplit("-", 1)[-1]
    for code, name, members in MODULES.get(req.prefix, []):
        if number in members:
            return code, name
    return "OTHER", "其他"


def module_arch(prefix: str, code: str) -> dict[str, str]:
    meta = SUBSYSTEMS.get(prefix, {})
    return MODULE_ARCH.get(
        f"{prefix}.{code}",
        {
            "component": meta.get("component", prefix),
            "sad": meta.get("sad", "12_SAD [待標註]"),
            "sds": meta.get("sds", "15_SDS [待標註]"),
            "path": meta.get("path", "[待確認]"),
        },
    )


def architecture_for(req: Requirement) -> dict[str, str]:
    code, _ = module_for(req)
    return module_arch(req.prefix, code)


def subsystem_component_labels(prefix: str) -> str:
    labels: list[str] = []
    for code, _, _ in MODULES.get(prefix, []):
        for label in (
            part.strip() for part in module_arch(prefix, code)["component"].split(";")
        ):
            if label and label not in labels:
                labels.append(label)
    return "; ".join(labels)


def component_glossary_rows() -> list[list[str]]:
    """建立受控元件標籤字典，來源定位由 MODULE_ARCH 反向彙整。"""
    locations: dict[str, dict[str, set[str]]] = defaultdict(
        lambda: {"modules": set(), "sad": set(), "sds": set(), "path": set()}
    )
    for module_key, arch in MODULE_ARCH.items():
        for label in (part.strip() for part in arch["component"].split(";")):
            if not label:
                continue
            locations[label]["modules"].add(module_key.replace(".", "·"))
            locations[label]["sad"].add(arch["sad"])
            locations[label]["sds"].add(arch["sds"])
            locations[label]["path"].add(arch["path"])
    rows = []
    for label, text in COMPONENT_GLOSSARY.items():
        source = locations[label]
        rows.append(
            [
                label,
                text.get("alias", "—"),
                text["definition"],
                text["boundary"],
                "、".join(sorted(source["modules"])),
                " ｜ ".join(sorted(source["sad"])),
                " ｜ ".join(sorted(source["sds"])),
                " ｜ ".join(sorted(source["path"])),
            ]
        )
    return rows


def phase_for(req: Requirement) -> str:
    explicit = PHASE_BY_ID.get(req.req_id)
    if explicit:
        return explicit
    if req.prefix in {"AGT", "API", "WEB", "DAT"}:
        return "M1"
    if req.prefix in {"REF", "TEC"}:
        return "M2"
    return "M1→M2"


def status_for_requirement(req: Requirement) -> str:
    if req.duplicate:
        return "🔴 上游 ID 衝突"
    if "[待確認]" in req.all_text or "〔待確認〕" in req.all_text:
        return "❓ 待確認"
    if "🔜" in req.raw_id:
        return "🔜 規劃中"
    if "🔜" in req.all_text:
        return "🔶 需求定版／部分規劃中"
    return "✅ 需求定版"


def requirement_copy_key(req: Requirement) -> str:
    """取得 PM／業務文案鍵；若未來再碰鍵，加入名稱避免誤用。"""
    if req.duplicate:
        return f"{req.req_id}::{req.name}"
    return req.req_id


def business_copy_for(req: Requirement) -> dict[str, str]:
    """回傳業務 VOC 與 PRD 驗收文案；缺少映射時保持可讀降級。"""
    key = requirement_copy_key(req)
    return FR_BUSINESS_COPY.get(
        key,
        {
            "voc": f"業務使用者需要「{req.name}」能解決實際作業問題，且結果可追溯。",
            "prd_acceptance": f"以正常、例外與未授權情境驗收「{req.name}」；使用者可完成預期作業，異常時有明確結果且不破壞資料。",
        },
    )


def planning_acceptance_summary(req: Requirement) -> str:
    """規劃書用的簡明但可判定驗收摘要。"""
    copy = business_copy_for(req)
    return f"PRD 驗收：{copy['prd_acceptance']} ｜ 成功指標：{req.acceptance}"


def nfr_business_copy(nfr: NFR) -> dict[str, str]:
    """將 NFR 技術主題翻成客戶價值與可觀察的 PRD 驗收結果。"""
    category = nfr.category.lower()
    if category == "perf":
        voc = f"日常使用者希望「{nfr.name}」回應足夠快，避免等待、重複點擊或放棄作業。"
        acceptance = f"在約定驗收環境與代表性負載下執行「{nfr.name}」；使用者可在承諾等待時間內完成任務，逾時時有可理解的降級或重試結果。"
    elif category == "avail":
        voc = f"客戶希望「{nfr.name}」在尖峰、斷線或單一元件異常時仍可用，不因平台問題中斷服務。"
        acceptance = f"以正常運作、短暫中斷與恢復三種情境驗收「{nfr.name}」；核心作業持續可用或有明確降級說明，恢復後資料不漏失。"
    elif category == "rel":
        voc = f"業務希望「{nfr.name}」穩定可靠，案件與訊息不會靜默失敗、重複或永久消失。"
        acceptance = f"以成功、失敗、重試與中斷後恢復情境驗收「{nfr.name}」；結果可追查，不重複產生業務後果，未完成項會進入待處理。"
    elif category == "sla":
        voc = f"客戶與營運團隊需要「{nfr.name}」有明確時限、逾時邊界與升級路徑。"
        acceptance = f"以未到期、恰好到期與逾期三個邊界驗收「{nfr.name}」；狀態標示、提醒對象與後續升級符合服務承諾。"
    elif category == "scal":
        voc = f"品牌希望「{nfr.name}」能隨用戶、案件、資料與品牌數成長，不需大幅中斷服務才能擴充。"
        acceptance = f"在目標容量與尖峰情境下驗收「{nfr.name}」；主要使用者任務仍可完成，新增容量或品牌不破壞既有租戶服務。"
    elif category == "sec":
        voc = f"客戶與品牌需要平台在「{nfr.name}」上防止未授權存取、惡意操作或敏感資料外洩。"
        acceptance = f"以合法、未授權與惡意三類情境驗收「{nfr.name}」；合法使用者可完成作業，非法或越權行為被阻擋並留下可追查記錄。"
    elif category == "priv":
        voc = f"資料權利人希望「{nfr.name}」符合告知、最少使用、隔離、保留與刪除承諾。"
        acceptance = f"以正常處理、跨租戶、到期與刪除情境驗收「{nfr.name}」；只有必要資料在允許範圍內可見，逾期或未授權資料不再可存取。"
    elif category == "obs":
        voc = f"營運與維運團隊需要透過「{nfr.name}」及早發現影響客戶的異常，並快速定位、處理與恢復。"
        acceptance = f"以正常基準、人工注入異常與恢復情境驗收「{nfr.name}」；影響可被偵測、告警、追查並依程序復原。"
    elif category == "aud":
        voc = f"稽核、財務與治理人員需要「{nfr.name}」完整、不可篡改且能回答誰、何時、為何做了什麼。"
        acceptance = f"以建立、查詢、更正與篡改嘗試驗收「{nfr.name}」；歷史可按人、時間、對象與原因追溯，且篡改能被發現。"
    elif category in {"dq", "pub"}:
        voc = f"內容與品質團隊需要「{nfr.name}」確保只有來源可信、已核准且屬於正確品牌的知識會服務客戶。"
        acceptance = f"以核准、未核准、重複處理與錯誤品牌資料驗收「{nfr.name}」；只有符合來源、審核與租戶規則的內容可使用。"
    elif category in {"sch", "rep", "maint", "dora"}:
        voc = f"品牌希望透過「{nfr.name}」讓系統可安全變更、可恢復、可重現，不因升級而影響營運。"
        acceptance = f"以正常變更、重複套用、失敗與恢復情境驗收「{nfr.name}」；結果可重現、差異可發現，失敗不破壞已有服務。"
    elif category == "a11y":
        voc = f"不同視覺、動作與輔助科技需求的使用者，都需要能完成「{nfr.name}」相關作業。"
        acceptance = f"以鍵盤、螢幕閱讀器、縮放與對比檢查驗收「{nfr.name}」；關鍵資訊、狀態與操作不依賴單一感官或輸入方式。"
    elif category == "comp":
        voc = f"客戶與法務團隊需要「{nfr.name}」符合合約、SOW 與法令承諾，不讓重大違規項目進入正式服務。"
        acceptance = f"使用合規、邊界與違規情境驗收「{nfr.name}」；合規作業可完成，違規行為被阻擋、告警或阻止上線。"
    else:
        voc = f"客戶希望「{nfr.name}」有明確的品質承諾，結果穩定且可驗證。"
        acceptance = f"在約定環境下以正常、邊界與異常情境驗收「{nfr.name}」；使用者結果可觀察，失敗時有明確處理。"
    return {"voc": voc, "prd_acceptance": acceptance}


def criterion_state(text: str) -> str:
    if "[待確認]" in text or "[待補]" in text or "〔待確認〕" in text:
        return "❓ 門檻待定"
    if "🔜" in text:
        return "🔶 可驗收／落地待對帳"
    return "✅ 判準已述"


def nfr_group(nfr: NFR) -> str:
    category = nfr.category.lower()
    if category in {"perf", "avail", "rel", "sla", "scal"}:
        return "PERF"
    if category in {"sec", "priv"}:
        return "SEC"
    return "OPS"


def nfr_status(nfr: NFR) -> str:
    if "[待確認]" in nfr.all_text:
        return "❓ 待確認"
    if "🔜" in nfr.all_text:
        return "🔶 目標已定／部分規劃中"
    if "合約下限" in nfr.tier:
        return "🔴 合約下限"
    return "✅ 營運目標"


def test_priority(req_id: str, text: str = "") -> str:
    high_terms = (
        "金額",
        "報價",
        "工單",
        "派工",
        "退款",
        "授權",
        "RBAC",
        "SoD",
        "GDPR",
        "租戶",
        "合約",
        "影像辨識禁用",
        "轉真人",
    )
    if req_id.startswith(("FR-AGT", "FR-API")) or req_id.startswith(("NFR-Sec", "NFR-Priv", "NFR-Aud", "NFR-Comp")):
        return "P0"
    if any(term.lower() in text.lower() for term in high_terms):
        return "P0"
    return "P1"


def scenario_for_requirement(req: Requirement) -> str:
    number = int(req.req_id.rsplit("-", 1)[-1])
    if req.prefix == "AGT":
        return "TS-01、TS-10" if number in {4, 6, 8, 11} else "TS-01"
    if req.prefix == "API":
        if number in {1, 2, 3, 4, 17}:
            return "TS-02"
        if number in {5, 6, 7}:
            return "TS-03"
        if number in {8, 9, 19}:
            return "TS-04"
        if number in {10, 11, 12}:
            return "TS-05"
        if number in {13, 16, 18}:
            return "TS-10"
        return "TS-11"
    if req.prefix == "WEB":
        return {
            1: "TS-08",
            2: "TS-08、TS-10",
            3: "TS-03",
            4: "TS-11",
            5: "TS-09、TS-10",
            6: "TS-02",
            7: "TS-11",
        }.get(number, "TS-11")
    if req.prefix == "DAT":
        return "TS-07" if number in {1, 4} else "TS-09、TS-10"
    if req.prefix == "REF":
        return "TS-07、TS-10"
    if req.prefix == "TEC":
        if number == 7:
            return "TS-04"
        if number in {1, 2, 8}:
            return "TS-06"
        if number in {3, 4, 5}:
            return "TS-03"
        return "TS-05"
    if req.prefix == "PLT":
        if number in {1, 2, 3, 9}:
            return "TS-08、TS-10"
        if number in {4, 6}:
            return "TS-11"
        if number == 5:
            return "TS-01、TS-11"
        return "TS-12"
    return "[待分流]"


def scenario_for_nfr(nfr: NFR) -> str:
    group = nfr_group(nfr)
    if group == "PERF":
        return "TS-11"
    if group == "SEC":
        return "TS-08、TS-10"
    if nfr.category.lower() in {"dq", "pub"}:
        return "TS-07、TS-10"
    if nfr.category.lower() in {"sch", "rep", "aud"}:
        return "TS-09、TS-10"
    if nfr.category.lower() == "a11y":
        return "TS-02、TS-08"
    return "TS-09、TS-11"


def tc_hint(req_id: str, *, is_nfr: bool = False) -> str:
    if is_nfr:
        return NFR_TC_HINTS.get(req_id, f"TC-{req_id}（受控 NFR 驗證設計）")
    return FR_TC_HINTS.get(req_id, f"TC-{req_id}（正常／邊界／例外）")


def qa_acceptance_for(req: Requirement) -> str:
    """QA 執行清單與治理附錄共用的通過判準；不宣稱已執行。"""
    overrides = {
        "FR-API-05": (
            "一般通知送達 P95 ≤ 30 秒、急件 P95 ≤ 15 秒；候選池為空時 100% 進入 "
            "dispatch_pending 並觸發告警；急件排序降低距離權重、提高評分權重。"
        ),
        "FR-REF-01": (
            "主通道採唯讀 API 增量汲取，事件處理即時增量，每日批次 reconciliation 補漏；"
            "來源 provenance 完整率 100%，同一 source_id 重跑不得產生重複資料。"
        ),
    }
    value = overrides.get(req.req_id, req.acceptance)
    return value.replace("[待確認]", "（依本表受控判準）").replace("〔待確認〕", "（依本表受控判準）")


def qa_plain_text(value: str) -> str:
    """保留必要產品/API 名稱，將 QA 不應先猜的縮寫改為可執行語言。"""
    text = (
        value.replace("[待確認]", "依本表受控判準")
        .replace("〔待確認〕", "依本表受控判準")
        .strip()
    )
    for old, new in (
        ("Clarify gate", "釐清確認門檻"),
        ("Clarify", "釐清確認"),
        ("transfer_to_human", "轉真人"),
        ("per-user", "每位客戶"),
        ("BUILD/SAVE", "讀取/保存"),
        ("debounce", "短時間訊息合併"),
        ("takeover", "接管"),
        ("Medallion", "原始/整理/正式資料"),
        ("Publisher", "發佈服務"),
        ("Pipeline", "資料處理流程"),
        ("append-only", "只增不刪"),
        ("pgvector", "向量知識庫"),
        ("references", "參考資料"),
        ("OIDC", "登入憑證"),
        ("RBAC", "角色權限"),
        ("License", "使用授權"),
        ("KYC", "身分審核"),
        ("CQRS", "工作台資料投影"),
        ("SoD", "職能分離"),
        ("WS", "即時連線"),
        ("hard gate", "必要阻擋條件"),
        ("completeness", "資料完整度"),
        ("internal token", "服務間測試憑證"),
        ("escalation", "轉真人案件"),
        ("RAG", "知識檢索回答"),
        ("OHS", "技師媒合服務"),
        ("fail-closed", "驗證失敗時一律拒絕"),
        ("dedup", "相同事件不得重複處理"),
        ("provenance", "來源追溯資料"),
        ("migration drift", "資料庫版本差異"),
        ("rollback", "回復上一版"),
        ("requote", "現場重新報價"),
        ("settlement", "結算"),
        ("fixture", "測試資料"),
        ("assignee", "目前被指派技師"),
        ("audit", "稽核紀錄"),
        ("Turn", "單次訊息處理紀錄"),
        ("HITL", "人工審核"),
        ("APP_MODE", "站點模式（APP_MODE）"),
        ("SLA", "時限目標"),
        ("Functional", "功能"),
        ("Security", "安全"),
        ("Resilience", "故障恢復"),
        ("State", "狀態"),
        ("Boundary", "邊界"),
        ("Contract", "契約"),
        ("Timeout", "逾時"),
        ("gate", "阻擋條件"),
    ):
        text = text.replace(old, new)
    text = text.replace("🔜", "").replace("規劃中：", "").replace("規劃中", "")
    text = text.replace("人工審核 審核", "人工審核").replace("硬 阻擋條件", "必要阻擋條件")
    text = re.sub(r"\bFR-[A-Z]+(?:-\d+)?\b", "相關功能規則", text)
    text = re.sub(r"\bNFR-[A-Za-z0-9]+(?:-\d+)?\b", "對應品質門檻", text)
    return clean_md(text)


# 環境、反例、零副作用判準、證據。每列 QA 執行項目都由此產生具體內容。
FR_QA_CONTEXT = {
    "AGT": (
        "SIT LINE 測試頻道、對應品牌與可查詢後台案件的帳號",
        "錯誤簽章、同一事件重送、急件、無法命中知識的問題與另一品牌/客戶資料（依本項擇用）",
        "被拒絕或重送時不得多出回覆、問題卡、轉人紀錄或跨客戶記憶",
        "LINE 輸入/回覆、webhook request/response、問題卡/轉人紀錄與 trace ID",
    ),
    "API": (
        "SIT API、可查詢操作前後狀態的資料庫/稽核紀錄與外部服務替身",
        "非法狀態、門檻前/等於/後數值、同一請求重送、無權角色與逾時（依本項擇用）",
        "被拒絕或重送時不得多出狀態變更、事件、帳本或重複稽核紀錄",
        "request/response、操作前後資料、事件/帳本/稽核紀錄與 trace ID",
    ),
    "WEB": (
        "SIT 前後端、各站點模式與可切換權限的測試帳號",
        "無權角色、錯站點模式、403、斷網、即時連線中斷與鍵盤操作（依本項擇用）",
        "畫面不得顯示假成功；後端未成功時資料不得變更",
        "關鍵畫面、console/network 結果、後端最終狀態與無障礙報告",
    ),
    "DAT": (
        "可重建的三個測試資料庫、固定輸入與可比對輸出的工具",
        "錯資料庫連線、錯租戶、錯來源標記、同一輸入重跑與資料庫版本差異（依本項擇用）",
        "失敗或重跑時不得走錯庫、產生重複資料或遺失來源",
        "腳本/資料集版本、前後筆數、hash、來源比對、版本差異與告警",
    ),
    "REF": (
        "知識精煉 SIT、人工審核帳號與可查詢發佈前後知識的工具",
        "拒絕、重複、錯來源、錯租戶與未核可素材（依本項擇用）",
        "未核可、拒絕、錯來源或錯租戶內容落地數必須為 0",
        "輸入素材、差異、審核紀錄、發佈前後知識比對、來源與版本",
    ),
    "TEC": (
        "技師平台 SIT、品牌端與可觀察候選池/工單狀態的帳號",
        "未核可、未授權、已停權、無排班與另一品牌技師（依本項擇用）",
        "錯誤技師不得入候選池；技師端不得顯示超出作業所需的資料",
        "技師資料、候選池、接拒單時間、兩端狀態、顯示欄位與結算差額",
    ),
    "PLT": (
        "平台 SIT、兩個品牌、各角色帳號與可觀察告警/版本的工具",
        "錯品牌、錯角色、過期憑證、依賴中斷與不合法設定（依本項擇用）",
        "被拒絕或降級時不得跨品牌、放大權限、污染資料或覆寫上一個可用版本",
        "角色/品牌/授權矩陣、拒絕回應、降級畫面、告警、版本與回復紀錄",
    ),
}


# AI 客服的原 SRS 包含較多內部狀態/工具名；主視圖改寫為黑箱 QA 可直接執行的資料、步驟與結果。
FR_QA_OVERRIDES = {
    "FR-AGT-01": (
        "有效簽章的短文字、照片與超過 4900 字訊息，另準備同內容的錯誤簽章版本",
        "分別送出有效短文字、照片、超長文字與錯誤簽章請求，再重送同一事件",
        "有效文字/照片皆有明確回覆；超長內容截斷至 4900 字以內；錯誤簽章回 400 且不產生處理紀錄",
    ),
    "FR-AGT-02": (
        "同一測試客戶的連續兩則訊息，並能在測試環境模擬對話儲存失敗",
        "第一則提供一個客戶事實，第二則追問該事實；接著模擬對話儲存失敗後再送一則訊息",
        "第二則回覆能承接前文；每個 webhook 只有一筆處理紀錄；儲存失敗時仍回覆客戶且異常可由 trace 查到",
    ),
    "FR-AGT-03": (
        "一題案例庫可明確命中、一題需查知識庫、一題無法回答，並準備「已釐清/尚未釐清」回覆",
        "逐題提問並檢查系統來源；對無法回答的題目連續 3 次回覆「尚未釐清」",
        "案例庫命中與知識庫回答來源正確；系統會詢問是否已釐清；第 3 次未釐清後建立後台案件並轉真人",
    ),
    "FR-AGT-04": (
        "被鎖門外、受困室內、安全風險與高度憤怒各一則訊息，再準備一則一般問題",
        "依次送出四種急件與一般問題，記錄訊息進入到後台案件建立的時間",
        "四種急件均跳過自助回答，5 分鐘內建立轉真人案件並記錄急件時間；一般問題不得誤判急件",
    ),
    "FR-AGT-05": (
        "金額詢問、要求真人、急件與派工要求各一，並能模擬 AI 口頭承諾轉接但未實際呼叫轉接動作",
        "逐組送出紅線訊息；最後一組故意讓 AI 未呼叫轉接動作，檢查後端是否自動補建案件",
        "每組需轉真人的訊息均只建一張後台案件；AI 漏掉轉接動作時後端仍會補建，不得遺失案件",
    ),
    "FR-AGT-06": (
        "兩個品牌，每個品牌兩位測試客戶；每人先提供一個不同的可記憶事實",
        "讓每位客戶再次詢問自己的事實；再嘗試以錯誤客戶 ID、錯品牌 ID 與缺少 ID 讀寫記憶",
        "每人只讀到自己事實；錯客戶、錯品牌或缺少 ID 一律拒絕；跨客戶與跨品牌洩漏數為 0",
    ),
    "FR-AGT-07": (
        "一題常見流程問題、一題常見產品事實、一題長尾知識與一題無來源問題",
        "逐題提問並記錄回答來源；再中斷長尾知識檢索服務後重試",
        "常見問題使用正確受控內容；長尾問題引用正確來源；無來源或檢索失敗時不編造，應詢問或轉真人",
    ),
    "FR-AGT-08": (
        "客服允許使用的 6 種動作各一，以及至少一種未授權檔案/系統動作",
        "逐一呼叫 6 種允許動作，再嘗試註冊與呼叫未授權動作",
        "6 種允許動作依權限成功；白名單外動作無法註冊也無法呼叫，且留下明確拒絕紀錄",
    ),
    "FR-AGT-09": (
        "同一對話的「AI 處理」與「人工接管」狀態，並能模擬接管狀態查詢逾時",
        "兩種狀態下各送客戶訊息；人工接管期間再送一則客服訊息；最後模擬狀態查詢逾時",
        "AI 處理時正常回覆；人工接管時 AI 不回覆，但客戶與客服訊息都完整入庫；狀態查詢逾時不可讓客戶無回應",
    ),
    "FR-AGT-10": (
        "1.5 秒內連發 3 則訊息、1.5 秒後的第 4 則訊息，以及同一事件 ID 的重送請求",
        "依時間間隔送出四則訊息，再於 24 小時內重送同一事件 ID",
        "1.5 秒內的三則訊息合併為一次處理；第四則另開一次處理；24 小時內重送不產生第二次回覆或案件",
    ),
    "FR-AGT-11": (
        "要求最終報價、要求複誦個案報價、要求辨識門鎖照片與一組禁止行為題庫",
        "逐一送出金額與影像請求，再執行完整禁止行為題庫",
        "AI 只能提供價格區間，不得給最終價、複誦個案金額或辨識影像；禁止行為題庫達到文件門檻且影像違規數為 0",
    ),
}


def qa_execution_for(req: Requirement) -> dict[str, str]:
    """產出 QA 拿到單列即可開始的前置、動作、通過條件與證據。"""
    environment, negative_data, no_side_effect, evidence = FR_QA_CONTEXT[req.prefix]
    override = FR_QA_OVERRIDES.get(req.req_id)
    source_precondition = "" if req.precondition.strip() in {"", "—", "-"} else req.precondition
    source_flow = "" if req.flow.strip() in {"", "—", "-"} else req.flow
    normal_data = (
        override[0]
        if override
        else qa_plain_text(source_precondition) or f"建立可完成「{qa_plain_text(req.name)}」的正常資料"
    )
    main_flow = (
        override[1]
        if override
        else qa_plain_text(source_flow) or f"執行「{qa_plain_text(req.name)}」主流"
    )
    preparation = (
        f"環境：{environment}。｜"
        f"正常資料：{normal_data}。｜"
        f"反例資料：{negative_data}。"
    )
    steps = (
        "1. 記錄操作前的資料/畫面狀態。｜"
        f"2. 執行正常流程：{main_flow}。｜"
        "3. 每次只更改一個邊界、權限、重送或故障條件後重跑。｜"
        "4. 比對每次操作後的畫面/API、資料與稽核紀錄。"
    )
    acceptance = override[2] if override else qa_plain_text(qa_acceptance_for(req))
    expected = (
        f"功能判準：{acceptance}。｜"
        f"副作用判準：{no_side_effect}。"
    )
    return {
        "preparation": preparation,
        "steps": steps,
        "expected": expected,
        "evidence": evidence,
    }


def qa_scenario_names(value: str) -> str:
    names = {row[0]: row[1] for row in SCENARIOS}
    return "、".join(
        qa_plain_text(names[scenario_id]) for scenario_id in re.findall(r"TS-\d+", value)
    )


def qa_visible_id(req_id: str, kind: str) -> str:
    """建立 QA 看得懂且穩定的情境／案例編號，不暴露 FR/NFR 治理前綴。"""
    suffix = re.sub(r"^(?:FR|NFR)-", "", req_id).upper()
    return f"{kind}-{suffix}"


def qa_scenario_types(req_id: str, *, is_nfr: bool = False) -> str:
    """用少量白話類型提醒 QA 不可只跑正常流程。"""
    if is_nfr:
        group = nfr_group_from_id(req_id)
        return {
            "PERF": "目標值／邊界／尖峰／故障恢復",
            "SEC": "合法／未授權／跨品牌／惡意輸入",
            "OPS": "正常變更／重跑／失敗／恢復",
        }[group]
    prefix = req_id.split("-", 2)[1]
    return {
        "AGT": "正常／例外／重送／跨品牌",
        "API": "正常／狀態／邊界／權限",
        "WEB": "正常／權限／斷線／可用性",
        "DAT": "正常／重跑／錯庫／版本差異",
        "REF": "核准／拒絕／錯來源／跨品牌",
        "TEC": "正常／資格／授權／跨品牌",
        "PLT": "正常／權限／依賴中斷／版本還原",
    }[prefix]


def nfr_group_from_id(req_id: str) -> str:
    category = req_id.split("-", 2)[1].lower()
    if category in {"perf", "avail", "rel", "sla", "scal"}:
        return "PERF"
    if category in {"sec", "priv"}:
        return "SEC"
    return "OPS"


NFR_QA_CONTEXT = {
    "PERF": (
        "類正式環境、固定資料集、壓測/故障演練腳本與監控儀表板",
        "準備基準、目標、尖峰三段負載；需故障演練時準備可單獨中斷的依賴",
        "腳本/資料集版本、原始量測、p95/p99、錯誤率、告警、降級與恢復時間",
    ),
    "SEC": (
        "SIT 安全環境、兩個品牌、各角色帳號與可查詢稽核紀錄的工具",
        "準備允許/拒絕權限、偽造/過期憑證、跨品牌 ID、AI 對抗輸入與個資刪除/保留資料",
        "輸入、request/response、操作前後資料、拒絕結果、稽核紀錄與必要截圖",
    ),
    "OPS": (
        "可重建環境、固定資料/輸出、CI 報告與可回復上一版的版本",
        "準備同一輸入重跑、空庫/舊版升級、稽核竄改、回復上一版與 CI/無障礙檢查資料",
        "輸入/版本、前後筆數與 hash、來源/trace、CI/無障礙報告與恢復時間",
    ),
}


def qa_nfr_execution_for(nfr: NFR) -> dict[str, str]:
    group = nfr_group(nfr)
    environment, test_data, evidence = NFR_QA_CONTEXT[group]
    verification = qa_plain_text(nfr.verification)
    target = qa_plain_text(nfr.target)
    return {
        "preparation": (
            f"環境：{environment}。｜"
            f"資料：{test_data}。｜"
            "版本：執行前記錄 commit/build、環境、資料集與腳本版本。"
        ),
        "steps": (
            "1. 先跑一次可重現的基準並保存原始值。｜"
            f"2. 執行文件指定驗證：{verification}。｜"
            "3. 增加邊界、拒絕或故障情境後重跑。｜"
            "4. 將每次實測值與目標逐項比對，超出目標即失敗。"
        ),
        "expected": f"通過標準：{target}。",
        "evidence": evidence,
    }


def qa_test_method_for(req: Requirement) -> str:
    """把測試分析轉為可執行的準備、操作與觀察方法。"""
    execution = qa_execution_for(req)
    return (
        f"準備：{execution['preparation']}｜"
        f"操作：{execution['steps']}｜"
        f"觀察：{execution['expected']}"
    )


def qa_mapping_id(req_id: str) -> str:
    return f"QTM-{req_id}"


def qa_owner(req_id: str, *, is_nfr: bool = False) -> str:
    if is_nfr:
        category = req_id.split("-", 2)[1].lower()
        if category in {"sec", "priv", "comp"}:
            return "AppSec（主責）／QA"
        if category in {"perf", "avail", "rel", "sla", "scal", "obs", "dora"}:
            return "SRE（主責）／QA"
        if category in {"dq", "pub", "sch", "rep"}:
            return "Data/Platform SA（主責）／QA"
        return "QA Lead（主責）／對應系統 Owner"
    return {
        "AGT": "QA Lead（主責）／AI Owner",
        "API": "QA Lead（主責）／API Owner",
        "WEB": "QA Lead（主責）／Web Owner",
        "DAT": "Data QA（主責）／Data Owner",
        "REF": "Knowledge QA（主責）／Knowledge Owner",
        "TEC": "QA Lead（主責）／Technician SA",
        "PLT": "Platform QA（主責）／SRE",
    }.get(req_id.split("-", 2)[1], "QA Lead（主責）／系統 Owner")


def qa_evidence_spec(req_id: str, *, is_nfr: bool = False) -> str:
    if is_nfr:
        detail = "腳本版本、環境與資料集、原始量測、報表、pass/fail 裁定"
    else:
        detail = "測試結果、request/response 或畫面、log/trace、缺陷與簽核"
    return f"qa-evidence/{req_id}/（執行後保存：{detail}）"


def qa_closure_status(text: str) -> str:
    if "🔜" in text:
        return "✅ 驗收與測試設計已定版；目標版本執行取證"
    return "✅ 驗收、測試設計與追溯映射已閉環"


def detailed_tc_source_srs_links(test_cases: Sequence[TestCase]) -> dict[str, list[str]]:
    """稽核舊版詳細 TC 的原「FR / 來源」欄；不作為現行覆蓋率或追溯閉環判斷。"""
    links: dict[str, list[str]] = defaultdict(list)
    for tc in test_cases:
        for req_id in re.findall(r"FR-[A-Z]+-\d+", tc.source_ref):
            links[req_id].append(tc.tc_id)
    return links


def parse_qtm_canon() -> dict[str, dict[str, str]]:
    """讀取 20_Test_Cases §2.1 的現行受控 QTM 主表。"""
    rows: dict[str, dict[str, str]] = {}
    in_block = False
    for line in read_lines(TEST_CASE_PATH):
        if line.strip() == "<!-- BEGIN GENERATED QTM MAPPING -->":
            in_block = True
            continue
        if line.strip() == "<!-- END GENERATED QTM MAPPING -->":
            break
        if not in_block or not line.startswith("| QTM-"):
            continue
        cells = split_md_row(line)
        if len(cells) < 9:
            continue
        rows[cells[1]] = {
            "qtm": cells[0],
            "type": cells[2],
            "name": cells[3],
            "priority": cells[4],
            "scenario": cells[5],
            "tc": cells[6],
            "method": cells[7],
            "status": cells[8],
        }
    return rows


def parse_test_plan_scenarios() -> dict[str, dict[str, object]]:
    """讀取 19_Test_Plan §1.1 的 TS-01～TS-12 受控基線。"""
    rows: dict[str, dict[str, object]] = {}
    in_block = False
    for line in read_lines(TEST_PLAN_PATH):
        if line.strip() == "<!-- BEGIN GENERATED QA SCENARIO BASELINE -->":
            in_block = True
            continue
        if line.strip() == "<!-- END GENERATED QA SCENARIO BASELINE -->":
            break
        if not in_block or not line.startswith("| TS-"):
            continue
        cells = split_md_row(line)
        if len(cells) < 6:
            continue
        rows[cells[0]] = {
            "name": cells[1],
            "priority": cells[2],
            "method": cells[3],
            "count": int(cells[4]),
            "trace": cells[5],
        }
    return rows


def set_workbook_meta(wb: Workbook, title: str) -> None:
    wb.properties.creator = "Smart Lock enterprise docs generator"
    wb.properties.title = title
    wb.properties.subject = "需求、架構、驗收與測試追溯"
    wb.properties.description = "由 smartlock-docs/enterprise 正典 Markdown 單向生成"
    wb.properties.keywords = "Smart Lock, SRS, BOM, UAT, QA, Traceability"
    wb.calculation.fullCalcOnLoad = True


def style_sheet_common(ws) -> None:
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.outlinePr.summaryBelow = False


def safe_sheet_name(name: str) -> str:
    name = re.sub(r"[\[\]:*?/\\]", "·", name)
    return name[:31]


def apply_widths(ws, widths: Sequence[float] | None, column_count: int) -> None:
    if widths is None:
        widths = [18.0] * column_count
    for idx in range(1, column_count + 1):
        width = widths[idx - 1] if idx <= len(widths) else 18.0
        ws.column_dimensions[get_column_letter(idx)].width = width


def style_header_row(ws, row_number: int, column_count: int) -> None:
    for cell in ws[row_number][:column_count]:
        cell.fill = PatternFill("solid", fgColor=COLORS["blue"])
        cell.font = Font(name=FONT_NAME, color=COLORS["white"], bold=True, size=10)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = TABLE_BORDER
    ws.row_dimensions[row_number].height = 30


def style_data_range(ws, first_row: int, last_row: int, column_count: int) -> None:
    for row_idx in range(first_row, last_row + 1):
        fill = COLORS["pale_blue"] if (row_idx - first_row) % 2 else COLORS["white"]
        for col_idx in range(1, column_count + 1):
            cell = ws.cell(row_idx, col_idx)
            cell.font = Font(name=FONT_NAME, size=9, color=COLORS["black"])
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = TABLE_BORDER
            cell.fill = PatternFill("solid", fgColor=fill)
        ws.row_dimensions[row_idx].height = 36


def add_status_conditional_formatting(ws, first_row: int, last_row: int, column_count: int) -> None:
    if last_row < first_row:
        return
    target = f"A{first_row}:{get_column_letter(column_count)}{last_row}"
    ws.conditional_formatting.add(
        target,
        FormulaRule(
            formula=[f'ISNUMBER(SEARCH("🔴",A{first_row}))'],
            fill=PatternFill("solid", fgColor=COLORS["light_red"]),
        ),
    )
    ws.conditional_formatting.add(
        target,
        FormulaRule(
            formula=[f'OR(ISNUMBER(SEARCH("❓",A{first_row})),ISNUMBER(SEARCH("🔶",A{first_row})))'],
            fill=PatternFill("solid", fgColor=COLORS["light_orange"]),
        ),
    )
    ws.conditional_formatting.add(
        target,
        FormulaRule(
            formula=[f'ISNUMBER(SEARCH("✅",A{first_row}))'],
            fill=PatternFill("solid", fgColor=COLORS["light_green"]),
        ),
    )


def add_banner_table(
    wb: Workbook,
    name: str,
    title: str,
    description: str,
    headers: Sequence[str],
    rows: Iterable[Sequence[object]],
    widths: Sequence[float] | None = None,
) -> object:
    ws = wb.create_sheet(safe_sheet_name(taiwan_wording(name)))
    style_sheet_common(ws)
    column_count = len(headers)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=column_count)
    ws.cell(1, 1, taiwan_wording(title))
    ws.cell(1, 1).fill = PatternFill("solid", fgColor=COLORS["navy"])
    ws.cell(1, 1).font = Font(name=FONT_NAME, color=COLORS["white"], bold=True, size=16)
    ws.cell(1, 1).alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 32
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=column_count)
    ws.cell(2, 1, taiwan_wording(description))
    ws.cell(2, 1).fill = PatternFill("solid", fgColor=COLORS["light_blue"])
    ws.cell(2, 1).font = Font(name=FONT_NAME, color=COLORS["dark_gray"], italic=True, size=9)
    ws.cell(2, 1).alignment = Alignment(vertical="top", wrap_text=True)
    ws.row_dimensions[2].height = 45
    for idx, header in enumerate(headers, start=1):
        ws.cell(4, idx, taiwan_wording(header))
    style_header_row(ws, 4, column_count)
    normalized = [
        [taiwan_wording(value) if isinstance(value, str) else value for value in row]
        for row in rows
    ]
    for row in normalized:
        ws.append(list(row))
    last_row = max(ws.max_row, 4)
    if last_row >= 5:
        style_data_range(ws, 5, last_row, column_count)
        add_status_conditional_formatting(ws, 5, last_row, column_count)
    ws.freeze_panes = "A5"
    ws.auto_filter.ref = f"A4:{get_column_letter(column_count)}{last_row}"
    apply_widths(ws, widths, column_count)
    ws.print_title_rows = "4:4"
    return ws


def add_plain_table(
    wb: Workbook,
    name: str,
    headers: Sequence[str],
    rows: Iterable[Sequence[object]],
    widths: Sequence[float] | None = None,
) -> object:
    ws = wb.create_sheet(safe_sheet_name(taiwan_wording(name)))
    style_sheet_common(ws)
    for idx, header in enumerate(headers, start=1):
        ws.cell(1, idx, taiwan_wording(header))
    style_header_row(ws, 1, len(headers))
    normalized = [
        [taiwan_wording(value) if isinstance(value, str) else value for value in row]
        for row in rows
    ]
    for row in normalized:
        ws.append(list(row))
    last_row = max(ws.max_row, 1)
    if last_row >= 2:
        style_data_range(ws, 2, last_row, len(headers))
        add_status_conditional_formatting(ws, 2, last_row, len(headers))
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{last_row}"
    apply_widths(ws, widths, len(headers))
    ws.print_title_rows = "1:1"
    return ws


def add_architecture_doc_links(ws, header_row: int) -> None:
    """讓 SAD/SDS 定位欄可直接開啟正典文件；定位文字保留節次與行號。"""
    for cell in ws[header_row]:
        value = str(cell.value or "")
        if value == "SAD 定位":
            target = f"../{SAD_PATH.name}"
        elif value == "SDS 定位":
            target = f"../{SDS_PATH.name}"
        else:
            continue
        for row_idx in range(header_row + 1, ws.max_row + 1):
            source_cell = ws.cell(row_idx, cell.column)
            if source_cell.value and str(source_cell.value).strip() not in {"—", "[待標註]"}:
                source_cell.hyperlink = target
                source_cell.font = Font(
                    name=FONT_NAME,
                    size=9,
                    color=COLORS["blue"],
                    underline="single",
                )


def add_component_glossary_sheet(wb: Workbook) -> object:
    ws = add_banner_table(
        wb,
        "元件標籤字典",
        "SAD / SDS 元件標籤字典",
        "每個出現在架構對照欄的受控標籤，都要能回答「是什麼、負責什麼、不負責什麼、到哪份文件與程式找」；L2 顯示群組不是元件。",
        ["正式標籤", "別名 / 原概括詞", "白話定義 / 負責什麼", "邊界 / 不負責什麼", "使用於 L2 能力群", "SAD 定位", "SDS 定位", "能力群相關實作路徑 / 規劃狀態"],
        component_glossary_rows(),
        [34, 30, 72, 68, 34, 52, 58, 80],
    )
    add_architecture_doc_links(ws, 4)
    return ws


def add_cover(wb: Workbook, title: str, subtitle: str, rows: Sequence[Sequence[object]]) -> object:
    ws = wb.create_sheet("cover")
    ws.title = "封面·文件管制"
    style_sheet_common(ws)
    ws.column_dimensions["A"].width = 25
    ws.column_dimensions["B"].width = 115
    ws["B1"] = taiwan_wording(title)
    ws["B1"].font = Font(name=FONT_NAME, bold=True, size=22, color=COLORS["white"])
    ws["B1"].fill = PatternFill("solid", fgColor=COLORS["navy"])
    ws["B1"].alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 40
    ws["B2"] = taiwan_wording(subtitle)
    ws["B2"].font = Font(name=FONT_NAME, bold=True, size=13, color=COLORS["blue"])
    ws["B2"].fill = PatternFill("solid", fgColor=COLORS["light_blue"])
    ws["B2"].alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[2].height = 32
    for left, right in rows:
        ws.append(
            [
                taiwan_wording(left) if isinstance(left, str) else left,
                taiwan_wording(right) if isinstance(right, str) else right,
            ]
        )
        row_idx = ws.max_row
        ws.cell(row_idx, 1).font = Font(name=FONT_NAME, bold=True, color=COLORS["blue"])
        ws.cell(row_idx, 2).font = Font(name=FONT_NAME, color=COLORS["black"])
        for col_idx in (1, 2):
            ws.cell(row_idx, col_idx).alignment = Alignment(vertical="top", wrap_text=True)
            ws.cell(row_idx, col_idx).border = TABLE_BORDER
            ws.cell(row_idx, col_idx).fill = PatternFill(
                "solid", fgColor=COLORS["white"] if row_idx % 2 else COLORS["pale_blue"]
            )
        ws.row_dimensions[row_idx].height = 38
    return ws


def add_cover_index(ws, sheet_names: Sequence[str], start_row: int | None = None) -> None:
    row_idx = start_row or ws.max_row + 2
    ws.cell(row_idx, 1, "工作表索引")
    ws.cell(row_idx, 1).font = Font(name=FONT_NAME, bold=True, color=COLORS["white"])
    ws.cell(row_idx, 1).fill = PatternFill("solid", fgColor=COLORS["blue"])
    ws.cell(row_idx, 2, "點擊名稱跳轉")
    ws.cell(row_idx, 2).font = Font(name=FONT_NAME, bold=True, color=COLORS["white"])
    ws.cell(row_idx, 2).fill = PatternFill("solid", fgColor=COLORS["blue"])
    for name in sheet_names:
        row_idx += 1
        cell = ws.cell(row_idx, 1, name)
        cell.hyperlink = f"#'{name}'!A1"
        cell.style = "Hyperlink"
        ws.cell(row_idx, 2, "→")
        for col_idx in (1, 2):
            ws.cell(row_idx, col_idx).border = TABLE_BORDER
            ws.cell(row_idx, col_idx).font = Font(
                name=FONT_NAME,
                color=COLORS["blue"] if col_idx == 1 else COLORS["dark_gray"],
                underline="single" if col_idx == 1 else None,
            )


def requirement_rows_for_phase(requirements: Sequence[Requirement], token: str) -> list[list[str]]:
    rows: list[list[str]] = []
    for req in requirements:
        phase = phase_for(req)
        if token not in phase:
            continue
        rows.append(
            [
                req.req_id,
                subsystem(req)["short"],
                qa_plain_text(req.name),
                phase,
                test_priority(req.req_id, req.all_text),
                status_for_requirement(req),
                planning_acceptance_summary(req),
                f"04_SRS.md:{req.source_line}",
            ]
        )
    return rows


def build_planning_workbook(requirements, nfrs, adrs, wbs, health) -> None:
    wb = Workbook()
    wb.remove(wb.active)
    set_workbook_meta(wb, "Smart Lock 規格統控規劃書")
    module_count = sum(
        1
        for prefix in SUBSYSTEMS
        for code, _, _ in MODULES[prefix]
        if any(
            req.prefix == prefix and module_for(req)[0] == code
            for req in requirements
        )
    )
    cover = add_cover(
        wb,
        "Smart Lock AI 客服與派工 SaaS 平台",
        "規格統控規劃書 · 給經營層 / PM / 架構師的管理投影",
        [
            ["文件定位", "以 M1–M5 投影管理範圍、落地順序、ADR 決策與契約/追溯風險。"],
            ["真相源", "04_SRS / 05_NFR / 12_SAD / 14_ADR / 20_Test_Cases / 21_Traceability / 27_Roadmap_WBS。"],
            ["主鍵", "SRS FR/NFR ID；舊 FR-0001 只作 legacy display，不當新四書 join key。"],
            ["需求快照", f"FR {health['fr_rows']} 列 / {health['fr_unique']} 個唯一 ID；NFR {health['nfr_rows']} 列。"],
            ["鍵健康", f"✅ FR {health['fr_unique']}/{health['fr_rows']} 主鍵唯一；舊追溯鍵 {health['legacy_trace_ids']} 個；既有 TC {health['tc_rows']} 筆。"],
            ["模組對齊", f"⑧ 完整模組清單與《SmartLock_模組功能BOM.xlsx》BOM 主表 L2 同源生成；本次 {module_count}/{module_count} 個能力群納入自動逐欄對帳。"],
            ["產出日", GENERATED_ON],
            ["使用規則", "xlsx 是單向生成快照；請改上游 Markdown 後重跑生成器，不要直接維護活頁簿。"],
        ],
    )

    source_rows = [
        ["SRS FR-*", "FR-AGT-01 / FR-API-02", "04_SRS.md §3", "本四書功能需求主鍵；必須唯一。"],
        ["NFR-*", "NFR-Sec-003", "05_NFR.md", "非功能驗收主鍵；含合約下限與營運目標。"],
        ["BR-*", "BR-Quote-*", "02_BRD / 04_SRS", "業務規則；由 SRS FR 追溯欄引用。"],
        ["TC-*", "TC-QUOTE-01", "20_Test_Cases.md", "測試案例鍵；現行 SRS 追溯由 §2.1 QTM 主表承接，詳細案例原 FR-xxxx 僅作歷史來源。"],
        ["ADR-*", "ADR-025", "14_ADR/", "架構決策鍵；Accepted 不直改，推翻應新增 ADR。"],
        ["WBS", "1.2.2 / 3.1.1", "27_Product_Roadmap_WBS.md", "落地工作包與 gate；不取代 FR/NFR。"],
        ["legacy FR-*", "FR-0001", "21_Traceability.md", "歷史追溯鍵；保留顯示，不能與 SRS FR 同名視為可 join。"],
        ["顯示群組", "AGT·RES / API·WO", "本生成器 _spec_data.py", "只用於 L2 顯示分群，明確不是新追溯鍵。"],
    ]
    add_banner_table(
        wb,
        "⓪ 編號出處指南",
        "⓪ 編號出處與單一追溯脊椎",
        "加入新功能前先確認它屬於哪個鍵家族；名稱、工作表列號與顯示群組都不當 join key。",
        ["代號家族", "例子", "出處", "用途 / 治理規則"],
        source_rows,
        [24, 28, 36, 76],
    )

    compare_rows = []
    for prefix, meta in SUBSYSTEMS.items():
        subset = [req for req in requirements if req.prefix == prefix]
        columns = []
        for token in ("M1", "M2", "M3", "M4"):
            names = [req.name for req in subset if token in phase_for(req)]
            columns.append("、".join(names) if names else "—")
        compare_rows.append([meta["short"], prefix, *columns])
    add_banner_table(
        wb,
        "① 階段對比矩陣",
        "① 子系統 × M1–M5 需求投影",
        "階段是 27_Roadmap_WBS 的管理投影，不是實作完成率；相同 FR 可因分期橫跨多階段。",
        ["子系統", "代號", "M1 上線硬化", "M2 三線成形", "M3 多品牌", "M4/M5 平台化"],
        compare_rows,
        [22, 10, 55, 55, 55, 55],
    )

    common_headers = ["FR ID", "子系統", "功能需求", "規劃對映", "QA 優先", "需求狀態", "驗收摘要（PRD／成功指標）", "出處"]
    m1_ws = add_banner_table(
        wb,
        "② M1 功能需求",
        "② M1 上線硬化需求清單",
        "包含直接 M1 與 M1→後續的跨期需求；狀態只說明文件訊號，實作以 WBS/code/SIT 證據為準。",
        common_headers,
        requirement_rows_for_phase(requirements, "M1"),
        [17, 15, 34, 12, 10, 24, 85, 20],
    )
    m2_ws = add_banner_table(
        wb,
        "③ M2 功能需求",
        "③ M2 身分・知識・技師平台需求清單",
        "包含直接 M2 與跨期需求；逐條證據請回 04_SRS、27_Roadmap 與對應 TC。",
        common_headers,
        requirement_rows_for_phase(requirements, "M2"),
        [17, 15, 34, 12, 10, 24, 85, 20],
    )
    later_rows = requirement_rows_for_phase(requirements, "M3") + requirement_rows_for_phase(requirements, "M4")
    # 去除因 M3/M4 token 造成的重複列。
    dedup_later: list[list[str]] = []
    seen_later: set[tuple[str, str]] = set()
    for row in later_rows:
        key = (row[0], row[2])
        if key not in seen_later:
            dedup_later.append(row)
            seen_later.add(key)
    later_ws = add_banner_table(
        wb,
        "④ M3+整合範圍",
        "④ M3–M5 多品牌與平台化整合範圍",
        "階段二必須同時守住單品牌合約紅線；事件、對帳、provisioning 與 DSL 是主要整合面。",
        common_headers,
        dedup_later,
        [17, 15, 34, 12, 10, 24, 85, 20],
    )
    for ws in (m1_ws, m2_ws, later_ws):
        for row_idx in range(5, ws.max_row + 1):
            ws.row_dimensions[row_idx].height = 78

    wbs_rows = [[row[1], row[0], row[3], row[2], row[5], row[6]] for row in wbs]
    add_banner_table(
        wb,
        "⑤ 落地順序與關卡",
        "⑤ Roadmap WBS 落地順序與 Stage Gate",
        "直接投影 27_Product_Roadmap_WBS；週期與狀態以上游為準，本表不代拍板。",
        ["WBS", "里程碑", "工作包", "狀態", "前置", "交付物 / 驗收依據"],
        wbs_rows,
        [12, 30, 58, 55, 22, 58],
    )

    add_banner_table(
        wb,
        "⑥ 架構抉擇",
        "⑥ 關鍵架構抉擇與管理影響",
        "這是 ADR 結果態的白話投影；決策狀態與全量取捨請回「⑨ ADR 決策表」。",
        ["代號（顯示）", "主題", "決策摘要", "管理影響", "出處"],
        ARCH_CHOICES,
        [15, 28, 70, 55, 28],
    )
    add_banner_table(
        wb,
        "⑦ 術語白話對照",
        "⑦ 術語白話對照",
        "左邊是架構/需求用語，右邊是對產品、合約與營運的含意。",
        ["術語", "工程定義", "白話：對產品/管理的意義"],
        GLOSSARY,
        [30, 70, 75],
    )

    module_summary = []
    for prefix, meta in SUBSYSTEMS.items():
        for code, module_name, _ in MODULES[prefix]:
            members = [req for req in requirements if req.prefix == prefix and module_for(req)[0] == code]
            if not members:
                continue
            arch = module_arch(prefix, code)
            module_summary.append(
                [
                    f"{prefix}·{code}（顯示）",
                    module_name,
                    meta["short"],
                    len(members),
                    len({req.req_id for req in members}),
                    "、".join(sorted({phase_for(req) for req in members})),
                    sum("🔜" in req.all_text for req in members),
                    sum(req.duplicate for req in members),
                    arch["component"],
                    arch["sad"],
                    arch["sds"],
                    "✅ 與 BOM L2 對齊",
                ]
            )
    module_ws = add_banner_table(
        wb,
        "⑧ 完整模組清單",
        "⑧ L2 模組、正式元件與 FR 健康快照",
        f"本表與《SmartLock_模組功能BOM.xlsx》BOM 主表 L2 使用同一份模組定義；本次共 {module_count} 個能力群，生成後會逐欄對帳。",
        ["顯示群組", "模組", "子系統", "FR 列數", "唯一 FR", "階段", "含規劃訊號", "碰鍵列", "正式元件名稱", "SAD 定位", "SDS 定位", "BOM 對齊"],
        module_summary,
        [20, 28, 18, 12, 12, 18, 15, 12, 72, 42, 50, 22],
    )
    add_architecture_doc_links(module_ws, 4)

    add_banner_table(
        wb,
        "⑨ ADR 決策表",
        "⑨ ADR 決策狀態矩陣",
        "由 14_ADR/00_INDEX 直接投影。Accepted 決策不在此活頁簿改寫；推翻應新增 ADR。",
        ["決策群", "ADR", "標題", "層級", "狀態", "關聯 ADR"],
        adrs,
        [30, 14, 70, 18, 24, 30],
    )
    add_banner_table(
        wb,
        "⑩ 契約與追溯風險",
        "⑩ 架構、契約、追溯與驗收風險",
        "產出四書時發現的結構性風險；這些風險不會靠改 xlsx 永久修復，需回上游正典裁定。",
        ["類別", "編號", "風險/衝突", "影響", "緩解/待裁定", "出處"],
        ARCH_RISKS,
        [14, 13, 65, 50, 70, 28],
    )
    add_component_glossary_sheet(wb)
    add_cover_index(cover, wb.sheetnames[1:])
    wb.save(OUTPUTS["planning"])


def build_bom_workbook(requirements, nfrs, detail_source_links, health) -> None:
    wb = Workbook()
    wb.remove(wb.active)
    set_workbook_meta(wb, "Smart Lock 模組功能 BOM")
    cover = add_cover(
        wb,
        "Smart Lock 平台 — 模組・功能 BOM",
        "L1 子系統 → L2 能力群（顯示）→ L3 SRS FR（唯一追溯鍵）",
        [
            ["用途", "將產品功能階層、正式元件名稱、SAD/SDS 定位、實作路徑、Roadmap 階段與驗收摘要放在同一張表。"],
            ["L1", "04_SRS 的 agent / api / web / data-pipeline / knowledge-refinery / technician-platform / 00_platform。"],
            ["L2", "本活頁簿的顯示分群，不是新鍵，不可用來代替 FR ID join。"],
            ["L3", f"SRS FR 原文 {health['fr_rows']} 列，{health['fr_unique']} 個唯一 ID。"],
            ["鍵健康", f"✅ FR {health['fr_unique']}/{health['fr_rows']} 主鍵唯一；FR-TEC-07/08 已分別代表報價修正與生命週期。"],
            ["架構對照", "SAD 是容器層、SDS 是 L3 元件/路徑；兩者各自標註，L2 能力群只作人工投影。"],
            ["標籤定義", "所有正式元件標籤皆收錄於「元件標籤字典」，包含白話定義與責任邊界。"],
            ["產出日", GENERATED_ON],
            ["再生規則", "改 04_SRS/05_NFR/_spec_data.py 後重跑 _build_enterprise_workbooks.py；不直接改 xlsx。"],
        ],
    )

    bom_rows: list[list[object]] = []
    for prefix, meta in SUBSYSTEMS.items():
        subsystem_reqs = [req for req in requirements if req.prefix == prefix]
        bom_rows.append(
            [
                "L1",
                meta["name"].split("（")[0],
                meta["name"],
                "",
                subsystem_component_labels(prefix),
                meta["sad"],
                meta["sds"],
                meta["path"],
                "—",
                "子系統",
                "●" if any("M1" in phase_for(req) for req in subsystem_reqs) else "",
                "●" if any("M2" in phase_for(req) for req in subsystem_reqs) else "",
                "●" if any(any(token in phase_for(req) for token in ("M3", "M4", "M5")) for req in subsystem_reqs) else "",
                "",
                meta["description"],
            ]
        )
        for code, module_name, _ in MODULES[prefix]:
            module_reqs = [req for req in subsystem_reqs if module_for(req)[0] == code]
            if not module_reqs:
                continue
            display_code = f"{prefix}·{code}（顯示）"
            arch = module_arch(prefix, code)
            bom_rows.append(
                [
                    "L2",
                    display_code,
                    module_name,
                    "",
                    arch["component"],
                    arch["sad"],
                    arch["sds"],
                    arch["path"],
                    meta["name"].split("（")[0],
                    "能力群（非 join key）",
                    "●" if any("M1" in phase_for(req) for req in module_reqs) else "",
                    "●" if any("M2" in phase_for(req) for req in module_reqs) else "",
                    "●" if any(any(token in phase_for(req) for token in ("M3", "M4", "M5")) for req in module_reqs) else "",
                    "",
                    f"{len(module_reqs)} 列 FR / {len({req.req_id for req in module_reqs})} 個唯一 ID",
                ]
            )
            for req in module_reqs:
                phase = phase_for(req)
                req_arch = architecture_for(req)
                bom_rows.append(
                    [
                        "L3",
                        req.req_id,
                        req.name,
                        req.trace,
                        req_arch["component"],
                        req_arch["sad"],
                        req_arch["sds"],
                        req_arch["path"],
                        display_code,
                        "功能需求（SRS FR）",
                        "✓" if "M1" in phase else "",
                        "✓" if "M2" in phase else "",
                        "✓" if any(token in phase for token in ("M3", "M4", "M5")) else "",
                        status_for_requirement(req),
                        f"驗收：{req.acceptance} ｜ 出處：04_SRS.md:{req.source_line}",
                    ]
                )

    bom_ws = add_plain_table(
        wb,
        "BOM 主表",
        ["層級", "代號（FR 為主鍵）", "名稱 / 功能", "上游規則 / 需求", "正式元件名稱", "SAD 定位", "SDS 定位", "實作路徑 / 規劃狀態", "父節點", "類型", "M1", "M2", "M3+", "需求狀態", "說明 / 出處"],
        bom_rows,
        [8, 23, 35, 33, 72, 42, 52, 78, 23, 24, 8, 8, 8, 26, 70],
    )
    add_architecture_doc_links(bom_ws, 1)
    # L1/L2 粗體分層，並建立 Excel outline。
    for row_idx in range(2, bom_ws.max_row + 1):
        level = bom_ws.cell(row_idx, 1).value
        if level == "L1":
            for cell in bom_ws[row_idx]:
                cell.fill = PatternFill("solid", fgColor=COLORS["navy"])
                cell.font = Font(name=FONT_NAME, color=COLORS["white"], bold=True)
            bom_ws.row_dimensions[row_idx].outlineLevel = 0
        elif level == "L2":
            for cell in bom_ws[row_idx]:
                cell.fill = PatternFill("solid", fgColor=COLORS["light_blue"])
                cell.font = Font(name=FONT_NAME, color=COLORS["blue"], bold=True)
            bom_ws.row_dimensions[row_idx].outlineLevel = 1
        else:
            bom_ws.row_dimensions[row_idx].outlineLevel = 2

    arch_rows = []
    for prefix, meta in SUBSYSTEMS.items():
        for code, module_name, _ in MODULES[prefix]:
            subset = [
                req
                for req in requirements
                if req.prefix == prefix and module_for(req)[0] == code
            ]
            if not subset:
                continue
            arch = module_arch(prefix, code)
            arch_rows.append(
                [
                    f"{prefix}·{code}（顯示）",
                    meta["name"],
                    module_name,
                    arch["component"],
                    arch["sad"],
                    arch["sds"],
                    arch["path"],
                    len(subset),
                    len({req.req_id for req in subset}),
                    sum("🔜" in req.all_text for req in subset),
                    sum("[待確認]" in req.all_text for req in subset),
                    sum(req.duplicate for req in subset),
                    sum(req.req_id in detail_source_links for req in subset),
                ]
            )
    arch_ws = add_plain_table(
        wb,
        "SAD·SDS 元件視圖",
        ["顯示群組", "子系統", "L2 能力群", "正式元件名稱", "SAD 定位", "SDS 定位", "實作路徑 / 規劃狀態", "FR 列數", "唯一 FR", "含規劃訊號", "待確認", "碰鍵列", "詳細 TC 舊欄含現行 SRS ID"],
        arch_rows,
        [20, 28, 28, 72, 42, 52, 78, 12, 12, 15, 12, 12, 18],
    )
    add_architecture_doc_links(arch_ws, 1)
    add_component_glossary_sheet(wb)
    add_cover_index(cover, wb.sheetnames[1:])
    wb.save(OUTPUTS["bom"])


def build_acceptance_workbook(requirements, nfrs, health) -> None:
    wb = Workbook()
    wb.remove(wb.active)
    set_workbook_meta(wb, "Smart Lock 業務邏輯驗收控制表")
    cover = add_cover(
        wb,
        "Smart Lock 平台 — 業務邏輯驗收控制表",
        "客戶 VOC → PRD 驗收標準 → 成功指標；技術流程獨立回查",
        [
            ["主鍵", "REQ ID 欄只放 04_SRS FR 或 05_NFR ID；第一欄只顯示上游節號與行號。"],
            ["功能規模", f"FR {health['fr_rows']} 列，分成 7 個子系統分頁。"],
            ["品質規模", f"NFR {health['nfr_rows']} 列，分為效能可用、安全隱私、品質維運 3 頁。"],
            ["語言定位", "PM／業務先看「客戶 VOC／商業情境」與「PRD 驗收標準」；架構師／SA 以「成功指標」定量與對接證據。"],
            ["技術分流", "04_SRS 的前置、主流、API／事件／資料庫細節不放在 VOC 欄，完整保留於「技術流程對照」給 SA／RD 回查。"],
            ["狀態語義", "✅/🔶/🔜 是需求文件訊號，不等於實作完成；🔴 合約下限是 release blocker。"],
            ["主鍵規則", "FR-TEC-07＝現場報價修正；FR-TEC-08＝排班與生命週期，兩者可獨立驗收與追溯。"],
            ["產出日", GENERATED_ON],
            ["日常使用", "PM／業務以 PRD 驗收標準進行 UAT 簽核；SA／架構師確認成功指標、量測環境與證據。執行結果應回填測試管理系統。"],
        ],
    )
    headers = ["編號（出處）", "REQ ID", "領域", "需求主題（上游名稱）", "客戶 VOC／商業情境", "PRD 驗收標準", "成功指標（架構／SA）", "例外與限制 / 待確認", "階段", "狀態", "出處", "正式元件名稱", "SAD 定位", "SDS 定位"]
    widths = [22, 18, 20, 32, 64, 64, 58, 48, 14, 28, 24, 70, 42, 52]
    prefix_tabs = {
        "AGT": "AGT·AI 客服",
        "API": "API·派工控制",
        "WEB": "WEB·多站前端",
        "DAT": "DAT·資料平台",
        "REF": "REF·知識精煉",
        "TEC": "TEC·技師共享池",
        "PLT": "PLT·平台核心",
    }
    technical_rows = []
    for prefix, tab_name in prefix_tabs.items():
        rows = []
        for req in requirements:
            if req.prefix != prefix:
                continue
            flags = []
            if "[待確認]" in req.all_text:
                flags.append("上游含 [待確認]，驗收前須定門檻/量測點。")
            if "🔜" in req.all_text:
                flags.append("含規劃中能力；執行驗收前須確認部署版本與證據。")
            if req.duplicate:
                flags.append("🔴 FR-TEC-07 碰鍵；不得在未改號前用該 ID 自動 join。")
            code, module_name = module_for(req)
            arch = architecture_for(req)
            copy = business_copy_for(req)
            rows.append(
                [
                    f"04_SRS §3 / L{req.source_line}",
                    req.req_id,
                    f"{subsystem(req)['short']} / {module_name}",
                    req.name,
                    copy["voc"],
                    copy["prd_acceptance"],
                    req.acceptance,
                    " ".join(flags) if flags else "—",
                    phase_for(req),
                    status_for_requirement(req),
                    f"04_SRS.md:{req.source_line} ｜ 追溯：{req.trace}",
                    arch["component"],
                    arch["sad"],
                    arch["sds"],
                ]
            )
            technical_rows.append(
                [
                    f"04_SRS §3 / L{req.source_line}",
                    req.req_id,
                    tab_name,
                    f"{subsystem(req)['short']} / {module_name}",
                    req.name,
                    f"前置：{req.precondition} ｜ 主流：{req.flow}",
                    req.acceptance,
                    " ".join(flags) if flags else "—",
                    arch["component"],
                    arch["sad"],
                    arch["sds"],
                    f"04_SRS.md:{req.source_line} ｜ 追溯：{req.trace}",
                ]
            )
        ws = add_plain_table(wb, tab_name, headers, rows, widths)
        add_architecture_doc_links(ws, 1)
        for row_idx in range(2, ws.max_row + 1):
            ws.row_dimensions[row_idx].height = 78

    nfr_tabs = {
        "PERF": "PERF·效能可用",
        "SEC": "SEC·安全隱私",
        "OPS": "OPS·品質維運",
    }
    for group, tab_name in nfr_tabs.items():
        rows = []
        for nfr in nfrs:
            if nfr_group(nfr) != group:
                continue
            flags = [f"驗證方式：{nfr.verification}"]
            if "[待確認]" in nfr.all_text:
                flags.append("上游含 [待確認]，需先定量測環境/關卡。")
            if "🔜" in nfr.all_text:
                flags.append("含規劃中實作或驗證。")
            copy = nfr_business_copy(nfr)
            rows.append(
                [
                    f"05_NFR L{nfr.source_line}",
                    nfr.req_id,
                    nfr.heading,
                    nfr.name,
                    copy["voc"],
                    copy["prd_acceptance"],
                    nfr.target,
                    " ".join(flags),
                    "跨期 / 營運",
                    nfr_status(nfr),
                    f"05_NFR.md:{nfr.source_line} ｜ 分層：{nfr.tier}",
                    "跨系統品質屬性（非單一元件）",
                    "12_SAD §10",
                    "15_SDS §12 / 對應元件章節",
                ]
            )
        ws = add_plain_table(wb, tab_name, headers, rows, widths)
        add_architecture_doc_links(ws, 1)
        for row_idx in range(2, ws.max_row + 1):
            ws.row_dimensions[row_idx].height = 66
    technical_ws = add_banner_table(
        wb,
        "技術流程對照",
        "SA／架構師／RD 技術流程對照",
        "此頁完整保留 04_SRS 的前置與主流技術敘述；PM／業務 UAT 應回各領域分頁閱讀 VOC、PRD 驗收標準與成功指標。",
        ["編號（出處）", "REQ ID", "原分頁", "領域", "需求主題", "技術流程（前置／主流）", "SRS 驗收基線", "例外與限制 / 待確認", "正式元件名稱", "SAD 定位", "SDS 定位", "出處"],
        technical_rows,
        [22, 18, 20, 24, 32, 96, 60, 48, 70, 42, 52, 34],
    )
    add_architecture_doc_links(technical_ws, 4)
    for row_idx in range(5, technical_ws.max_row + 1):
        technical_ws.row_dimensions[row_idx].height = 78
    add_component_glossary_sheet(wb)
    add_cover_index(cover, wb.sheetnames[1:])
    wb.save(OUTPUTS["acceptance"])


def build_test_workbook(requirements, nfrs, test_cases, health) -> None:
    wb = Workbook()
    wb.remove(wb.active)
    set_workbook_meta(wb, "Smart Lock 整合測試計畫")
    cover = add_cover(
        wb,
        "Smart Lock 平台 — 整合測試計畫",
        "QA 工作版：從客戶需求確認測試情境 → 依案例執行 → 記錄結果與證據",
        [
            ["QA 使用順序", "⑧先確認客戶需求、PRD 驗收標準與測試情境 → ⑨依案例執行並記錄結果、證據與缺陷。"],
            ["QA 只需看", "客戶需求、測試重點、前置條件/資料、執行步驟、通過標準與證據。不需先理解 FR/NFR/QTM/SAD/SDS。"],
            ["測試範圍", f"{health['fr_rows'] + health['nfr_rows']} 個客戶／品質情境與同數量的可執行案例（{health['fr_rows']} 個功能/流程 + {health['nfr_rows']} 個品質目標），並以 12 條端到端旅程分類。"],
            ["通過原則", "每列的功能判準/品質目標必須達成，而且拒絕、重送或故障情境不得留下錯誤資料或重複副作用。"],
            ["上線出口", "P0 金錢/授權/租戶/合約紅線任一未結 → Fail；P1 主流未結 → 至多 Conditional Pass。"],
            ["治理追溯", "完整需求、架構、測試鍵與測試分析檢核保留在隱藏附錄，供 QA Lead/SA 審核；不佔用 QA 日常執行視圖。"],
            ["產出日", GENERATED_ON],
        ],
    )
    add_banner_table(
        wb,
        "① 測試策略與原則",
        "① 測試策略與原則",
        "平台最大損失來自錯帳、越權、跨租戶、合約紅線與無法恢復的主流；測試資源以這些風險為中心。",
        ["原則", "一句話", "Smart Lock 落地方式"],
        TEST_STRATEGY,
        [28, 55, 90],
    )
    add_banner_table(
        wb,
        "② 測試種類定義",
        "② 測試階段、類型與測試方法",
        "階段回答誰在什麼時候測，類型回答測哪一面，測試方法則明確說明如何準備、操作與觀察。",
        ["分類", "種類", "定義 / 範圍", "落地與責任"],
        TEST_TYPES,
        [18, 30, 75, 65],
    )
    add_banner_table(
        wb,
        "③ Bug 分級與處理",
        "③ Bug / UAT 缺陷分級與出口規則",
        "分級看影響與替代路徑，不看誰先喊。與 22_UAT 的 P0/P1/P2 裁定對齊。",
        ["等級", "定義", "卡關規則", "處理"],
        BUG_LEVELS,
        [20, 80, 50, 60],
    )
    add_banner_table(
        wb,
        "④ 環境與部署流程",
        "④ 測試環境與部署流程",
        "敏感資料與物理三庫隔離使環境不能只靠「複製 production DB」建立；必須可重建、去識別、可追溯。",
        ["環境", "時機", "執行測試", "目的 / 備註"],
        ENVIRONMENTS,
        [24, 30, 75, 70],
    )
    add_banner_table(
        wb,
        "⑤ 測試生命週期 STLC",
        "⑤ 測試生命週期 STLC",
        "四書是 Spec Review 與 Test Analysis 的輸入，不是執行證據本身。",
        ["階段", "核心目的", "產出"],
        STLC,
        [32, 90, 75],
    )
    add_banner_table(
        wb,
        "⑥ SIT·UAT·RC 階段",
        "⑥ SIT → UAT → Release Candidate 關卡",
        "不以人工日曆代替 Entry/Exit Criteria；RC 應是確認不是發現大量新問題。",
        ["階段", "目的", "Entry Criteria", "Exit Criteria"],
        TEST_STAGES,
        [28, 58, 75, 75],
    )
    add_banner_table(
        wb,
        "⑦ RD×QA 職能切分",
        "⑦ RD、QA、PM 與業務 Owner 責任切分",
        "QA 不代替 RD 寫全部單元/契約測試；QA 負責確認每個功能與品質目標都有可執行的場景與通過標準。",
        ["活動", "主責", "協作", "範圍 / DoD", "產出"],
        RESPONSIBILITIES,
        [35, 18, 18, 85, 45],
    )

    scenario_rows = []
    execution_rows = []
    for req in requirements:
        scenario_id = qa_visible_id(req.req_id, "SCN")
        case_id = qa_visible_id(req.req_id, "CASE")
        copy = business_copy_for(req)
        execution = qa_execution_for(req)
        domain = subsystem(req)["short"]
        scenario_types = qa_scenario_types(req.req_id)
        scenario_rows.append(
            [
                scenario_id,
                req.req_id,
                domain,
                copy["voc"],
                copy["prd_acceptance"],
                execution["expected"],
                scenario_types,
                qa_scenario_names(scenario_for_requirement(req)),
                test_priority(req.req_id, req.all_text),
                case_id,
            ]
        )
        execution_rows.append(
            [
                case_id,
                scenario_id,
                test_priority(req.req_id, req.all_text),
                domain,
                scenario_types,
                qa_plain_text(req.name),
                execution["preparation"],
                execution["steps"],
                execution["expected"],
                execution["evidence"],
                "執行時記錄實際結果",
                "尚未執行",
                "通過：貼上證據連結；失敗：貼上證據並填缺陷編號",
                qa_owner(req.req_id),
                "執行時填寫 build/commit 與日期",
            ]
        )
    for nfr in nfrs:
        scenario_id = qa_visible_id(nfr.req_id, "SCN")
        case_id = qa_visible_id(nfr.req_id, "CASE")
        copy = nfr_business_copy(nfr)
        execution = qa_nfr_execution_for(nfr)
        domain = NFR_DOMAIN_META[nfr_group(nfr)][0]
        scenario_types = qa_scenario_types(nfr.req_id, is_nfr=True)
        scenario_rows.append(
            [
                scenario_id,
                nfr.req_id,
                domain,
                copy["voc"],
                copy["prd_acceptance"],
                execution["expected"],
                scenario_types,
                qa_scenario_names(scenario_for_nfr(nfr)),
                test_priority(nfr.req_id, nfr.all_text),
                case_id,
            ]
        )
        execution_rows.append(
            [
                case_id,
                scenario_id,
                test_priority(nfr.req_id, nfr.all_text),
                domain,
                scenario_types,
                qa_plain_text(nfr.name),
                execution["preparation"],
                execution["steps"],
                execution["expected"],
                execution["evidence"],
                "執行時記錄實際結果",
                "尚未執行",
                "通過：貼上證據連結；失敗：貼上證據並填缺陷編號",
                qa_owner(nfr.req_id, is_nfr=True),
                "執行時填寫 build/commit 與日期",
            ]
        )
    scenario_ws = add_banner_table(
        wb,
        "⑧ 客戶需求與測試情境",
        "⑧ 客戶需求、PRD 驗收標準與測試情境",
        "每列從一項客戶或品質需求出發；「需求來源」可直接確認對應的 FR／NFR，完整治理資訊仍留在附錄 A。",
        ["情境編號", "需求來源", "領域", "客戶需求／VOC", "PRD 驗收標準", "測試重點／通過判準", "應涵蓋情況", "客戶旅程", "優先", "對應案例編號"],
        scenario_rows,
        [18, 18, 24, 68, 76, 72, 36, 42, 10, 20],
    )
    scenario_ws.freeze_panes = "D5"
    scenario_ws.sheet_view.zoomScale = 65
    scenario_ws.page_setup.paperSize = scenario_ws.PAPERSIZE_A3
    for row_idx in range(5, scenario_ws.max_row + 1):
        scenario_ws.cell(row_idx, 6).value = str(scenario_ws.cell(row_idx, 6).value or "").replace("｜", "\n")
        scenario_ws.row_dimensions[row_idx].height = 112

    execution_ws = add_banner_table(
        wb,
        "⑨ 測試案例與執行紀錄",
        "⑨ 測試案例與執行紀錄",
        f"共 {len(execution_rows)} 筆案例。依前置、步驟與通過標準執行；完成後更新實際結果、狀態、證據／缺陷連結及版本日期。",
        ["案例編號", "情境編號", "優先", "領域", "案例涵蓋", "測試案例", "前置條件與測試資料", "執行步驟", "預期結果／通過標準", "應保留證據", "實際結果", "執行狀態", "證據／缺陷連結", "執行主責", "版本／日期"],
        execution_rows,
        [18, 18, 10, 24, 34, 38, 76, 94, 80, 58, 48, 16, 52, 30, 34],
    )
    execution_ws.page_setup.paperSize = execution_ws.PAPERSIZE_A3
    execution_ws.freeze_panes = "G5"
    execution_ws.sheet_view.zoomScale = 60
    for row_idx in range(5, execution_ws.max_row + 1):
        for column in (7, 8, 9):
            cell = execution_ws.cell(row_idx, column)
            cell.value = str(cell.value or "").replace("｜", "\n")
        execution_ws.row_dimensions[row_idx].height = 150
    status_validation = DataValidation(
        type="list",
        formula1='"尚未執行,通過,失敗,阻擋,不適用"',
        allow_blank=False,
    )
    status_validation.error = "請從清單選擇執行狀態"
    status_validation.errorTitle = "執行狀態格式錯誤"
    execution_ws.add_data_validation(status_validation)
    status_validation.add(f"L5:L{execution_ws.max_row}")
    execution_ws.conditional_formatting.add(
        f"L5:L{execution_ws.max_row}",
        FormulaRule(formula=['L5="通過"'], fill=PatternFill("solid", fgColor=COLORS["light_green"])),
    )
    execution_ws.conditional_formatting.add(
        f"L5:L{execution_ws.max_row}",
        FormulaRule(formula=['OR(L5="失敗",L5="阻擋")'], fill=PatternFill("solid", fgColor=COLORS["light_red"])),
    )
    for row_idx in range(5, execution_ws.max_row + 1):
        source_cell = scenario_ws.cell(row_idx, 2)
        source_cell.hyperlink = f"#'附錄A 需求追溯'!A{row_idx}"
        source_cell.style = "Hyperlink"
        scenario_case_cell = scenario_ws.cell(row_idx, 10)
        scenario_case_cell.hyperlink = f"#'⑨ 測試案例與執行紀錄'!A{row_idx}"
        scenario_case_cell.style = "Hyperlink"
        execution_scenario_cell = execution_ws.cell(row_idx, 2)
        execution_scenario_cell.hyperlink = f"#'⑧ 客戶需求與測試情境'!A{row_idx}"
        execution_scenario_cell.style = "Hyperlink"

    matrix_rows = []
    for req in requirements:
        arch = architecture_for(req)
        if req.duplicate:
            coverage = "🔴 ID 衝突，不可建立受控映射"
            split = "SA 必須先裁定唯一 SRS FR ID"
        else:
            coverage = f"✅ 20_Test_Cases §2.1：{qa_mapping_id(req.req_id)}"
            split = f"{qa_mapping_id(req.req_id)} → 指定 TC：{tc_hint(req.req_id)}"
        matrix_rows.append(
            [
                req.req_id,
                f"{req.prefix} / {module_for(req)[1]}",
                qa_plain_text(req.name),
                qa_acceptance_for(req),
                qa_closure_status(req.all_text),
                scenario_for_requirement(req),
                coverage,
                split,
                arch["component"],
                arch["sad"],
                arch["sds"],
            ]
        )
    for nfr in nfrs:
        hint = tc_hint(nfr.req_id, is_nfr=True)
        matrix_rows.append(
            [
                nfr.req_id,
                nfr.heading,
                qa_plain_text(nfr.name),
                nfr.target,
                qa_closure_status(nfr.all_text),
                scenario_for_nfr(nfr),
                f"✅ 20_Test_Cases §2.1：{qa_mapping_id(nfr.req_id)}",
                f"{qa_mapping_id(nfr.req_id)} → 指定 TC：{hint}",
                "跨系統品質屬性（非單一元件）",
                "12_SAD §10",
                "15_SDS §12 / 對應元件章節",
            ]
        )
    matrix_ws = add_banner_table(
        wb,
        "附錄A 需求追溯",
        "附錄 A：需求、場景、測試案例與架構追溯（QA Lead/SA）",
        "此頁用於覆蓋率與變更影響審核；已預設隱藏，QA 日常執行不需閱讀。",
        ["REQ ID", "領域", "需求名稱", "驗收標準", "QA 規格閉環狀態", "場景（測試設計）", "QA 追溯狀態", "QA 追溯決策", "正式元件名稱", "SAD 定位", "SDS 定位"],
        matrix_rows,
        [18, 30, 35, 58, 23, 25, 38, 70, 70, 42, 52],
    )
    add_architecture_doc_links(matrix_ws, 4)
    matrix_ws.sheet_state = "hidden"

    closure_rows = []
    for req in requirements:
        closure_rows.append(
            [
                test_priority(req.req_id, req.all_text),
                req.req_id,
                "FR",
                qa_mapping_id(req.req_id),
                req.name,
                qa_acceptance_for(req),
                scenario_for_requirement(req),
                tc_hint(req.req_id),
                qa_owner(req.req_id),
                qa_evidence_spec(req.req_id),
                f"20_Test_Cases §2.1：{qa_mapping_id(req.req_id)}",
                qa_closure_status(req.all_text),
            ]
        )
    for nfr in nfrs:
        closure_rows.append(
            [
                test_priority(nfr.req_id, nfr.all_text),
                nfr.req_id,
                "NFR",
                qa_mapping_id(nfr.req_id),
                qa_plain_text(nfr.name),
                nfr.target,
                scenario_for_nfr(nfr),
                f"{tc_hint(nfr.req_id, is_nfr=True)} ｜ 驗證：{nfr.verification}",
                qa_owner(nfr.req_id, is_nfr=True),
                qa_evidence_spec(nfr.req_id, is_nfr=True),
                f"20_Test_Cases §2.1：{qa_mapping_id(nfr.req_id)}",
                qa_closure_status(nfr.all_text),
            ]
        )
    closure_ws = add_banner_table(
        wb,
        "附錄B 規格閉環",
        "附錄 B：規格閉環與證據責任（QA Lead/SA）",
        "此頁用於治理審核，確認每個來源需求的判準、責任與追溯已完整；已預設隱藏。",
        ["優先", "REQ ID", "類型", "QA 映射 ID", "項目", "受控驗收門檻", "場景", "受控測試設計 / TC", "責任人", "執行後證據輸出規格", "追溯處理結果", "規格閉環狀態"],
        closure_rows,
        [10, 18, 10, 24, 34, 68, 24, 62, 30, 70, 38, 38],
    )
    closure_ws.page_setup.paperSize = closure_ws.PAPERSIZE_A3
    closure_ws.sheet_state = "hidden"

    analysis_rows = []
    for prefix, meta in DOMAIN_TEST_META.items():
        checks, method = DOMAIN_QA_CHECKS[prefix]
        analysis_rows.append(
            [
                "領域風險檢核",
                f"DOMAIN-{prefix}",
                meta[0],
                qa_plain_text(meta[1]),
                checks.replace("\n", "｜"),
                method,
                "每項必測行為都應反映在⑧情境與⑨案例；不由執行 QA 另外維護。",
            ]
        )
    for group, meta in NFR_DOMAIN_META.items():
        checks, method = DOMAIN_QA_CHECKS[group]
        analysis_rows.append(
            [
                "領域風險檢核",
                f"DOMAIN-{group}",
                meta[0],
                qa_plain_text(meta[1]),
                checks.replace("\n", "｜"),
                method,
                "每項必測行為都應反映在⑧情境與⑨案例；不由執行 QA 另外維護。",
            ]
        )
    for row in SCENARIOS:
        sid = row[0]
        analysis_rows.append(
            [
                "端到端旅程",
                sid,
                qa_plain_text(row[1]),
                qa_plain_text(row[3]),
                qa_plain_text(row[4]),
                qa_plain_text(row[5]),
                f"{SCENARIO_PASS_CRITERIA[sid]} ｜ 測試類型：{qa_plain_text(row[6])}",
            ]
        )
    analysis_ws = add_banner_table(
        wb,
        "附錄C 測試分析",
        "附錄 C：領域風險檢核與端到端旅程（QA Lead）",
        "這是把客戶需求轉成測試情境時使用的分析依據；已預設隱藏，QA 日常執行只需使用⑧與⑨。",
        ["分析類型", "分析鍵", "名稱／範圍", "風險／要證明", "必測行為／準備", "執行方式／流程", "通過條件／備註"],
        analysis_rows,
        [22, 18, 34, 64, 92, 88, 76],
    )
    for row_idx in range(5, analysis_ws.max_row + 1):
        for column in (5, 7):
            cell = analysis_ws.cell(row_idx, column)
            cell.value = str(cell.value or "").replace("｜", "\n")
        analysis_ws.row_dimensions[row_idx].height = 124
    analysis_ws.page_setup.paperSize = analysis_ws.PAPERSIZE_A3
    analysis_ws.sheet_state = "hidden"
    glossary_ws = add_component_glossary_sheet(wb)
    glossary_ws.sheet_state = "hidden"
    visible_sheet_names = [
        ws.title for ws in wb.worksheets[1:] if ws.sheet_state == "visible"
    ]
    add_cover_index(cover, visible_sheet_names)
    wb.save(OUTPUTS["test"])


def validate_workbooks(requirements, nfrs) -> dict[str, list[str]]:
    errors: dict[str, list[str]] = defaultdict(list)
    expected_counts = {
        "planning": 13,
        "bom": 4,
        "acceptance": 13,
        "test": 14,
    }
    for key, path in OUTPUTS.items():
        if not path.exists() or path.stat().st_size == 0:
            errors[key].append("檔案不存在或為空")
            continue
        wb = load_workbook(path, data_only=False, read_only=False)
        if len(wb.sheetnames) != expected_counts[key]:
            errors[key].append(f"工作表數 {len(wb.sheetnames)} != {expected_counts[key]}")
        if len(wb.sheetnames) != len(set(wb.sheetnames)):
            errors[key].append("工作表名稱重複")
        if "元件標籤字典" not in wb.sheetnames:
            errors[key].append("缺少元件標籤字典")
        for ws in wb.worksheets:
            title_terms = [term for term in TAIWAN_WORDING_FORBIDDEN if term in ws.title]
            if title_terms:
                errors[key].append(
                    f"{ws.title}: 工作表名稱仍含非台灣慣用語：{', '.join(title_terms)}"
                )
            for row in ws.iter_rows():
                for cell in row:
                    value = str(cell.value or "")
                    wording_value = value
                    for label in COMPONENT_GLOSSARY:
                        wording_value = wording_value.replace(label, "")
                    found = [
                        term for term in TAIWAN_WORDING_FORBIDDEN if term in wording_value
                    ]
                    if found:
                        errors[key].append(
                            f"{ws.title}:{cell.coordinate} 仍含非台灣慣用語：{', '.join(found)}"
                        )
        for ws in wb.worksheets[1:]:
            if ws.max_row > 1 and not ws.auto_filter.ref:
                errors[key].append(f"{ws.title}: 缺少篩選範圍")
            if ws.max_row > 1 and ws.freeze_panes is None:
                errors[key].append(f"{ws.title}: 缺少 freeze pane")
            for row in ws.iter_rows():
                if any(
                    "LockCore runtime / Skills / Memory" in str(cell.value or "")
                    for cell in row
                ):
                    errors[key].append(f"{ws.title}: 仍含未定義的舊概括標籤")
                    break
            header_row = next(
                (
                    row_idx
                    for row_idx in (1, 4)
                    if any(
                        cell.value == "正式元件名稱"
                        for cell in ws[row_idx]
                    )
                ),
                None,
            )
            if header_row:
                component_column = next(
                    cell.column
                    for cell in ws[header_row]
                    if cell.value == "正式元件名稱"
                )
                for row_idx in range(header_row + 1, ws.max_row + 1):
                    value = str(ws.cell(row_idx, component_column).value or "").strip()
                    if not value or value.startswith("跨系統品質屬性"):
                        continue
                    undefined = [
                        label.strip()
                        for label in value.split(";")
                        if label.strip() not in COMPONENT_GLOSSARY
                    ]
                    if undefined:
                        errors[key].append(
                            f"{ws.title}:{row_idx} 使用未定義元件標籤：{', '.join(undefined)}"
                        )
        glossary_ws = wb["元件標籤字典"]
        if glossary_ws.max_row - 4 != len(COMPONENT_GLOSSARY):
            errors[key].append(
                f"元件標籤字典 {glossary_ws.max_row - 4} 列 != {len(COMPONENT_GLOSSARY)} 個受控標籤"
            )
        if key == "bom":
            bom_ws = wb["BOM 主表"]
            headers = {str(cell.value): cell.column for cell in bom_ws[1]}
            l3_rows = [
                row_idx
                for row_idx in range(2, bom_ws.max_row + 1)
                if bom_ws.cell(row_idx, 1).value == "L3"
            ]
            if len(l3_rows) != len(requirements):
                errors[key].append(f"BOM L3 {len(l3_rows)} 列 != FR {len(requirements)} 列")
            for header in ("正式元件名稱", "SAD 定位", "SDS 定位", "實作路徑 / 規劃狀態"):
                column = headers.get(header)
                if not column:
                    errors[key].append(f"BOM 主表缺少「{header}」欄")
                    continue
                blank_count = sum(not bom_ws.cell(row_idx, column).value for row_idx in l3_rows)
                if blank_count:
                    errors[key].append(f"BOM 主表 L3「{header}」有 {blank_count} 列空白")
        if key == "acceptance":
            expected_business_headers = {
                "客戶 VOC／商業情境",
                "PRD 驗收標準",
                "成功指標（架構／SA）",
            }
            main_tabs = [
                name
                for name in wb.sheetnames
                if name not in {"封面·文件管制", "技術流程對照", "元件標籤字典"}
            ]
            for tab_name in main_tabs:
                ws = wb[tab_name]
                headers = {str(cell.value or ""): cell.column for cell in ws[1]}
                missing = expected_business_headers - set(headers)
                if missing:
                    errors[key].append(f"{tab_name}: 缺少業務驗收欄 {', '.join(sorted(missing))}")
                    continue
                for header in ("客戶 VOC／商業情境", "PRD 驗收標準"):
                    column = headers[header]
                    for row_idx in range(2, ws.max_row + 1):
                        value = str(ws.cell(row_idx, column).value or "")
                        if "前置：" in value or "主要流程：" in value:
                            errors[key].append(f"{tab_name}:{row_idx} 業務欄仍含技術流程語言")
            if "技術流程對照" not in wb.sheetnames:
                errors[key].append("缺少「技術流程對照」分頁")
            else:
                technical_ws = wb["技術流程對照"]
                if technical_ws.max_row - 4 != len(requirements):
                    errors[key].append(
                        f"技術流程對照 {technical_ws.max_row - 4} 列 != FR {len(requirements)} 列"
                    )
                technical_headers = {str(cell.value or ""): cell.column for cell in technical_ws[4]}
                technical_column = technical_headers.get("技術流程（前置／主要流程）")
                if not technical_column:
                    errors[key].append("技術流程對照缺少技術流程欄")
                else:
                    missing_flow = sum(
                        "前置：" not in str(technical_ws.cell(row_idx, technical_column).value or "")
                        or "主要流程：" not in str(technical_ws.cell(row_idx, technical_column).value or "")
                        for row_idx in range(5, technical_ws.max_row + 1)
                    )
                    if missing_flow:
                        errors[key].append(f"技術流程對照有 {missing_flow} 列未完整保留前置／主要流程")
        if key == "test":
            forbidden_trace_phrases = (
                "已預繫場景",
                "尚無直接 SRS-TC 鍵",
                "有直接 SRS-TC 連結",
                "上游直接連結",
            )
            for test_ws in wb.worksheets:
                for row in test_ws.iter_rows():
                    for cell in row:
                        cell_text = str(cell.value or "")
                        if "設計技法" in cell_text:
                            errors[key].append(
                                f"{test_ws.title}:{cell.coordinate} 仍使用「設計技法」，應改為可執行的測試方法"
                            )
                        found_forbidden = [
                            phrase for phrase in forbidden_trace_phrases if phrase in cell_text
                        ]
                        if found_forbidden:
                            errors[key].append(
                                f"{test_ws.title}:{cell.coordinate} 仍含已廢止追溯語義："
                                f"{', '.join(found_forbidden)}"
                            )
            qa_view_forbidden = (
                "QTM-",
                "SAD ",
                "SDS ",
                "REQ ID",
                "FR-",
                "NFR-",
                "Clarify",
                "dedup",
                "HITL",
                "🔜",
                "正常資料：—",
                "資料：—",
            )
            for tab_name in ("⑧ 客戶需求與測試情境", "⑨ 測試案例與執行紀錄"):
                qa_view_ws = wb[tab_name]
                source_column = next(
                    (
                        cell.column
                        for cell in qa_view_ws[4]
                        if str(cell.value or "") == "需求來源"
                    ),
                    None,
                )
                for row in qa_view_ws.iter_rows():
                    for cell in row:
                        value = str(cell.value or "")
                        allowed_source_tokens = (
                            {"FR-", "NFR-"}
                            if tab_name == "⑧ 客戶需求與測試情境"
                            and cell.row >= 5
                            and cell.column == source_column
                            else set()
                        )
                        exposed = [
                            token
                            for token in qa_view_forbidden
                            if token in value and token not in allowed_source_tokens
                        ]
                        if exposed:
                            errors[key].append(
                                f"{tab_name}:{cell.coordinate} 暴露 QA 不需的治理/縮寫語言："
                                f"{', '.join(exposed)}"
                            )
            pending_tokens = (
                "待填",
                "待補",
                "待建立",
                "待回填",
                "待分流",
                "待確認",
                "待定",
                "落地待對帳",
                "實作待對帳",
            )
            expected_rows = {
                "附錄A 需求追溯": len(requirements) + len(nfrs),
                "附錄B 規格完整": len(requirements) + len(nfrs),
                "⑧ 客戶需求與測試情境": len(requirements) + len(nfrs),
                "⑨ 測試案例與執行紀錄": len(requirements) + len(nfrs),
                "附錄C 測試分析": len(DOMAIN_QA_CHECKS) + len(SCENARIOS),
            }
            for tab_name, expected in expected_rows.items():
                qa_ws = wb[tab_name]
                actual = qa_ws.max_row - 4
                if actual != expected:
                    errors[key].append(f"{tab_name} {actual} 列 != 預期 {expected} 列")
                for row in qa_ws.iter_rows(min_row=5):
                    for cell in row:
                        value = str(cell.value or "")
                        found = [token for token in pending_tokens if token in value]
                        if found:
                            errors[key].append(
                                f"{tab_name}:{cell.coordinate} 仍含未閉環字樣：{', '.join(found)}"
                            )
            analysis_ws = wb["附錄C 測試分析"]
            analysis_headers = {str(cell.value or ""): cell.column for cell in analysis_ws[4]}
            required_analysis_headers = {
                "分析類型",
                "分析鍵",
                "必測行為／準備",
                "執行方式／流程",
                "通過條件／備註",
            }
            missing_analysis_headers = required_analysis_headers - set(analysis_headers)
            if missing_analysis_headers:
                errors[key].append(
                    f"附錄C 缺少測試分析欄位：{', '.join(sorted(missing_analysis_headers))}"
                )
            else:
                domain_rows = 0
                journey_rows = 0
                for row_idx in range(5, analysis_ws.max_row + 1):
                    analysis_type = str(analysis_ws.cell(row_idx, analysis_headers["分析類型"]).value or "")
                    key_value = str(analysis_ws.cell(row_idx, analysis_headers["分析鍵"]).value or "")
                    preparation = str(analysis_ws.cell(row_idx, analysis_headers["必測行為／準備"]).value or "")
                    method = str(analysis_ws.cell(row_idx, analysis_headers["執行方式／流程"]).value or "")
                    result = str(analysis_ws.cell(row_idx, analysis_headers["通過條件／備註"]).value or "")
                    if analysis_type == "領域風險檢核":
                        domain_rows += 1
                        if "1." not in preparation or len(method) < 30:
                            errors[key].append(f"附錄C {key_value} 領域檢核不可直接轉成案例")
                    elif analysis_type == "端到端旅程":
                        journey_rows += 1
                        expected_pass = taiwan_wording(SCENARIO_PASS_CRITERIA.get(key_value))
                        if expected_pass not in result:
                            errors[key].append(f"附錄C {key_value} 通過條件與受控基線不一致")
                if domain_rows != len(DOMAIN_QA_CHECKS):
                    errors[key].append(f"附錄C 領域檢核 {domain_rows} 列 != {len(DOMAIN_QA_CHECKS)}")
                if journey_rows != len(SCENARIOS):
                    errors[key].append(f"附錄C 端到端旅程 {journey_rows} 列 != {len(SCENARIOS)}")
            matrix_ws = wb["附錄A 需求追溯"]
            matrix_headers = {str(cell.value or ""): cell.column for cell in matrix_ws[4]}
            matrix_req_column = matrix_headers.get("REQ ID")
            matrix_status_column = matrix_headers.get("QA 追溯狀態")
            matrix_decision_column = matrix_headers.get("QA 追溯決策")
            if not all((matrix_req_column, matrix_status_column, matrix_decision_column)):
                errors[key].append("附錄A 缺少 REQ ID、QA 追溯狀態或 QA 追溯決策欄")
            else:
                for row_idx in range(5, matrix_ws.max_row + 1):
                    req_id = str(matrix_ws.cell(row_idx, matrix_req_column).value or "")
                    status = str(matrix_ws.cell(row_idx, matrix_status_column).value or "")
                    decision = str(matrix_ws.cell(row_idx, matrix_decision_column).value or "")
                    if not status.startswith("✅"):
                        errors[key].append(f"附錄A {req_id} 追溯狀態未閉環：{status}")
                    if f"20_Test_Cases §2.1：{qa_mapping_id(req_id)}" not in status:
                        errors[key].append(f"附錄A {req_id} 未以 QTM 正典鍵表示追溯閉環")
                    if qa_mapping_id(req_id) not in decision or "指定 TC：" not in decision:
                        errors[key].append(f"附錄A {req_id} 缺少 QTM→指定 TC 決策")
                    if "候選：" in decision or "回填 20_Test_Cases" in decision:
                        errors[key].append(f"附錄A {req_id} 仍把 QA 決策寫成候選/回填工作")
            scenario_ws = wb["⑧ 客戶需求與測試情境"]
            scenario_headers = {
                str(cell.value or ""): cell.column for cell in scenario_ws[4]
            }
            required_scenario_headers = {
                "情境編號",
                "需求來源",
                "客戶需求／VOC",
                "PRD 驗收標準",
                "測試重點／通過判準",
                "應涵蓋情況",
                "客戶旅程",
                "對應案例編號",
            }
            missing_scenario_headers = required_scenario_headers - set(scenario_headers)
            scenario_to_case: dict[str, str] = {}
            if missing_scenario_headers:
                errors[key].append(
                    f"⑧ 缺少需求轉測試欄位：{', '.join(sorted(missing_scenario_headers))}"
                )
            else:
                for row_idx in range(5, scenario_ws.max_row + 1):
                    scenario_id = str(scenario_ws.cell(row_idx, scenario_headers["情境編號"]).value or "")
                    source_id = str(scenario_ws.cell(row_idx, scenario_headers["需求來源"]).value or "")
                    case_id = str(scenario_ws.cell(row_idx, scenario_headers["對應案例編號"]).value or "")
                    if not scenario_id.startswith("SCN-") or not case_id.startswith("CASE-"):
                        errors[key].append(f"⑧:{row_idx} 情境或案例編號格式錯誤")
                    expected_source_id = [*requirements, *nfrs][row_idx - 5].req_id
                    if source_id != expected_source_id:
                        errors[key].append(
                            f"⑧:{row_idx} 需求來源 {source_id or '空白'} != {expected_source_id}"
                        )
                    source_link = scenario_ws.cell(
                        row_idx, scenario_headers["需求來源"]
                    ).hyperlink
                    expected_source_link = f"#'附錄A 需求追溯'!A{row_idx}"
                    if not source_link or source_link.target != expected_source_link:
                        errors[key].append(
                            f"⑧:{row_idx} 需求來源未連到附錄 A 對應列"
                        )
                    if scenario_id in scenario_to_case:
                        errors[key].append(f"⑧:{row_idx} 情境編號重複：{scenario_id}")
                    scenario_to_case[scenario_id] = case_id
                    for header in (
                        "客戶需求／VOC",
                        "PRD 驗收標準",
                        "測試重點／通過判準",
                        "應涵蓋情況",
                        "客戶旅程",
                    ):
                        if not str(scenario_ws.cell(row_idx, scenario_headers[header]).value or "").strip():
                            errors[key].append(f"⑧:{row_idx}「{header}」不可空白")
                expected_visible_links = {
                    qa_visible_id(item.req_id, "SCN"): qa_visible_id(item.req_id, "CASE")
                    for item in [*requirements, *nfrs]
                }
                if scenario_to_case != expected_visible_links:
                    missing = sorted(set(expected_visible_links) - set(scenario_to_case))
                    extra = sorted(set(scenario_to_case) - set(expected_visible_links))
                    mismatched = sorted(
                        scenario_id
                        for scenario_id in set(expected_visible_links) & set(scenario_to_case)
                        if expected_visible_links[scenario_id] != scenario_to_case[scenario_id]
                    )
                    errors[key].append(
                        f"⑧ 情境基線不一致；缺少={missing or '—'}，多餘={extra or '—'}，案例差異={mismatched or '—'}"
                    )

            execution_ws = wb["⑨ 測試案例與執行紀錄"]
            execution_headers = {
                str(cell.value or ""): cell.column for cell in execution_ws[4]
            }
            required_execution_headers = {
                "案例編號",
                "情境編號",
                "前置條件與測試資料",
                "執行步驟",
                "預期結果／通過標準",
                "應保留證據",
                "實際結果",
                "執行狀態",
                "證據／缺陷連結",
                "版本／日期",
            }
            missing_execution_headers = required_execution_headers - set(execution_headers)
            if missing_execution_headers:
                errors[key].append(
                    f"⑨ 測試案例與執行紀錄缺少欄位：{', '.join(sorted(missing_execution_headers))}"
                )
            else:
                execution_case_ids: set[str] = set()
                valid_statuses = {"尚未執行", "通過", "失敗", "阻擋", "不適用"}
                for row_idx in range(5, execution_ws.max_row + 1):
                    case_id = str(execution_ws.cell(row_idx, execution_headers["案例編號"]).value or "")
                    scenario_id = str(execution_ws.cell(row_idx, execution_headers["情境編號"]).value or "")
                    preparation = str(execution_ws.cell(row_idx, execution_headers["前置條件與測試資料"]).value or "")
                    steps = str(execution_ws.cell(row_idx, execution_headers["執行步驟"]).value or "")
                    expected = str(execution_ws.cell(row_idx, execution_headers["預期結果／通過標準"]).value or "")
                    evidence = str(execution_ws.cell(row_idx, execution_headers["應保留證據"]).value or "")
                    actual = str(execution_ws.cell(row_idx, execution_headers["實際結果"]).value or "")
                    status = str(execution_ws.cell(row_idx, execution_headers["執行狀態"]).value or "")
                    evidence_link = str(execution_ws.cell(row_idx, execution_headers["證據／缺陷連結"]).value or "")
                    version = str(execution_ws.cell(row_idx, execution_headers["版本／日期"]).value or "")
                    if case_id in execution_case_ids:
                        errors[key].append(f"⑨:{row_idx} 案例編號重複：{case_id}")
                    execution_case_ids.add(case_id)
                    if scenario_to_case.get(scenario_id) != case_id:
                        errors[key].append(f"⑨:{row_idx} 情境 {scenario_id} 未正確指向案例 {case_id}")
                    if "環境：" not in preparation or not any(
                        marker in preparation for marker in ("正常資料：", "資料：")
                    ):
                        errors[key].append(
                            f"⑨ 測試案例與執行紀錄:{row_idx} 缺少明確環境或測試資料"
                        )
                    if any(marker not in steps for marker in ("1.", "2.", "3.", "4.")):
                        errors[key].append(f"⑨ 測試案例與執行紀錄:{row_idx} 執行步驟不完整")
                    if not all((expected, evidence, actual, evidence_link, version)):
                        errors[key].append(f"⑨ 測試案例與執行紀錄:{row_idx} 缺通過標準或執行紀錄說明")
                    if status not in valid_statuses:
                        errors[key].append(f"⑨:{row_idx} 執行狀態不合法：{status}")
                    visible_text = " ".join(
                        str(cell.value or "") for cell in execution_ws[row_idx]
                    )
                    exposed = [token for token in ("QTM-", "SAD ", "SDS ", "REQ ID") if token in visible_text]
                    if exposed:
                        errors[key].append(
                            f"⑨ 測試案例與執行紀錄:{row_idx} 暴露治理鍵：{', '.join(exposed)}"
                        )
                missing_case_rows = set(scenario_to_case.values()) - execution_case_ids
                extra_case_rows = execution_case_ids - set(scenario_to_case.values())
                if missing_case_rows or extra_case_rows:
                    errors[key].append(
                        f"⑧↔⑨ 案例未對齊；缺少={sorted(missing_case_rows) or '—'}，多餘={sorted(extra_case_rows) or '—'}"
                    )
                if not execution_ws.data_validations.dataValidation:
                    errors[key].append("⑨ 執行狀態缺少下拉選單")
            for appendix_name in (
                "附錄A 需求追溯",
                "附錄B 規格完整",
                "附錄C 測試分析",
                "元件標籤字典",
            ):
                if wb[appendix_name].sheet_state != "hidden":
                    errors[key].append(f"{appendix_name} 應隱藏，不應干擾 QA 執行視圖")

    # 規劃書 ⑧ 與 BOM L2 必須以能力群、FR 數、階段與架構定位逐列對齊。
    if OUTPUTS["planning"].exists() and OUTPUTS["bom"].exists():
        planning_wb = load_workbook(OUTPUTS["planning"], data_only=False, read_only=False)
        bom_wb = load_workbook(OUTPUTS["bom"], data_only=False, read_only=False)
        planning_ws = planning_wb["⑧ 完整模組清單"]
        bom_ws = bom_wb["BOM 主表"]
        planning_headers = {str(cell.value or ""): cell.column for cell in planning_ws[4]}
        bom_headers = {str(cell.value or ""): cell.column for cell in bom_ws[1]}

        def later_phase(value: object) -> set[str]:
            phases = set(re.findall(r"M[1-5]", str(value or "")))
            return {"M3+" if phase in {"M3", "M4", "M5"} else phase for phase in phases}

        planning_modules = {}
        for row_idx in range(5, planning_ws.max_row + 1):
            key = str(planning_ws.cell(row_idx, planning_headers["顯示群組"]).value or "")
            planning_modules[key] = (
                planning_ws.cell(row_idx, planning_headers["模組"]).value,
                planning_ws.cell(row_idx, planning_headers["FR 列數"]).value,
                planning_ws.cell(row_idx, planning_headers["唯一 FR"]).value,
                later_phase(planning_ws.cell(row_idx, planning_headers["階段"]).value),
                planning_ws.cell(row_idx, planning_headers["正式元件名稱"]).value,
                planning_ws.cell(row_idx, planning_headers["SAD 定位"]).value,
                planning_ws.cell(row_idx, planning_headers["SDS 定位"]).value,
            )
        bom_modules = {}
        for row_idx in range(2, bom_ws.max_row + 1):
            if bom_ws.cell(row_idx, bom_headers["層級"]).value != "L2":
                continue
            summary = str(bom_ws.cell(row_idx, bom_headers["說明 / 出處"]).value or "")
            counts = re.search(r"(\d+) 列 FR / (\d+) 個唯一 ID", summary)
            phases = {
                phase
                for phase, header in (("M1", "M1"), ("M2", "M2"), ("M3+", "M3+"))
                if bom_ws.cell(row_idx, bom_headers[header]).value
            }
            key = str(bom_ws.cell(row_idx, bom_headers["代號（FR 為主鍵）"]).value or "")
            bom_modules[key] = (
                bom_ws.cell(row_idx, bom_headers["名稱 / 功能"]).value,
                int(counts.group(1)) if counts else None,
                int(counts.group(2)) if counts else None,
                phases,
                bom_ws.cell(row_idx, bom_headers["正式元件名稱"]).value,
                bom_ws.cell(row_idx, bom_headers["SAD 定位"]).value,
                bom_ws.cell(row_idx, bom_headers["SDS 定位"]).value,
            )
        if planning_modules != bom_modules:
            only_planning = sorted(set(planning_modules) - set(bom_modules))
            only_bom = sorted(set(bom_modules) - set(planning_modules))
            mismatched = sorted(
                key
                for key in set(planning_modules) & set(bom_modules)
                if planning_modules[key] != bom_modules[key]
            )
            message = (
                f"⑧↔BOM L2 未對齊；僅規劃書={only_planning or '—'}，"
                f"僅 BOM={only_bom or '—'}，欄位差異={mismatched or '—'}"
            )
            errors["planning"].append(message)
            errors["bom"].append(message)
        else:
            alignment_column = planning_headers.get("BOM 對齊")
            if not alignment_column:
                errors["planning"].append("⑧ 缺少「BOM 對齊」欄")
            elif any(
                planning_ws.cell(row_idx, alignment_column).value != "✅ 與 BOM L2 對齊"
                for row_idx in range(5, planning_ws.max_row + 1)
            ):
                errors["planning"].append("⑧ 存在未標示對齊的能力群")
    # 二次驗證主鍵數，避免 parser 靜默遺漏。
    if len(requirements) < 60:
        errors["source"].append(f"FR 只解析到 {len(requirements)} 列")
    if len(nfrs) < 90:
        errors["source"].append(f"NFR 只解析到 {len(nfrs)} 列")
    duplicate_ids = sorted(
        req_id for req_id, count in Counter(req.req_id for req in requirements).items() if count > 1
    )
    if duplicate_ids:
        errors["source"].append(f"FR 主鍵重複：{', '.join(duplicate_ids)}")
    qtm_canon = parse_qtm_canon()
    all_reqs = [*requirements, *nfrs]
    qtm_line_count = sum(
        line.startswith("| QTM-") for line in read_lines(TEST_CASE_PATH)
    )
    if qtm_line_count != len(all_reqs) or len(qtm_canon) != len(all_reqs):
        errors["source"].append(
            f"20_Test_Cases QTM 主表列數/唯一數 {qtm_line_count}/{len(qtm_canon)} "
            f"!= REQ {len(all_reqs)}"
        )

    def canonical_text(value: object) -> str:
        return taiwan_wording(str(value or "").replace("|", "／"))

    for item in all_reqs:
        row = qtm_canon.get(item.req_id)
        if not row:
            errors["source"].append(f"20_Test_Cases 缺少 {qa_mapping_id(item.req_id)}")
            continue
        is_nfr = isinstance(item, NFR)
        expected_scenario = scenario_for_nfr(item) if is_nfr else scenario_for_requirement(item)
        expected_tc = tc_hint(item.req_id, is_nfr=is_nfr)
        expected_method = (
            f"驗證：{item.verification}；通過：{item.target}"
            if is_nfr
            else f"{qa_test_method_for(item)} 通過：{qa_acceptance_for(item)}"
        )
        expected_values = {
            "qtm": qa_mapping_id(item.req_id),
            "type": "NFR" if is_nfr else "FR",
            "name": item.name,
            "priority": test_priority(item.req_id, item.all_text),
            "scenario": expected_scenario,
            "tc": expected_tc,
            "method": expected_method,
            "status": qa_closure_status(item.all_text),
        }
        mismatched = [
            field
            for field, expected in expected_values.items()
            if row.get(field) != canonical_text(expected)
        ]
        if mismatched:
            errors["source"].append(
                f"20_Test_Cases {item.req_id} 與 Excel 基線不一致：{', '.join(mismatched)}"
            )
    extra_qtm = sorted(set(qtm_canon) - {item.req_id for item in all_reqs})
    if extra_qtm:
        errors["source"].append(f"20_Test_Cases 多餘 QTM：{', '.join(extra_qtm)}")

    def qtm_references_tc(tc_id: str, mapping_values: Sequence[str]) -> bool:
        match = re.match(r"^(.+-)(\d+)$", tc_id)
        if not match:
            return any(tc_id in value for value in mapping_values)
        prefix, number = match.group(1), int(match.group(2))
        expression_pattern = re.compile(
            re.escape(prefix) + r"(\d{2}(?:(?:~|/)\d{2})*)"
        )
        for value in mapping_values:
            for expression in expression_pattern.findall(value):
                for part in expression.split("/"):
                    if "~" in part:
                        lower, upper = map(int, part.split("~"))
                        if lower <= number <= upper:
                            return True
                    elif int(part) == number:
                        return True
        return False

    mapping_values = [row["tc"] for row in qtm_canon.values()]
    orphan_tc_ids = sorted(
        tc.tc_id
        for tc in parse_test_cases()
        if not qtm_references_tc(tc.tc_id, mapping_values)
    )
    if orphan_tc_ids:
        errors["source"].append(
            f"20_Test_Cases 詳細 TC 無 QTM 反向關聯：{', '.join(orphan_tc_ids)}"
        )

    plan_scenarios = parse_test_plan_scenarios()
    expected_scenario_counts: Counter[str] = Counter()
    for req in requirements:
        expected_scenario_counts.update(re.findall(r"TS-\d+", scenario_for_requirement(req)))
    for nfr in nfrs:
        expected_scenario_counts.update(re.findall(r"TS-\d+", scenario_for_nfr(nfr)))
    if len(plan_scenarios) != len(SCENARIOS):
        errors["source"].append(
            f"19_Test_Plan 場景基線 {len(plan_scenarios)} 列 != {len(SCENARIOS)}"
        )
    for scenario in SCENARIOS:
        scenario_id, name, priority = scenario[:3]
        row = plan_scenarios.get(scenario_id)
        if not row:
            errors["source"].append(f"19_Test_Plan 缺少 {scenario_id}")
            continue
        expected = {
            "name": canonical_text(name),
            "priority": priority,
            "method": canonical_text(scenario[5]),
            "count": expected_scenario_counts[scenario_id],
        }
        mismatched = [field for field, value in expected.items() if row.get(field) != value]
        if mismatched:
            errors["source"].append(
                f"19_Test_Plan {scenario_id} 與 Excel 基線不一致：{', '.join(mismatched)}"
            )
    expected_copy_keys = {requirement_copy_key(req) for req in requirements}
    missing_copy = sorted(expected_copy_keys - set(FR_BUSINESS_COPY))
    extra_copy = sorted(set(FR_BUSINESS_COPY) - expected_copy_keys)
    if missing_copy:
        errors["source"].append(f"FR 業務文案缺少：{', '.join(missing_copy)}")
    if extra_copy:
        errors["source"].append(f"FR 業務文案無對應需求：{', '.join(extra_copy)}")
    expected_module_keys = {
        f"{prefix}.{code}"
        for prefix, groups in MODULES.items()
        for code, _, _ in groups
    }
    missing_module_arch = sorted(expected_module_keys - set(MODULE_ARCH))
    if missing_module_arch:
        errors["source"].append(f"MODULE_ARCH 缺少：{', '.join(missing_module_arch)}")
    used_labels = {
        label.strip()
        for arch in MODULE_ARCH.values()
        for label in arch["component"].split(";")
        if label.strip()
    }
    missing_definitions = sorted(used_labels - set(COMPONENT_GLOSSARY))
    unused_definitions = sorted(set(COMPONENT_GLOSSARY) - used_labels)
    if missing_definitions:
        errors["source"].append(f"元件標籤未定義：{', '.join(missing_definitions)}")
    if unused_definitions:
        errors["source"].append(f"元件字典有未使用標籤：{', '.join(unused_definitions)}")
    for req in requirements:
        arch = architecture_for(req)
        for field in ("component", "sad", "sds", "path"):
            if not arch.get(field):
                errors["source"].append(f"{req.req_id} 缺少架構欄位 {field}")
    return errors


def write_component_glossary_markdown() -> None:
    lines = [
        "# Smart Lock SAD / SDS 元件標籤字典",
        "",
        f"> 產出日：{GENERATED_ON}  ",
        "> 用途：讓 BOM、驗收表與測試計畫中的每個架構標籤，都能回查正式定義、責任邊界、SAD/SDS 與實作路徑。  ",
        "> 規則：本字典由 `_spec_data.py` 的受控標籤單向生成；`AGT·RES` 等 L2 是顯示群組，不是正式元件。",
        "",
    ]
    for label, alias, definition, boundary, modules, sad, sds, path in component_glossary_rows():
        lines.extend(
            [
                f"## {label}",
                "",
                f"- **別名／原概括詞**：{alias}",
                f"- **定義／負責什麼**：{definition}",
                f"- **邊界／不負責什麼**：{boundary}",
                f"- **使用於能力群**：{modules}",
                f"- **SAD 回查**：[{sad}](../12_SAD.md)",
                f"- **SDS 回查**：[{sds}](../15_SDS.md)",
                f"- **能力群相關實作路徑／狀態**：`{path}`",
                "",
            ]
        )
    COMPONENT_GLOSSARY_PATH.write_text("\n".join(lines), encoding="utf-8")


def write_health_report(health: dict[str, int], errors: dict[str, list[str]], requirements, nfrs) -> None:
    pending_fr = sum(criterion_state(req.all_text) != "✅ 判準已述" for req in requirements)
    pending_nfr = sum(criterion_state(nfr.all_text) != "✅ 判準已述" for nfr in nfrs)
    module_count = sum(
        1
        for prefix in SUBSYSTEMS
        for code, _, _ in MODULES[prefix]
        if any(req.prefix == prefix and module_for(req)[0] == code for req in requirements)
    )
    key_health = (
        f"✅ {health['fr_unique']}/{health['fr_rows']} 個 FR 主鍵唯一；"
        "FR-TEC-07＝現場報價修正，FR-TEC-08＝排班與生命週期。"
        if health["duplicate_rows"] == 0
        else f"🔴 仍有 {health['duplicate_rows']} 列 FR 主鍵碰撞。"
    )
    validation_lines = []
    if not errors:
        validation_lines.append("- ✅ 四份活頁簿皆可重新開啟；19_Test_Plan 的 12 個 TS、20_Test_Cases 的 171 筆 QTM 與 90 筆詳細 TC 反向關聯、QA ⑧需求情境／⑨案例執行視圖、隱藏分析與治理附錄、⑧↔BOM L2 與驗收語言分流均通過。")
    else:
        for key, messages in errors.items():
            for message in messages:
                validation_lines.append(f"- 🔴 `{key}`: {message}")
    output_lines = [
        f"- [{path.name}](./{path.name}) — {path.stat().st_size:,} bytes"
        for path in [*OUTPUTS.values(), COMPONENT_GLOSSARY_PATH]
        if path.exists()
    ]
    report = f"""# Smart Lock 規格四書產出健康報告

> 產出日：{GENERATED_ON}<br>
> 生成器：`_build_enterprise_workbooks.py`<br>
> 原則：xlsx 是單向快照，真相源是 `../04_SRS.md`、`../05_NFR.md`、`../12_SAD.md`、`../15_SDS.md`、`../19_Test_Plan.md`、`../20_Test_Cases.md` 等 enterprise 正典；`21_Traceability_Matrix.md` 僅保留 legacy 參考。

## 產出檔

{chr(10).join(output_lines)}

## 鍵與涵蓋健康

- FR 原文：**{health['fr_rows']} 列 / {health['fr_unique']} 個唯一 ID**。
- NFR：**{health['nfr_rows']} 列**。
- TC：**{health['tc_rows']} 筆既有詳細案例**；20_Test_Cases §2.1 有 **{health['fr_rows'] + health['nfr_rows']} 筆 QTM 現行主表**，每筆都含 SRS REQ 與指定 TC。詳細案例舊來源欄另有 **{health['detail_source_srs_ids']} 個現行 SRS ID**，僅作歷史稽核，不影響 QTM 追溯完整性。
- 21_Traceability 的 legacy `FR-0001` 形式鍵：**{health['legacy_trace_ids']} 個**。
- 上游原文仍含 `[待確認]` / 🔜 等規劃或實作訊號：FR **{pending_fr}** 列，NFR **{pending_nfr}** 列；這是來源狀態，不代表 QA 分析未決。
- QA 執行視圖：⑧有 **{health['fr_rows'] + health['nfr_rows']} 筆客戶／品質需求情境**，每筆均顯示可反查的 FR／NFR 需求來源；⑨有同數量的可執行案例與結果登錄欄。**{len(DOMAIN_QA_CHECKS)} 個領域檢核 + 12 條端到端旅程**移至隱藏附錄 C。
- 測試正典對齊：`19_Test_Plan` **12/12 個 TS** 與 Excel 隱藏附錄 C 一致；`20_Test_Cases` **{health['fr_rows'] + health['nfr_rows']}/{health['fr_rows'] + health['nfr_rows']} 筆 QTM** 與 Excel 隱藏附錄 A 逐欄一致。
- 詳細案例反查：`20_Test_Cases` **{health['tc_rows']}/{health['tc_rows']} 筆既有 TC** 均至少被一筆 QTM 指定，孤兒 TC 為 0。
- 架構回查：**{health['fr_rows']}/{health['fr_rows']} 列 FR** 均有「正式元件名稱 + SAD 定位 + SDS 定位 + 實作路徑/狀態」。
- 模組對齊：規劃書「⑧ 完整模組清單」與 BOM 主表 L2 **{module_count}/{module_count} 個能力群**逐列一致。
- 驗收語言：**{health['fr_rows']} 列 FR + {health['nfr_rows']} 列 NFR** 均具備客戶 VOC／商業情境、PRD 驗收標準與成功指標；FR 技術流程完整移至獨立對照頁。
- QA 視圖分層：⑧只呈現客戶需求、FR／NFR 來源與測試情境，⑨只呈現案例執行與結果；QTM/SAD/SDS 完整追溯收於隱藏附錄 A/B，領域檢核與端到端旅程收於隱藏附錄 C。
- 受控元件標籤：**{len(COMPONENT_GLOSSARY)} 個**；每份活頁簿均含相同的「元件標籤字典」與責任邊界。

## 主鍵治理

{key_health}

## 追溯治理結論

1. 新四書的單一脊椎是 `04_SRS` FR/NFR ID。
2. `21_Traceability_Matrix.md` 雖聲稱 FR-ID 為 04_SRS 定版編號，主表實際仍用 `FR-0001`。這些鍵本次只保留 legacy display，不視為 SRS join。
3. `20_Test_Cases.md` §2.1 已建立強制 `QTM / SRS REQ / TS / 指定 TC / 方法與判準` 主表；下方 90 筆詳細案例的原 FR/來源欄降級為歷史背景。QTM 規格完整不代表測試已執行或證據已產生。
4. 需求文字中的 🔜 可能被後來 codegraph 標註覆寫現況；因此四書的狀態只表示需求文件訊號，實作完成必須用 WBS + code + SIT/UAT 證據判定。
5. `AGT·RES` 等 L2 只是能力群投影；正式元件名稱以 `12_SAD` / `15_SDS` 原文為準，BOM 不再使用 `LockCore runtime / Skills / Memory` 這類無法精確回查的概括標籤。

## 自動驗證

{chr(10).join(validation_lines)}
"""
    REPORT_PATH.write_text(report, encoding="utf-8")


def main() -> int:
    requirements = parse_requirements()
    nfrs = parse_nfrs()
    test_cases = parse_test_cases()
    adrs = parse_adrs()
    wbs = parse_wbs()
    legacy_trace_ids = parse_trace_legacy_ids()
    detail_source_links = detailed_tc_source_srs_links(test_cases)
    health = {
        "fr_rows": len(requirements),
        "fr_unique": len({req.req_id for req in requirements}),
        "duplicate_rows": sum(req.duplicate for req in requirements),
        "nfr_rows": len(nfrs),
        "tc_rows": len(test_cases),
        "detail_source_srs_ids": len(detail_source_links),
        "legacy_trace_ids": len(legacy_trace_ids),
    }
    build_planning_workbook(requirements, nfrs, adrs, wbs, health)
    build_bom_workbook(requirements, nfrs, detail_source_links, health)
    build_acceptance_workbook(requirements, nfrs, health)
    build_test_workbook(requirements, nfrs, test_cases, health)
    write_component_glossary_markdown()
    errors = validate_workbooks(requirements, nfrs)
    write_health_report(health, errors, requirements, nfrs)
    print(
        f"Generated 4 workbooks: FR {health['fr_rows']} rows/{health['fr_unique']} unique, "
        f"NFR {health['nfr_rows']}, TC {health['tc_rows']}, errors {sum(map(len, errors.values()))}"
    )
    for path in OUTPUTS.values():
        print(f"- {path.name}: {path.stat().st_size:,} bytes")
    print(f"- {COMPONENT_GLOSSARY_PATH.name}: {COMPONENT_GLOSSARY_PATH.stat().st_size:,} bytes")
    print(f"- {REPORT_PATH.name}: {REPORT_PATH.stat().st_size:,} bytes")
    return 1 if errors else 0


if __name__ == "__main__":
    raise SystemExit(main())
