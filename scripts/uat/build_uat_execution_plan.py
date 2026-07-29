#!/usr/bin/env python3
"""由四書 xlsx 生成《UAT 執行追蹤表》—— 可執行的 UAT 計畫載體。

**這支腳本不定義驗收準則**。準則、角色、KPI 門檻、缺陷分級與簽核裁定一律
來自 `smartlock-docs/enterprise/22_UAT_Report.md`（v1.1，正典）。本檔只做
22 沒有、而實際執行必須有的三件事：

1. **排程與相依** —— 22 §4 列了九支腳本但沒說誰先誰後，而它們有硬相依。
   最典型的是 UAT-08：其步驟 1 要求「同一技師在**兩品牌**皆可派工、且其中
   一筆**已接單**」，這隱含 UAT-09 要跑兩次（兩個品牌）、UAT-07 的授權要
   覆蓋兩個品牌、UAT-02 要在第二品牌再跑到接單為止。照 01→09 的編號順序
   執行，跑到 UAT-08 才會發現前置不存在。
2. **案例對照** —— 130 支案例經 SC 對映到腳本，走查時才知道該勾哪幾支。
3. **孤兒案例收容** —— 130 支裡有 34 支不隸屬任何 SC（NFR / RBAC / 效能 /
   注入 / a11y 等橫切關注點），走查腳本天生涵蓋不到。不另立一組收容，
   「九支腳本全綠」會被讀成「全部驗完」，而其中 20 支是 P0。

用法：
    .venv/bin/python scripts/uat/build_uat_execution_plan.py
    .venv/bin/python scripts/uat/build_uat_execution_plan.py --out <path.xlsx>

來源（唯讀，絕不寫回）：
    smartlock-docs/enterprise/規格統控整理/SmartLock_整合測試計畫.xlsx
    smartlock-docs/enterprise/規格統控整理/SmartLock_業務邏輯驗收控制表.xlsx
"""
from __future__ import annotations

import argparse
import re
import sys
from collections import defaultdict
from pathlib import Path

try:
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError:  # pragma: no cover
    sys.exit("需要 openpyxl：.venv/bin/pip install openpyxl")

ROOT = Path(__file__).resolve().parents[2]
CANON = ROOT / "smartlock-docs" / "enterprise" / "規格統控整理"
WB_TEST = CANON / "SmartLock_整合測試計畫.xlsx"
WB_ACC = CANON / "SmartLock_業務邏輯驗收控制表.xlsx"
DEFAULT_OUT = ROOT / "docs" / "uat" / "SmartLock_UAT執行追蹤表.xlsx"

SHEET_UAT = "⑤ UAT 走查腳本（UAT-01–UAT-09）"
SHEET_TC = "② 測試案例主表"
SHEET_SC = "② 旅程驗收主表"
SHEET_PERSONA = "④ 旅程 × Persona（誰在走）"

# ── 樣式 ──────────────────────────────────────────────────────────────────
H_FILL = PatternFill("solid", fgColor="1F3864")
H_FONT = Font(color="FFFFFF", bold=True, size=10)
WAVE_FILL = PatternFill("solid", fgColor="D9E2F3")
WARN_FILL = PatternFill("solid", fgColor="FFF2CC")
BLOCK_FILL = PatternFill("solid", fgColor="FCE4EC")
WRAP = Alignment(vertical="top", wrap_text=True)
TOP = Alignment(vertical="top")

# ── 波次定義 ───────────────────────────────────────────────────────────────
# 這是本檔唯一的「人為判斷」部分：相依關係由各腳本步驟文字推導，理由逐條寫明，
# 讓讀者能反駁而不是只能相信。
WAVES: list[dict] = [
    {
        "wave": "W0",
        "name": "前置閘（不走查）",
        "items": ["22 §3 #1 DB migration 同步", "22 §3 #9 服務憑證與 secret", "22 §3 #12 金額與費率簽核"],
        "depends": "—",
        "why": "#12 未簽核就跑 UAT-02／UAT-04＝拿草稿費率驗金額，事後費率轉正必須整段重跑。"
               "#1 未過則任何資料面異常都無法區分是程式錯還是 schema 沒同步。",
        "parallel": "三項可平行",
    },
    {
        "wave": "W1",
        "name": "開站",
        "items": ["UAT-09（品牌 A）", "UAT-09（品牌 B）"],
        "depends": "W0",
        "why": "UAT-08 步驟 1 要求同一技師在**兩品牌**皆可派工，故 UAT-09 必須跑兩次。"
               "只開一個品牌，跑到 W5 才會發現 UAT-08 無法開始。",
        "parallel": "兩品牌可平行，但品牌 B 只需跑到可登入＋LINE 綁定",
    },
    {
        "wave": "W2",
        "name": "主體就位",
        "items": ["UAT-05（品牌 A）", "UAT-07（授權至品牌 A＋B）"],
        "depends": "W1",
        "why": "UAT-07 的品牌授權必須同時覆蓋 A 與 B，否則 W5 的 UAT-08 沒有「兩品牌皆可派工」的前提。"
               "UAT-05 只需品牌 A。",
        "parallel": "兩支可平行（不同角色、不同資料面）",
    },
    {
        "wave": "W3",
        "name": "進線",
        "items": ["UAT-01（品牌 A）"],
        "depends": "W2",
        "why": "產出下游要吃的兩種狀態：confirmed 問題卡（給 UAT-02）與 escalated 對話（給 UAT-06）。"
               "須刻意保留至少一則未解除接管的對話。",
        "parallel": "單線",
    },
    {
        "wave": "W4",
        "name": "主流程",
        "items": ["UAT-02（品牌 A，完整）", "UAT-02'（品牌 B，只跑到技師接單）", "UAT-06"],
        "depends": "W3",
        "why": "UAT-02 吃 W3 的問題卡與 W2 的技師。UAT-02' 是 UAT-08 的前置（需要一筆已接單工單）"
               "而非獨立驗收，跑到 assigned 即可。UAT-06 吃 W3 保留的 escalated 對話。",
        "parallel": "三者可平行（不同品牌／不同對話）",
    },
    {
        "wave": "W5",
        "name": "合規與撤銷",
        "items": ["UAT-04", "UAT-08"],
        "depends": "W4",
        "why": "UAT-04 的退款／取消費點測需要 W4 產出的已收款、已完工工單。"
               "UAT-08 需要 W2 的雙品牌授權＋W4' 的已接單工單，缺一不可。",
        "parallel": "兩支可平行",
    },
    {
        "wave": "W6",
        "name": "知識螺旋",
        "items": ["UAT-03"],
        "depends": "W3、W4",
        "why": "SOP draft 由客服回饋觸發，需要 W3／W4 累積的真實對話與工單當素材，"
               "空庫跑不出有意義的雙審與家族覆核。",
        "parallel": "單線（家族稽核員為外部角色，需預約）",
    },
    {
        "wave": "W7",
        "name": "受控驗收（不走查）",
        "items": ["SC-09 月結與七帳本對帳", "SC-13 技師跨品牌對帳與佣金請領", "UAT-10 橫切受控驗收"],
        "depends": "W5、W6",
        "why": "SC-09／SC-13 宣告 uat: null——沒有人能「走一遍」給你看，改由案例與證據直接判定。"
               "UAT-10 收容 34 支不隸屬任何 SC 的橫切案例（NFR／RBAC／效能／注入／a11y）。",
        "parallel": "UAT-10 可與 SC-09／SC-13 平行",
        "blocked": "SC-09／SC-13 已知阻斷：BR-SETTLE-05 對帳閘門設計上無法通過"
                   "（閘門查 legacy public.settlements、月結讀寫 v2 saas.*，且只有 legacy "
                   "approve_reconciliation 會發 commission.accrued）。CR-0189 §8 該題仍待業主裁決，"
                   "未裁決前這兩條不應排入正式輪次，排了也只會卡在同一處。",
    },
]

# 腳本 → 該腳本失敗時需連帶重跑的下游（回歸範圍）
REGRESSION: dict[str, str] = {
    "UAT-09": "全部——品牌開站失敗則其下所有旅程都沒有可用租戶",
    "UAT-05": "無下游（配置治理不餵資料給其他腳本）",
    "UAT-07": "UAT-02、UAT-02'、UAT-08（技師身分是這三者的前置）",
    "UAT-01": "UAT-02、UAT-06、UAT-03（問題卡／對話是其素材）",
    "UAT-02": "UAT-04、SC-09、SC-13（金流與對帳吃其產出）",
    "UAT-02'": "UAT-08",
    "UAT-06": "無下游",
    "UAT-04": "SC-09（退款／取消影響帳本）",
    "UAT-08": "無下游",
    "UAT-03": "UAT-01（SOP 發布後 agent 行為改變，需重測知識類回答）",
}


# ── 讀取來源 ───────────────────────────────────────────────────────────────
def _rows(path: Path, sheet: str) -> list[tuple]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        return list(wb[sheet].iter_rows(values_only=True))
    finally:
        wb.close()


def load_scripts() -> tuple[dict, dict]:
    """回 (scripts, sc2uat)。scripts[uat] = {name, scs, n_sc, n_tc, steps[], accept}"""
    rows = _rows(WB_TEST, SHEET_UAT)
    scripts, sc2uat = {}, {}
    for r in rows[1:]:
        if not r or not r[0]:
            continue
        uat = str(r[0]).strip()
        if not re.fullmatch(r"UAT-\d+", uat):
            continue  # 「不走 UAT 走查」分隔列與「—」列另行處理
        scs = re.findall(r"SC-\d+", str(r[2] or ""))
        for sc in scs:
            sc2uat[sc] = uat
        steps = [s.strip() for s in str(r[5] or "").split("\n") if s.strip()]
        scripts[uat] = {
            "name": str(r[1] or ""), "scs": scs, "n_sc": r[3], "n_tc": r[4],
            "steps": steps, "accept": str(r[6] or ""),
        }
    # 「—」列（SC-09/SC-13）也要記進 sc2uat，標為受控驗收
    for r in rows[1:]:
        if r and str(r[0] or "").strip() == "—":
            for sc in re.findall(r"SC-\d+", str(r[2] or "")):
                sc2uat[sc] = "受控驗收"
    return scripts, sc2uat


def load_cases() -> list[dict]:
    rows = _rows(WB_TEST, SHEET_TC)
    hdr = list(rows[0])
    idx = {h: i for i, h in enumerate(hdr) if h}
    out = []
    for r in rows[1:]:
        if not r or not r[idx["TC ID"]]:
            continue
        out.append({
            "tc": str(r[idx["TC ID"]]),
            "chapter": str(r[idx["章節"]] or ""),
            "pre": str(r[idx["前置"]] or ""),
            "step": str(r[idx["步驟"]] or ""),
            "expect": str(r[idx["預期結果（判定基準）"]] or ""),
            "prio": str(r[idx["優先級"]] or ""),
            "aspect": str(r[idx["驗證面向"]] or ""),
            "path": str(r[idx["路徑類型"]] or ""),
            "req": str(r[idx["驗證哪些需求"]] or ""),
            "scs": re.findall(r"SC-\d+", str(r[idx["屬於哪條旅程腳本"]] or "")),
        })
    return out


def load_sc_meta() -> dict:
    rows = _rows(WB_ACC, SHEET_SC)
    hdr = list(rows[0])
    idx = {h: i for i, h in enumerate(hdr) if h}
    out = {}
    for r in rows[1:]:
        if not r or not r[0]:
            continue
        out[str(r[0])] = {
            "name": str(r[idx["旅程"]] or ""),
            "criteria": str(r[idx["PM 驗收標準（完成判定）"]] or ""),
            "fail": str(r[idx["什麼情況算失敗"]] or ""),
        }
    return out


def load_personas() -> dict:
    rows = _rows(WB_ACC, SHEET_PERSONA)
    hdr = list(rows[0])
    idx = {h: i for i, h in enumerate(hdr) if h}
    out = defaultdict(list)
    for r in rows[1:]:
        if not r or not r[0]:
            continue
        p = str(r[idx["角色"]] or "")
        if p and p not in out[str(r[0])]:
            out[str(r[0])].append(p)
    return out


# ── 輸出 ──────────────────────────────────────────────────────────────────
def _head(ws, headers: list[str], widths: list[int]) -> None:
    ws.append(headers)
    for c, (h, w) in enumerate(zip(headers, widths), start=1):
        cell = ws.cell(row=1, column=c)
        cell.fill, cell.font, cell.alignment = H_FILL, H_FONT, WRAP
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.freeze_panes = "A2"


def sheet_readme(wb, scripts, cases, orphans) -> None:
    ws = wb.create_sheet("① 怎麼用這本")
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 118
    rows = [
        ("UAT 執行追蹤表", ""),
        ("", ""),
        ("這本是什麼", "UAT 的**執行載體**。準則、角色、KPI 門檻、缺陷分級、簽核裁定一律以 "
                    "smartlock-docs/enterprise/22_UAT_Report.md（v1.1）為準，本表不重述也不得與之牴觸。"),
        ("這本補什麼", "22 §4 列了九支腳本卻沒說執行順序，而它們有硬相依；130 支案例沒有「這支走查該勾哪幾支」的對照；"
                    "另有 34 支案例不隸屬任何旅程，走查天生涵蓋不到。本表補這三件事。"),
        ("", ""),
        ("② 執行排程", "八個波次與其相依理由。**照編號 01→09 跑會在 UAT-08 卡住**——它要求同一技師在兩品牌"
                    "皆可派工且其中一筆已接單，前置橫跨 UAT-09/07/02。"),
        ("③ 腳本逐步", f"九支腳本 {sum(len(s['steps']) for s in scripts.values())} 個步驟逐列展開，一列一步、可勾選。"),
        ("④ 案例執行表", f"{len(cases)} 支案例，標明所屬走查腳本與波次區間（最早可驗～最晚須驗完）。無旅程歸屬者標 UAT-10。"),
        ("⑤ 回歸矩陣", "某支腳本判 Fail 時，需連帶重跑哪些下游。修完只重跑自己＝下游仍在用髒資料驗收。"),
        ("⑥ 缺陷登記", "P0/P1/P2 分級定義見 22 §6。本表只登記，不改變裁定規則。"),
        ("⑦ 簽核", "角色與簽核權見 22 §2/§7。本表逐腳本列出該由誰簽。"),
        ("", ""),
        ("⚠ 已知阻斷", "SC-09 月結對帳／SC-13 技師跨品牌對帳目前**設計上無法通過**（BR-SETTLE-05）："
                    "閘門查 legacy public.settlements，月結讀寫 v2 saas.*，且只有 legacy approve_reconciliation "
                    "會發 commission.accrued。CR-0189 §8 該題待業主裁決，未裁決前不應排入正式輪次。"),
        ("⚠ 涵蓋落差", f"{len(orphans)} 支案例（{sum(1 for c in orphans if c['prio']=='P0')} 支 P0）不隸屬任何 SC，"
                    "含 RBAC、跨租戶、注入、稽核鏈、隱私。九支走查全綠 ≠ 全部驗完，必須另跑 UAT-10。"),
        ("", ""),
        ("重新生成", ".venv/bin/python scripts/uat/build_uat_execution_plan.py"),
        ("不要手改", "本表由四書 xlsx 生成。手改會在下次重生時消失；要改內容請改四書或改生成器。"
                  "**執行結果欄例外**——那是執行時填的，重生前請先另存。"),
    ]
    for a, b in rows:
        ws.append([a, b])
    for r in range(1, ws.max_row + 1):
        ws.cell(row=r, column=1).font = Font(bold=True, size=10)
        ws.cell(row=r, column=2).alignment = WRAP
    ws["A1"].font = Font(bold=True, size=14)


def sheet_schedule(wb) -> None:
    ws = wb.create_sheet("② 執行排程")
    _head(ws, ["波次", "名稱", "本波項目", "相依", "為什麼是這個順序", "可否平行",
               "預定日", "實際完成", "狀態"], [7, 16, 34, 10, 62, 26, 11, 11, 10])
    for w in WAVES:
        ws.append([w["wave"], w["name"], "\n".join(w["items"]), w["depends"],
                   w["why"], w["parallel"], "", "", "☐"])
        r = ws.max_row
        for c in range(1, 10):
            ws.cell(row=r, column=c).alignment = WRAP
        ws.cell(row=r, column=1).fill = WAVE_FILL
        if w.get("blocked"):
            ws.append(["", "", "⚠ 已知阻斷", "", w["blocked"], "", "", "", ""])
            rr = ws.max_row
            for c in range(1, 10):
                ws.cell(row=rr, column=c).fill = BLOCK_FILL
                ws.cell(row=rr, column=c).alignment = WRAP


def sheet_steps(wb, scripts, sc_meta, personas) -> None:
    ws = wb.create_sheet("③ 腳本逐步")
    _head(ws, ["腳本", "名稱", "涵蓋旅程", "參與角色", "步驟#", "步驟內容",
               "驗收點", "結果", "執行日", "執行人", "缺陷 ID"],
          [9, 30, 20, 26, 7, 62, 46, 8, 11, 10, 11])
    for uat in sorted(scripts):
        s = scripts[uat]
        roles = sorted({p for sc in s["scs"] for p in personas.get(sc, [])})
        for i, step in enumerate(s["steps"], start=1):
            ws.append([uat if i == 1 else "", s["name"] if i == 1 else "",
                       "、".join(s["scs"]) if i == 1 else "",
                       "\n".join(roles) if i == 1 else "",
                       i, step, s["accept"] if i == 1 else "", "☐", "", "", ""])
            r = ws.max_row
            for c in range(1, 12):
                ws.cell(row=r, column=c).alignment = WRAP
            if i == 1:
                ws.cell(row=r, column=1).fill = WAVE_FILL


def sheet_cases(wb, cases, sc2uat, wave_of) -> None:
    ws = wb.create_sheet("④ 案例執行表")
    _head(ws, ["TC ID", "走查腳本", "最早可驗", "最晚須驗完", "優先級", "驗證面向",
               "路徑類型", "章節", "前置", "步驟", "預期結果（判定基準）", "對映旅程",
               "驗證需求", "結果", "執行日", "執行人", "缺陷 ID"],
          [21, 20, 9, 11, 8, 10, 13, 26, 26, 44, 52, 15, 22, 8, 11, 10, 11])
    for c in sorted(cases, key=lambda x: (x["chapter"], x["tc"])):
        uats = sorted({sc2uat.get(sc, "?") for sc in c["scs"]}) or ["UAT-10"]
        u = "、".join(uats)
        # 跨多支腳本的案例給區間而非單一值：
        #   最早可驗 = 其中最早那支腳本所在波次（該波跑到時「可以」順手勾）
        #   最晚須驗完 = 最後那支所在波次（到此波所有前置齊備，是不可錯過的期限）
        # 只給最早會在前置未齊時誤導執行者；只給最晚會讓排程無謂拖後。
        waves = [wave_of.get(x, "W7") for x in uats]
        ws.append([c["tc"], u, min(waves), max(waves), c["prio"], c["aspect"],
                   c["path"], c["chapter"], c["pre"], c["step"], c["expect"],
                   "、".join(c["scs"]) or "—", c["req"], "☐", "", "", ""])
        r = ws.max_row
        for col in range(1, 18):
            ws.cell(row=r, column=col).alignment = WRAP
        if not c["scs"]:
            ws.cell(row=r, column=2).fill = WARN_FILL


def sheet_regression(wb) -> None:
    ws = wb.create_sheet("⑤ 回歸矩陣")
    _head(ws, ["若這支 Fail", "需連帶重跑", "說明"], [14, 62, 46])
    for k, v in REGRESSION.items():
        ws.append([k, v, "修完只重跑自己＝下游仍在用修復前產生的資料驗收，等於沒驗。"])
        for c in range(1, 4):
            ws.cell(row=ws.max_row, column=c).alignment = WRAP


def sheet_defects(wb) -> None:
    ws = wb.create_sheet("⑥ 缺陷登記")
    _head(ws, ["缺陷 ID", "發現腳本", "發現案例", "級別(P0/P1/P2)", "現象",
               "預期", "影響旅程", "狀態", "負責人", "修復 commit", "重測結果"],
          [13, 12, 16, 15, 50, 44, 18, 10, 10, 15, 12])
    ws.append(["（範例）UAT-D-001", "UAT-02", "TC-WO-03", "P0",
               "完工缺 serial 仍可結案", "應被硬閘擋下", "SC-06", "open", "", "", ""])
    for c in range(1, 12):
        ws.cell(row=2, column=c).alignment = WRAP
        ws.cell(row=2, column=c).font = Font(italic=True, color="888888")


def sheet_signoff(wb, scripts) -> None:
    ws = wb.create_sheet("⑦ 簽核")
    _head(ws, ["腳本", "名稱", "簽核角色（依 22 §2）", "裁定", "附帶條件", "簽名", "日期"],
          [9, 32, 34, 12, 40, 12, 12])
    # 22 §2 明列的旅程簽核歸屬
    owners = {
        "UAT-01": "CS team（客服主管）／業主",
        "UAT-02": "品牌小編、師傅代表／業主",
        "UAT-03": "家族稽核員（合約 4.4(d)）／業主",
        "UAT-04": "CS team、法務／DPO／業主",
        "UAT-05": "租戶 Admin／業主",
        "UAT-06": "品牌小編、CS team／業主",
        "UAT-07": "師傅代表／業主",
        "UAT-08": "師傅代表、平台管理員／業主",
        "UAT-09": "業主（平台方）",
    }
    for uat in sorted(scripts):
        ws.append([uat, scripts[uat]["name"], owners.get(uat, "業主"), "", "", "", ""])
        for c in range(1, 8):
            ws.cell(row=ws.max_row, column=c).alignment = WRAP
    ws.append(["UAT-10", "橫切受控驗收（34 支無旅程案例）", "Release Manager／業主", "", "", "", ""])
    ws.append(["受控驗收", "SC-09 月結對帳、SC-13 技師跨品牌對帳", "業主（⚠ BR-SETTLE-05 未裁決前不可簽）", "", "", "", ""])
    for c in range(1, 8):
        ws.cell(row=ws.max_row, column=c).fill = BLOCK_FILL
        ws.cell(row=ws.max_row, column=c).alignment = WRAP


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", type=Path, default=DEFAULT_OUT)
    args = ap.parse_args()

    for p in (WB_TEST, WB_ACC):
        if not p.exists():
            sys.exit(f"FAIL: 找不到來源 {p}")

    scripts, sc2uat = load_scripts()
    cases = load_cases()
    sc_meta = load_sc_meta()
    personas = load_personas()
    orphans = [c for c in cases if not c["scs"]]

    # 波次對照（供 ④ 標注）
    wave_of: dict[str, str] = {}
    for w in WAVES:
        for it in w["items"]:
            m = re.match(r"(UAT-\d+)", it)
            if m:
                wave_of.setdefault(m.group(1), w["wave"])
    wave_of["UAT-10"] = "W7"
    wave_of["受控驗收"] = "W7"

    # 自我驗證：實算案例數必須對得上 ⑤ 宣告值，否則來源已漂移
    bucket: dict[str, set] = defaultdict(set)
    for c in cases:
        for sc in c["scs"]:
            bucket[sc2uat.get(sc, "?")].add(c["tc"])
    drift = [f"{u}: 宣告 {s['n_tc']} 實算 {len(bucket.get(u, set()))}"
             for u, s in scripts.items() if s["n_tc"] != len(bucket.get(u, set()))]
    if drift:
        sys.exit("FAIL: 案例數與 ⑤ 宣告值不一致（來源已漂移，先查四書）：\n  " + "\n  ".join(drift))

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    sheet_readme(wb, scripts, cases, orphans)
    sheet_schedule(wb)
    sheet_steps(wb, scripts, sc_meta, personas)
    sheet_cases(wb, cases, sc2uat, wave_of)
    sheet_regression(wb)
    sheet_defects(wb)
    sheet_signoff(wb, scripts)

    args.out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(args.out)

    try:  # --out 可指到 repo 外（暫存驗證用），此時印絕對路徑
        shown = args.out.resolve().relative_to(ROOT)
    except ValueError:
        shown = args.out.resolve()
    print(f"✅ 已生成 {shown}")
    print(f"   走查腳本 {len(scripts)} 支 / 步驟 {sum(len(s['steps']) for s in scripts.values())} 個")
    print(f"   案例 {len(cases)} 支（有旅程 {len(cases) - len(orphans)} / 無旅程 {len(orphans)}，"
          f"其中 P0 {sum(1 for c in orphans if c['prio'] == 'P0')} 支）")
    print(f"   波次 {len(WAVES)} 個；案例數與 ⑤ 宣告值逐支一致")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
