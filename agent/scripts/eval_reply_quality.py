"""平均抽題目跑 LockCore 客服 agent,評估回覆品質。

從 `agent/tests/AI_Blue_鎖匠AI客服_987題核心訓練集_v1_20260602.xlsx` 依「分類」分層抽樣,
每類抽 N 題送進真實 LockCore turn,再用 LLM-as-judge 比對標準答案打 5 維分數,輸出 CSV/JSON。

評分維度(每題 0.0/0.5/1.0):
    - intent_match       — agent 是否抓對意圖
    - key_info_coverage  — 標準答案核心要點命中率
    - followup_correct   — 缺資料追問是否符合規範
    - escalation_correct — 該轉真人/維持 AI 的判斷正確
    - safety_ok          — 無越界承諾(報價、可行性、保證)

用法:
    cd agent
    pip install -e ".[eval]"                                        # 第一次裝 openpyxl
    .venv/bin/python scripts/eval_reply_quality.py --per-category 3 # 每類 3 題 (≈84 題)
    .venv/bin/python scripts/eval_reply_quality.py --total 50       # 依類別比例共 50 題
    .venv/bin/python scripts/eval_reply_quality.py --per-category 2 \\
        --output evals/reply_quality_$(date +%Y%m%d).csv

跑真 LLM API,不進 pytest。
"""

from __future__ import annotations

import argparse
import asyncio
import csv
import json
import random
import sys
import tempfile
from collections import defaultdict
from dataclasses import asdict, dataclass, field
from pathlib import Path
from typing import Any

import openpyxl

# 讓 `python scripts/eval_reply_quality.py` 直接執行也找得到 lockcore
ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from lockcore.agent.loop import AgentLoop  # noqa: E402
from lockcore.app_config import (  # noqa: E402
    CS_TOOL_ALLOWLIST,
    build_escalation_store,
    build_memory_manager,
    build_provider,
    load_config,
)
from lockcore.bus.events import InboundMessage  # noqa: E402
from lockcore.bus.queue import MessageBus  # noqa: E402

DEFAULT_XLSX = ROOT / "tests" / "AI_Blue_鎖匠AI客服_987題核心訓練集_v1_20260602.xlsx"
SHEET_NAME = "01 核心訓練集987"
HEADER_ROW = 4
DATA_START_ROW = 5

# Excel column(1-based) → 欄位
COLS = {
    "core_id": 1,                # 核心題ID
    "category": 3,               # 分類
    "group_id": 4,               # 題組ID
    "user_question": 7,          # 使用者問法
    "intent": 9,                 # 意圖
    "standard_answer": 10,       # 標準答案
    "required_info": 11,         # 必抓資訊
    "followup_when_missing": 12, # 缺資料追問
    "transfer_to_human": 14,     # 是否轉真人
    "safety_boundary": 20,       # 安全邊界
}


@dataclass
class QAItem:
    core_id: str
    category: str
    group_id: str
    user_question: str
    intent: str
    standard_answer: str
    required_info: str
    followup_when_missing: str
    transfer_to_human: str
    safety_boundary: str


@dataclass
class EvalResult:
    core_id: str
    category: str
    user_question: str
    standard_answer: str
    agent_reply: str
    triggered_transfer: bool
    intent_match: float = 0.0
    key_info_coverage: float = 0.0
    followup_correct: float = 0.0
    escalation_correct: float = 0.0
    safety_ok: float = 0.0
    overall: float = 0.0
    judge_comment: str = ""
    error: str = ""


def load_qa_items(xlsx_path: Path) -> list[QAItem]:
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    if SHEET_NAME not in wb.sheetnames:
        raise RuntimeError(f"題庫缺少 sheet: {SHEET_NAME}; 現有: {wb.sheetnames}")
    ws = wb[SHEET_NAME]
    items: list[QAItem] = []
    for r in range(DATA_START_ROW, ws.max_row + 1):
        category = ws.cell(r, COLS["category"]).value
        question = ws.cell(r, COLS["user_question"]).value
        if not category or not question:
            continue
        items.append(
            QAItem(
                core_id=str(ws.cell(r, COLS["core_id"]).value or ""),
                category=str(category),
                group_id=str(ws.cell(r, COLS["group_id"]).value or ""),
                user_question=str(question),
                intent=str(ws.cell(r, COLS["intent"]).value or ""),
                standard_answer=str(ws.cell(r, COLS["standard_answer"]).value or ""),
                required_info=str(ws.cell(r, COLS["required_info"]).value or ""),
                followup_when_missing=str(ws.cell(r, COLS["followup_when_missing"]).value or ""),
                transfer_to_human=str(ws.cell(r, COLS["transfer_to_human"]).value or ""),
                safety_boundary=str(ws.cell(r, COLS["safety_boundary"]).value or ""),
            )
        )
    return items


def stratified_sample(
    items: list[QAItem], per_category: int | None, total: int | None, seed: int
) -> list[QAItem]:
    """分層抽樣:有 per_category 就每類抽 N;有 total 就依類別比例分配,至少 1。"""
    rng = random.Random(seed)
    by_cat: dict[str, list[QAItem]] = defaultdict(list)
    for it in items:
        by_cat[it.category].append(it)

    sampled: list[QAItem] = []
    if total is not None:
        cats = list(by_cat.keys())
        s = sum(len(by_cat[c]) for c in cats)
        for c in cats:
            n = max(1, round(total * len(by_cat[c]) / s))
            n = min(n, len(by_cat[c]))
            sampled.extend(rng.sample(by_cat[c], n))
        rng.shuffle(sampled)
        sampled = sampled[:total]
    else:
        n_per = per_category or 3
        for xs in by_cat.values():
            sampled.extend(rng.sample(xs, min(n_per, len(xs))))
    return sampled


JUDGE_PROMPT = """你是鎖匠 AI 客服的品質評審。比對「agent 實際回覆」與「標準規範」,
對 5 個維度各打 0 / 0.5 / 1 分(0=完全不符,0.5=部分,1=完全符合)。

【使用者問題】
{question}

【意圖(標準)】{intent}
【標準答案要點】{standard_answer}
【必抓資訊】{required_info}
【缺資料追問規範】{followup}
【是否轉真人(標準)】{escalation}
【安全邊界】{safety}

【Agent 實際回覆】
{agent_reply}

【Agent 是否實際觸發 transfer_to_human】{triggered_transfer}

請輸出 JSON,只回 JSON 不加任何其他文字:
{{
  "intent_match": 0.0,
  "key_info_coverage": 0.0,
  "followup_correct": 0.0,
  "escalation_correct": 0.0,
  "safety_ok": 0.0,
  "comment": "一句話總結(60 字內)"
}}

評分指引:
- intent_match: 回覆方向是否對應【意圖(標準)】
- key_info_coverage: 是否涵蓋【標準答案要點】的核心訊息
- followup_correct: 若缺【必抓資訊】是否依【缺資料追問規範】問;不缺則給 1.0
- escalation_correct: 對照【是否轉真人(標準)】與【Agent 是否實際觸發】是否一致
- safety_ok: 有無違反【安全邊界】(承諾報價、可行性、保證等)
"""


def _content(out: Any) -> str:
    if out is None:
        return ""
    return getattr(out, "content", None) or str(out)


def _parse_judge_json(text: str) -> dict[str, Any]:
    """處理常見的 ```json fence 與夾雜文字,只抓首個 { ... }。"""
    text = (text or "").strip()
    if text.startswith("```"):
        text = text.strip("`")
        if text[:4].lower() == "json":
            text = text[4:]
        text = text.strip()
    start, end = text.find("{"), text.rfind("}")
    if 0 <= start < end:
        text = text[start : end + 1]
    try:
        return json.loads(text)
    except json.JSONDecodeError:
        return {}


async def _run_one_turn(loop: AgentLoop, tenant: str, user_id: str, question: str) -> str:
    msg = InboundMessage(channel="cli", sender_id=user_id, chat_id=user_id, content=question)
    out = await loop._process_message(msg, session_key=f"{tenant}:{user_id}")
    return _content(out)


async def _judge_one(
    provider, model: str, item: QAItem, agent_reply: str, triggered_transfer: bool
) -> dict[str, Any]:
    prompt = JUDGE_PROMPT.format(
        question=item.user_question,
        intent=item.intent,
        standard_answer=item.standard_answer,
        required_info=item.required_info,
        followup=item.followup_when_missing,
        escalation=item.transfer_to_human,
        safety=item.safety_boundary,
        agent_reply=agent_reply or "(空回覆)",
        triggered_transfer="是" if triggered_transfer else "否",
    )
    resp = await provider.chat(
        messages=[{"role": "user", "content": prompt}],
        model=model,
        max_tokens=512,
        temperature=0.0,
    )
    parsed = _parse_judge_json(resp.content or "")
    if not parsed:
        parsed = {
            "intent_match": 0.0,
            "key_info_coverage": 0.0,
            "followup_correct": 0.0,
            "escalation_correct": 0.0,
            "safety_ok": 0.0,
            "comment": f"[judge parse fail] {(resp.content or '')[:120]}",
        }
    return parsed


def _avg(xs: list[float]) -> float:
    return sum(xs) / len(xs) if xs else 0.0


def _write_report(results: list[EvalResult], out_path: Path) -> None:
    out_path.parent.mkdir(parents=True, exist_ok=True)
    if out_path.suffix.lower() == ".json":
        out_path.write_text(
            json.dumps([asdict(r) for r in results], ensure_ascii=False, indent=2),
            encoding="utf-8",
        )
        return
    fieldnames = list(asdict(EvalResult(
        core_id="", category="", user_question="", standard_answer="",
        agent_reply="", triggered_transfer=False,
    )).keys())
    with out_path.open("w", encoding="utf-8", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()
        for r in results:
            writer.writerow(asdict(r))


def _print_summary(results: list[EvalResult]) -> None:
    n_err = sum(1 for r in results if r.error)
    ok = [r for r in results if not r.error]
    print("\n=== 總體 ===")
    print(f"  題數: {len(results)}    成功: {len(ok)}    錯誤: {n_err}")
    if ok:
        print(f"  overall            : {_avg([r.overall for r in ok]):.3f}")
        for k in (
            "intent_match", "key_info_coverage", "followup_correct",
            "escalation_correct", "safety_ok",
        ):
            print(f"  {k:18s}: {_avg([getattr(r, k) for r in ok]):.3f}")

    print("\n=== 各分類 overall(由低到高) ===")
    cat_scores: dict[str, list[float]] = defaultdict(list)
    for r in ok:
        cat_scores[r.category].append(r.overall)
    for c, xs in sorted(cat_scores.items(), key=lambda x: _avg(x[1])):
        print(f"  {c:24s} n={len(xs):3d}  avg={_avg(xs):.3f}")


async def main_async(args: argparse.Namespace) -> int:
    xlsx_path = Path(args.xlsx)
    if not xlsx_path.exists():
        print(f"[ERROR] 找不到題庫: {xlsx_path}", file=sys.stderr)
        return 2

    print(f"[1/4] 讀題庫: {xlsx_path}")
    items = load_qa_items(xlsx_path)
    print(f"      共 {len(items)} 題")

    sampled = stratified_sample(items, args.per_category, args.total, args.seed)
    print(f"[2/4] 分層抽樣: {len(sampled)} 題  (seed={args.seed})")
    cat_count: dict[str, int] = defaultdict(int)
    for it in sampled:
        cat_count[it.category] += 1
    for c, n in sorted(cat_count.items(), key=lambda x: -x[1]):
        print(f"        {c}: {n}")

    print(f"[3/4] 載入 LockCore config: {args.config or '<default>'}")
    cfg = load_config(args.config)
    provider = build_provider(cfg)
    mgr = build_memory_manager(cfg, provider)
    esc = build_escalation_store(cfg)
    workspace = Path(tempfile.mkdtemp(prefix="lockcore-eval-"))
    judge_model = args.judge_model or cfg.model
    print(f"      agent model: {cfg.model}")
    print(f"      judge model: {judge_model}")
    print(f"      tenant     : {cfg.tenant}")

    print(f"[4/4] 開始評估 ({len(sampled)} 題)...")
    results: list[EvalResult] = []
    for i, item in enumerate(sampled, 1):
        # 每題獨立 user_id + 獨立 AgentLoop:避免記憶/狀態相互污染
        user_id = f"eval-{item.core_id or i}"
        loop = AgentLoop(
            bus=MessageBus(),
            provider=provider,
            workspace=workspace,
            model=cfg.model,
            memory_manager=mgr,
            memory_tenant=cfg.tenant,
            escalation_store=esc,
            tool_allowlist=CS_TOOL_ALLOWLIST,
        )
        try:
            reply = await _run_one_turn(loop, cfg.tenant, user_id, item.user_question)
            triggered = bool(esc.list_for_user(cfg.tenant, user_id, limit=1))
            judge = await _judge_one(provider, judge_model, item, reply, triggered)
            scores = {
                k: float(judge.get(k, 0.0) or 0.0)
                for k in (
                    "intent_match", "key_info_coverage", "followup_correct",
                    "escalation_correct", "safety_ok",
                )
            }
            overall = sum(scores.values()) / 5.0
            r = EvalResult(
                core_id=item.core_id,
                category=item.category,
                user_question=item.user_question,
                standard_answer=item.standard_answer,
                agent_reply=reply,
                triggered_transfer=triggered,
                **scores,
                overall=overall,
                judge_comment=str(judge.get("comment", ""))[:200],
            )
            results.append(r)
            print(
                f"  [{i:3d}/{len(sampled)}] {item.core_id:>10s} "
                f"{item.category:<18s} → overall={overall:.2f}"
            )
        except Exception as e:  # noqa: BLE001 — 單題失敗繼續跑,錯誤寫到結果裡
            results.append(
                EvalResult(
                    core_id=item.core_id,
                    category=item.category,
                    user_question=item.user_question,
                    standard_answer=item.standard_answer,
                    agent_reply="",
                    triggered_transfer=False,
                    error=f"{type(e).__name__}: {e}",
                )
            )
            print(
                f"  [{i:3d}/{len(sampled)}] {item.core_id} {item.category} → ERROR {e}",
                file=sys.stderr,
            )

    out_path = Path(args.output)
    _write_report(results, out_path)
    print(f"\n寫入結果: {out_path}")
    _print_summary(results)
    return 0


def parse_args() -> argparse.Namespace:
    p = argparse.ArgumentParser(description="平均抽題評估 LockCore agent 回覆品質")
    p.add_argument("--xlsx", default=str(DEFAULT_XLSX), help="題庫 xlsx 路徑")
    p.add_argument("--config", default=None, help="lockcore config.toml 路徑")
    grp = p.add_mutually_exclusive_group()
    grp.add_argument("--per-category", type=int, default=None, help="每類抽 N 題(預設 3)")
    grp.add_argument("--total", type=int, default=None, help="總共抽 N 題(依類別比例)")
    p.add_argument("--seed", type=int, default=42, help="random seed")
    p.add_argument("--judge-model", default=None, help="評審用 LLM,預設與 agent 同")
    p.add_argument("--output", default="evals/reply_quality.csv",
                   help="輸出 CSV 或 JSON(看副檔名)")
    args = p.parse_args()
    if args.per_category is None and args.total is None:
        args.per_category = 3
    return args


def main() -> None:
    sys.exit(asyncio.run(main_async(parse_args())))


if __name__ == "__main__":
    main()
