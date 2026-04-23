"""將 `docs/manuals/50題測試題目.xlsx` 轉成 `fixtures/golden.yaml`。

重跑方式：
    python -m agent.evals.tools.convert_xlsx

客戶後續補題只要把新列加進同一張 sheet，再跑一次即可。
"""

from __future__ import annotations

import sys
from pathlib import Path
from typing import Any

import openpyxl
import yaml

CATEGORY_MAP: dict[str, str] = {
    "硬體維修技師 (Hardware Technician) 測試案例": "hardware_technician",
    "報價與客服專員 (Sales Representative) 測試案例": "sales_rep",
    "門市與規格助理 (Store Assistant) 測試案例": "store_assistant",
    "APP 設定專家 (APP Specialist) 測試案例": "app_specialist",
    "多意圖協作測試 (Multi-Intent)": "multi_intent",
    "圍籬與領域外測試 (Guardrails)": "guardrails",
}


def _slugify_category(label: str) -> str:
    key = CATEGORY_MAP.get(label)
    if key is None:
        raise ValueError(f"未知類別：{label!r}，請更新 CATEGORY_MAP")
    return key


def _iter_cases(ws: Any) -> list[dict[str, Any]]:
    """從 worksheet 萃取測試案例，自動處理 有 ID/無 ID 兩種格式。"""
    cases: list[dict[str, Any]] = []
    extra_counter = 1
    skipped: list[tuple[int, str]] = []
    for row_idx, row in enumerate(ws.iter_rows(values_only=True)):
        if row_idx == 0:
            continue
        category_label, case_id, question, expected = row[:4]
        if not question or not expected:
            continue
        if not category_label:
            reason = f"第 {row_idx + 1} 列缺『分類』欄：{str(question)[:30]!r}"
            skipped.append((row_idx + 1, reason))
            continue
        category = _slugify_category(category_label)
        if case_id is None:
            case_id = f"X-{extra_counter:02d}"
            extra_counter += 1
        cases.append(
            {
                "id": str(case_id).strip(),
                "category": category,
                "category_label": category_label,
                "question": str(question).strip(),
                "expected": str(expected).strip(),
                "multi_intent": category == "multi_intent",
                "guardrail": category == "guardrails",
            }
        )
    if skipped:
        print("⚠️  跳過以下列（欄位不完整，請修正 xlsx）：", file=sys.stderr)
        for _, reason in skipped:
            print(f"   - {reason}", file=sys.stderr)
    return cases


def convert(xlsx_path: Path, yaml_path: Path) -> int:
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    sheet_name = next((s for s in wb.sheetnames if "題" in s), wb.sheetnames[0])
    ws = wb[sheet_name]
    cases = _iter_cases(ws)
    yaml_path.parent.mkdir(parents=True, exist_ok=True)
    with yaml_path.open("w", encoding="utf-8") as fh:
        yaml.safe_dump(
            cases,
            fh,
            allow_unicode=True,
            sort_keys=False,
            default_flow_style=False,
            width=1000,
        )
    return len(cases)


def main() -> int:
    repo_root = Path(__file__).resolve().parents[3]
    xlsx_path = repo_root / "agent" / "docs" / "manuals" / "50題測試題目.xlsx"
    yaml_path = repo_root / "agent" / "evals" / "fixtures" / "golden.yaml"
    if not xlsx_path.exists():
        print(f"找不到：{xlsx_path}", file=sys.stderr)
        return 1
    count = convert(xlsx_path, yaml_path)
    print(f"轉出 {count} 題到 {yaml_path.relative_to(repo_root)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
